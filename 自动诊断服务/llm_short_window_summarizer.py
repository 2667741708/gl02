from __future__ import annotations

import argparse
import json
import os
import time
from datetime import datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from diagnosis_queue_service import build_queue
from service_config import PROJECT_ROOT
from store import DiagnosisStore


DESCRIPTION = "Send each 12-diagnosis queue to Ollama every 5 minutes, store the summary, and export a Word report."
EPILOG = """
Examples:
  python 自动诊断服务/llm_short_window_summarizer.py --queue-id dq_20260510_185000 --write --export-docx
  python 自动诊断服务/llm_short_window_summarizer.py --latest --model chiqiong-blast-furnace:latest --dry-run
  python 自动诊断服务/llm_short_window_summarizer.py --loop --interval-minutes 5
"""


def build_prompt(queue: dict) -> dict:
    items = queue.get("diagnosis_json") or []
    lines = [
        "你是高炉炉况短时综合诊断助手。请基于最近1小时、每5分钟一次的炉况诊断队列，给出简明综合判断。",
        "要求：先给总判断，再列出主要证据、风险变化、建议关注项。不要编造队列外数据。",
        "",
        f"队列ID：{queue.get('queue_id')}",
        f"时间范围：{queue.get('queue_start_ts')} 至 {queue.get('queue_end_ts')}",
        f"队列状态：{queue.get('status')}，诊断条数：{queue.get('diagnosis_count')}/{queue.get('expected_count')}",
        "",
        "诊断队列：",
    ]
    for item in items:
        lines.append(
            f"- {item.get('diagnosis_ts')} | 主:{item.get('main_label')} {item.get('main_score')} "
            f"置信:{item.get('main_confidence')} | 次:{item.get('secondary_label')} {item.get('secondary_score')}"
        )
    return {"messages": [{"role": "user", "content": "\n".join(lines)}], "queue_id": queue.get("queue_id")}


def call_ollama(prompt: dict, model: str, base_url: str) -> str:
    body = json.dumps({"model": model, "messages": prompt["messages"], "stream": False, "think": False}, ensure_ascii=False).encode("utf-8")
    req = Request(base_url.rstrip("/") + "/api/chat", data=body, headers={"Content-Type": "application/json"})
    with urlopen(req, timeout=180) as resp:
        payload = json.loads(resp.read().decode("utf-8", errors="replace"))
    return ((payload.get("message") or {}).get("content") or payload.get("response") or "").strip()


def fallback_summary(queue: dict, *, llm_unavailable: bool = False) -> str:
    labels = [str(item.get("main_label") or "unknown") for item in queue.get("diagnosis_json") or []]
    dominant = max(set(labels), key=labels.count) if labels else "unknown"
    suffix = "大模型摘要暂不可用，已使用确定性队列摘要。" if llm_unavailable else "当前为自动摘要占位，未调用大模型。"
    return f"最近1小时短时诊断队列以 {dominant} 为主，共 {len(labels)} 条诊断。{suffix}"


def summary_degradation(exc: BaseException) -> dict:
    """Build a safe, structured record for an unavailable Ollama summary."""
    if isinstance(exc, HTTPError):
        reason_code = "ollama_http_error"
    elif isinstance(exc, URLError):
        reason_code = "ollama_connection_error"
    elif isinstance(exc, TimeoutError):
        reason_code = "ollama_timeout"
    else:
        reason_code = "ollama_response_error"
    http_status = getattr(exc, "code", None)
    return {
        "stage": "llm_summary",
        "reason_code": reason_code,
        "error_type": type(exc).__name__,
        "message": str(exc),
        "http_status": int(http_status) if isinstance(http_status, int) else None,
        "fallback_kind": "deterministic_queue_summary",
    }


def export_docx(queue: dict, summary: str, target_dir: Path) -> tuple[str, str]:
    target_dir.mkdir(parents=True, exist_ok=True)
    stamp = str(queue.get("queue_end_ts") or datetime.now()).replace(":", "").replace(" ", "_").replace("-", "")
    md_path = target_dir / f"short_window_{stamp}.md"
    docx_path = target_dir / f"short_window_{stamp}.docx"
    rows = queue.get("diagnosis_json") or []
    md_lines = ["# 短时炉况综合判断", "", summary, "", "## 12次炉况判断队列", ""]
    for item in rows:
        md_lines.append(f"- {item.get('diagnosis_ts')} | {item.get('main_label')} | {item.get('main_score')} | {item.get('main_confidence')}")
    md_path.write_text("\n".join(md_lines), encoding="utf-8")
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，无法导出 Word。请先安装 python-docx。") from exc
    doc = Document()
    doc.add_heading("短时炉况综合判断", level=1)
    doc.add_paragraph(summary)
    doc.add_heading("12次炉况判断队列", level=2)
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    hdr[0].text = "时间"
    hdr[1].text = "主炉况"
    hdr[2].text = "主分数"
    hdr[3].text = "置信度"
    hdr[4].text = "次炉况"
    for item in rows:
        cells = table.add_row().cells
        cells[0].text = str(item.get("diagnosis_ts") or "")
        cells[1].text = str(item.get("main_label") or "")
        cells[2].text = str(item.get("main_score") or "")
        cells[3].text = str(item.get("main_confidence") or "")
        cells[4].text = str(item.get("secondary_label") or "")
    doc.save(docx_path)
    return str(docx_path), str(md_path)


def summarize(
    config_path: str | None = None,
    queue_id: str = "",
    latest: bool = False,
    model: str = "",
    ollama_url: str = "",
    export_word: bool = False,
    write: bool = False,
    dry_run: bool = False,
) -> dict:
    store = DiagnosisStore(config_path)
    store.ensure_schema()
    model = model or os.getenv("BF_LLM_MODEL") or "chiqiong-blast-furnace:latest"
    ollama_url = ollama_url or os.getenv("OLLAMA_BASE_URL") or "http://10.30.220.12:11434"
    queue = store.get_diagnosis_queue(queue_id) if queue_id else None
    if queue is None and latest:
        queue = store.latest_diagnosis_queue()
    if queue is None:
        queue = build_queue(config_path, write=write and not dry_run)
    prompt = build_prompt(queue)
    degradation = None
    if dry_run:
        summary = fallback_summary(queue)
    else:
        try:
            summary = call_ollama(prompt, model, ollama_url)
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            degradation = summary_degradation(exc)
            summary = fallback_summary(queue, llm_unavailable=True)
            prompt["summary_runtime"] = {
                "mode": "fallback",
                "degradation": degradation,
            }
    docx_path = ""
    md_path = ""
    if export_word and not dry_run:
        docx_path, md_path = export_docx(queue, summary, PROJECT_ROOT / "高炉前端数据" / "data" / "reports" / "short_window")
    payload = {
        "summary_id": "sws_" + str(queue["queue_id"]),
        "queue_id": queue["queue_id"],
        "queue_hash": queue["queue_hash"],
        "model_name": model,
        "prompt_json": prompt,
        "llm_summary": summary,
        "diagnosis_queue_json": queue.get("diagnosis_json") or [],
        "docx_path": docx_path,
        "markdown_path": md_path,
        "status": "dry_run" if dry_run else ("degraded" if degradation else "ok"),
    }
    saved = store.upsert_short_window_summary(payload) if write and not dry_run else payload
    result = {"ok": True, "degraded": bool(degradation), "summary": saved}
    if degradation:
        result["degradation"] = degradation
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=DESCRIPTION, epilog=EPILOG, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--latest", action="store_true", help="Summarize the latest persisted queue.")
    parser.add_argument("--queue-id", default="", help="Queue id/hash to summarize.")
    parser.add_argument("--loop", action="store_true", help="Run forever every --interval-minutes.")
    parser.add_argument("--interval-minutes", type=int, default=5, help="Loop interval.")
    parser.add_argument("--model", default=os.getenv("BF_LLM_MODEL") or "chiqiong-blast-furnace:latest", help="Approved production Ollama model.")
    parser.add_argument("--ollama-url", default=os.getenv("OLLAMA_BASE_URL") or "http://10.30.220.12:11434", help="Ollama base URL.")
    parser.add_argument("--export-docx", action="store_true", help="Export Word document.")
    parser.add_argument("--config", default="", help="Optional config.yaml path.")
    parser.add_argument("--write", action="store_true", help="Write summary metadata to PostgreSQL.")
    parser.add_argument("--dry-run", action="store_true", help="Build prompt/report plan without calling model or writing.")
    args = parser.parse_args()
    while True:
        result = summarize(args.config or None, args.queue_id, args.latest or not args.queue_id, args.model, args.ollama_url, args.export_docx, args.write, args.dry_run)
        print(json.dumps(result, ensure_ascii=False, default=str, indent=2))
        if not args.loop:
            return 0
        time.sleep(max(30, args.interval_minutes * 60))


if __name__ == "__main__":
    raise SystemExit(main())
