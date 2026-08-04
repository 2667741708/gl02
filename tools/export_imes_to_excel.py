# -*- coding: utf-8 -*-
"""Export IMES data from 220.12 PostgreSQL raw_rows to Excel files."""
from __future__ import annotations

import json
import os
import sys
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

EXCEL_AVAILABLE = False
try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    pass

try:
    import psycopg
except ImportError:
    import psycopg2 as psycopg


OUTPUT_DIR = Path(os.environ.get("IMES_EXPORT_DIR", r"F:\bf_export_packages\imes_excel_export"))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

PGUSER = os.environ.get("GL02_PGUSER")
PGPASSWORD = os.environ.get("GL02_PGPASSWORD")
if not PGUSER or not PGPASSWORD:
    print("ERROR: GL02_PGUSER/GL02_PGPASSWORD not set")
    sys.exit(1)

CONN_INFO = {
    "host": "127.0.0.1",
    "port": 5432,
    "dbname": "bf_trend",
    "user": PGUSER,
    "password": PGPASSWORD,
    "connect_timeout": 10,
    "options": "-c statement_timeout=120000",
}

TODAY = date.today().strftime("%Y%m%d")

# Friendly names for datasets
DATASET_LABELS = {
    "bf2_batch_all_rows": "批次投料详情_合并明细",
    "bf2_batch_mining": "批次投料详情_矿批",
    "bf2_batch_coke": "批次投料详情_焦批",
    "bf2_batch_input_detail_list_page_data2": "批次投料详情_合并明细",
    "bf2_batch_input_detail_list_page_data3": "批次投料详情_矿批",
    "bf2_batch_input_detail_list_page_data4": "批次投料详情_焦批",
    "bf2_output_list_cond_data": "生产实绩_主表",
    "bf2_output_list_page_data": "生产实绩_明细",
    "bf2_output_list_page_data2": "生产实绩_附表",
    "bf2_output_list_page_ldata": "生产实绩_L类明细",
    "bf2_output_list_page_zdata": "生产实绩_Z类明细",
    "bf2_heat_lab_list_cond_data_avg2": "炉次化验结果",
    "bf2_heat_lab_all_list_cond_data_avg2_new": "炉次化验结果_全量",
    "bf2_slag_lab_list_page_data": "炉渣检验_主表",
    "bf2_slag_lab_list_insp_data2": "炉渣检验_明细2",
    "bf2_slag_lab_list_insp_data3": "炉渣检验_明细3",
    "bf2_input_list_page_data": "原料投入管理",
    "bf2_input_list_insp_data": "原料投入_检验信息",
    "bf2_dosing_scheme_list_page_data": "配料方案",
    "bf2_dosing_scheme_list_page_detail_data": "配料方案_明细",
    "bf2_bin_list_page_data": "料仓维护",
    "bf2_bin_material_list_page_data": "料仓变料",
    "bf2_bin_material_list_page_fdata": "料仓变料_历史",
    "bf2_plan_month_list_day_data": "生产日计划",
    "bf2_plan_month_list_month_data": "生产月计划",
    "bf2_input_list_mate_coef_data": "原料投入_系数",
    "bf2_input_list_mate_coef_plan_data": "原料投入_计划系数",
    "bf2_input_list_plan_data": "原料投入_计划",
    "bf2_heat_lab_list_page_data": "炉次化验_明细",
    "bf2_heat_lab_list_page_data3": "炉次化验_附表",
}


def flatten_json(obj, prefix=""):
    """Flatten nested JSON into dot-notation keys, returning OrderedDict."""
    result = OrderedDict()
    if isinstance(obj, dict):
        for k, v in obj.items():
            new_key = f"{prefix}{k}" if prefix else k
            if isinstance(v, (dict, list)):
                result.update(flatten_json(v, f"{new_key}."))
            else:
                result[new_key] = v
    elif isinstance(obj, list):
        for i, item in enumerate(obj):
            new_key = f"{prefix}{i}"
            if isinstance(item, (dict, list)):
                result.update(flatten_json(item, f"{new_key}."))
            else:
                result[new_key] = item
    else:
        result[prefix.rstrip(".")] = obj
    return result


def write_excel(rows, columns, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columns)
    for row in rows:
        ws.append([str(v) if v is not None else "" for v in row])
    for col_idx, _ in enumerate(columns, 1):
        letter = get_column_letter(col_idx)
        if len(columns) > 30:
            ws.column_dimensions[letter].width = 14
        else:
            ws.column_dimensions[letter].width = 22
    if len(rows) > 0:
        ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
    return len(rows)


def write_csv(rows, columns, filepath):
    import csv
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([str(v) if v is not None else "" for v in row])
    return len(rows)


def export_dataset(conn, dataset_key: str, file_index: int) -> int:
    """Export one dataset from raw_rows, dynamically extracting JSON fields."""
    cur = conn.cursor()

    # Count rows
    cur.execute(
        "SELECT COUNT(*) FROM bf_imes.raw_rows WHERE dataset_key = %s",
        (dataset_key,)
    )
    count = cur.fetchone()[0]
    if count == 0:
        print(f"  [{dataset_key}] 0 rows, skipped")
        return 0

    # Fetch all rows for this dataset
    cur.execute(
        """SELECT workdate, workdate2, lot, charge, prodcentercode,
                  row_key, row_json, fetched_at
           FROM bf_imes.raw_rows
           WHERE dataset_key = %s
           ORDER BY workdate, row_key""",
        (dataset_key,)
    )

    # First pass: collect all unique JSON keys
    all_keys = OrderedDict()
    all_keys["workdate"] = True
    all_keys["workdate2"] = True
    all_keys["lot"] = True
    all_keys["charge"] = True
    all_keys["prodcentercode"] = True
    all_keys["row_key"] = True

    rows_data = []
    for row in cur:
        workdate, workdate2, lot, charge, prodcentercode, row_key, row_json, fetched_at = row
        flat = OrderedDict()
        flat["workdate"] = str(workdate) if workdate else ""
        flat["workdate2"] = str(workdate2) if workdate2 else ""
        flat["lot"] = str(lot) if lot else ""
        flat["charge"] = str(charge) if charge else ""
        flat["prodcentercode"] = str(prodcentercode) if prodcentercode else ""
        flat["row_key"] = str(row_key) if row_key else ""

        if isinstance(row_json, str):
            try:
                row_json = json.loads(row_json)
            except json.JSONDecodeError:
                row_json = {"_raw": str(row_json)}

        if isinstance(row_json, dict):
            # Extract all leaf keys
            flat_json = flatten_json(row_json)
            for k, v in flat_json.items():
                flat[k] = v
                all_keys[k] = True

        rows_data.append(flat)

    cur.close()

    # Build ordered column list
    columns = list(all_keys.keys())

    # Convert rows to tuples matching column order
    rows_out = []
    for rd in rows_data:
        rows_out.append(tuple(rd.get(c, "") for c in columns))

    # Write file
    label = DATASET_LABELS.get(dataset_key, dataset_key)
    safe_label = label.replace("/", "_").replace("\\", "_")[:60]
    filename = f"{file_index:02d}_{safe_label}_{TODAY}"

    if EXCEL_AVAILABLE:
        fp = OUTPUT_DIR / f"{filename}.xlsx"
        n = write_excel(rows_out, columns, fp)
    else:
        fp = OUTPUT_DIR / f"{filename}.csv"
        n = write_csv(rows_out, columns, fp)
    print(f"  [{dataset_key}] {fp.name}: {n} rows, {len(columns)} cols")
    return n


def main():
    print(f"Excel available: {EXCEL_AVAILABLE}")
    print(f"Output dir: {OUTPUT_DIR}")

    conn = psycopg.connect(**CONN_INFO)
    cur = conn.cursor()

    # Get all datasets
    cur.execute(
        "SELECT dataset_key, COUNT(*) cnt, MIN(workdate), MAX(workdate) "
        "FROM bf_imes.raw_rows GROUP BY dataset_key ORDER BY dataset_key"
    )
    datasets = cur.fetchall()
    print(f"Found {len(datasets)} datasets in raw_rows:")
    for ds_key, cnt, dmin, dmax in datasets:
        label = DATASET_LABELS.get(ds_key, ds_key)
        print(f"  {ds_key} ({label}): {cnt} rows, {dmin} ~ {dmax}")

    total = 0
    for i, (ds_key, cnt, dmin, dmax) in enumerate(datasets, 1):
        print(f"\n--- Exporting [{i}/{len(datasets)}] {ds_key} ---")
        total += export_dataset(conn, ds_key, i)

    # Summary
    summary_path = OUTPUT_DIR / f"export_summary_{TODAY}.json"
    summary = {
        "exported_at": datetime.now().isoformat(),
        "format": "xlsx" if EXCEL_AVAILABLE else "csv",
        "total_rows_exported": total,
        "datasets": len(datasets),
        "output_dir": str(OUTPUT_DIR),
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n=== DONE ===")
    print(f"Total: {total} rows from {len(datasets)} datasets")
    print(f"Summary: {summary_path}")

    cur.close()
    conn.close()


if __name__ == "__main__":
    main()
