"""Generate the permanent 8093 intelligent-assistant repair handbook.

OPS-8093-ASSISTANT-HEALTH-GUARD-FIX-20260804
OPS-8093-KNOWLEDGE-KEYWORD-MODE-20260805
OPS-8093-ASSISTANT-AUTO-RECOVERY-20260805
ERR-8093-ASSISTANT-DEPLOYMENT-COLLISION-20260806

The document intentionally contains no passwords, tokens, or database secrets.
Run this script after the traceability records change so the DOCX stays aligned
with AGENTS.md and the checked-in operational tools.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "docs" / "8093智能助手不可用原因与正式修复手册_20260804.docx"

TITLE = "8093 智能助手不可用原因与正式修复手册"
REQUIREMENT_ID = "OPS-8093-ASSISTANT-HEALTH-GUARD-FIX-20260804"
SECONDARY_REQUIREMENT_ID = "OPS-8093-KNOWLEDGE-KEYWORD-MODE-20260805"
TERTIARY_REQUIREMENT_ID = "OPS-8093-ASSISTANT-AUTO-RECOVERY-20260805"
INCIDENT_REQUIREMENT_ID = "ERR-8093-ASSISTANT-DEPLOYMENT-COLLISION-20260806"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name: str = "宋体", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_1, instr_text, fld_char_2))
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    for style_name, size in (("Normal", 10.5), ("Title", 24), ("Subtitle", 11), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "宋体"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)

    styles["Title"].font.color.rgb = RGBColor(24, 55, 70)
    styles["Heading 1"].font.color.rgb = RGBColor(24, 83, 91)
    styles["Heading 2"].font.color.rgb = RGBColor(55, 78, 88)
    styles["Heading 3"].font.color.rgb = RGBColor(55, 78, 88)

    normal = styles["Normal"].paragraph_format
    normal.space_after = Pt(5)
    normal.line_spacing = 1.18

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(f"{TITLE}  |  {REQUIREMENT_ID}")
    set_run_font(run, size=8)
    run.font.color.rgb = RGBColor(95, 105, 110)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def add_title_page(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("10.30.220.12 · 8093 生产预览服务 · 版本 1.3")

    document.add_paragraph()
    summary = document.add_table(rows=5, cols=2)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.style = "Table Grid"
    rows = (
        ("追踪编号", f"{REQUIREMENT_ID}；{SECONDARY_REQUIREMENT_ID}；{TERTIARY_REQUIREMENT_ID}；{INCIDENT_REQUIREMENT_ID}"),
        ("结论状态", "已确认根因、已正式修复、已完成一次真实 SSE 问答验收"),
        ("正式修复时间", "守卫：2026-08-04；keyword/连接池：2026-08-05；部署互斥与页面容错：2026-08-06"),
        ("真实问答验收", "2026-08-06 唯一一次 SSE 33.9933 秒；keyword 证据 2 条；全部保护检查通过"),
        ("适用范围", "8093 智能助手无回复、中途断流、状态全绿但当次问答失败"),
    )
    for row, values in zip(summary.rows, rows):
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], "DDE9EA")
        row.cells[0].paragraphs[0].runs[0].bold = True

    document.add_paragraph()
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = notice.add_run("内部运维文档：不含密码、Token 或数据库口令；不得用反复重启代替取证。")
    set_run_font(run, size=9, bold=True)
    run.font.color.rgb = RGBColor(153, 75, 35)

    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_bullets(document: Document, items: list[str], level: int = 0) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_code_block(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F4F5")
    set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
    paragraph = cell.paragraphs[0]
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=8.5)


def add_table(document: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]], widths: tuple[float, ...] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.text = label
        set_cell_shading(cell, "315D66")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Cm(widths[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[index])
            if row_index % 2:
                set_cell_shading(cell, "F3F7F7")


def build_document() -> Document:
    document = Document()
    configure_document(document)
    add_title_page(document)

    add_heading(document, "1. 结论：这次为什么会“无法回复”")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("2026-08-06 页面显示 Failed to fetch 时，8093 端口确实没有监听；直接原因是多个部署/恢复流程重叠，反复进入受控停止—启动窗口。")
    run.bold = True
    paragraph.add_run("SCM 在 09:57～10:29 记录了多次停启，stderr 只有受控停止触发的 KeyboardInterrupt，诊断期间后端/配置哈希仍在变化；没有对应的自然崩溃堆栈。因此本轮不是 PostgreSQL PoolTimeout 复发。历史取证结论仍然是：不是 27B 模型整体崩溃。2026-08-04 的另一条故障链是旧健康守卫把单次瞬时探测失败立即升级为重启；两类问题都能切断在途 /api/qa/chat SSE，但必须依据日志分别处理。")

    add_heading(document, "2. 之前的失败是否检查到", level=1)
    document.add_paragraph("是。2026-08-04 的历史健康日志与 runner 日志已逐项对齐，确认以下 7 组失败后均紧跟 restart_service_done，并出现对应的新代理进程启动：")
    add_table(
        document,
        ("时间", "探测失败", "后续证据", "判定"),
        [
            ("20:46", "8093 TCP 短时不可达", "立即重启并重新拉起代理", "守卫切断风险"),
            ("20:47", "8093 TCP 短时不可达", "立即重启并重新拉起代理", "守卫切断风险"),
            ("20:51", "8093 TCP 短时不可达", "立即重启并重新拉起代理", "守卫切断风险"),
            ("20:56", "8093 TCP 短时不可达", "立即重启并重新拉起代理", "守卫切断风险"),
            ("20:59", "/api/automation/status 不可达", "立即重启并重新拉起代理", "守卫切断风险"),
            ("21:15", "/api/ollama/status 与 /api/automation/status 不可达", "立即重启并重新拉起代理", "守卫切断风险"),
            ("21:48", "8093 TCP 短时不可达", "立即重启并重新拉起代理", "守卫切断风险"),
        ],
        (2.0, 5.0, 5.1, 3.5),
    )
    add_bullets(
        document,
        [
            "proxy_8093.service.err.log 的主要异常为受控停止引发的 KeyboardInterrupt，以及客户端连接被切断后的 BrokenPipeError。",
            "没有发现能解释上述 7 次时间点的模型崩溃堆栈。",
            "因此，“守卫误重启切断 SSE”是本轮已经证实的直接故障链；hybrid 检索与单模型驻留的组合是另一项独立风险，不能用它替代已经取得的重启证据。该风险已于 2026-08-05 通过显式 keyword 配置消除。",
        ],
    )

    add_heading(document, "3. 正式修复方案")
    add_table(
        document,
        ("配置项", "现行值", "作用"),
        [
            ("health.failureThreshold", "3", "服务仍在 Running 时，连续三次失败才进入重启判定"),
            ("health.serviceNotRunningFailureThreshold", "1", "服务确实不在 Running 时立即恢复，不牺牲真正故障的自愈速度"),
            ("health.preRestartBackoffSeconds", "15", "达到阈值后等待 15 秒并再次探测，瞬时抖动恢复则不重启"),
            ("health.restartCooldownSeconds", "600", "重启后 10 分钟内抑制重复重启，避免重启风暴"),
            ("跨轮次状态文件", "logs/proxy_8093.health.state.json", "保存连续失败次数和最近重启时间；成功检查将计数清零"),
            ("BF_QA_KNOWLEDGE_SEARCH_MODE", "keyword", "知识库使用 PostgreSQL 关键词/词法检索，不在共享 11434 上调用 embedding"),
        ],
        (5.0, 3.2, 8.0),
    )
    document.add_paragraph("现行状态机：")
    add_numbered(
        document,
        [
            "服务为 Running 且探测成功：记录 ok，并把 consecutiveFailures 清零。",
            "服务为 Running 且第 1、2 次连续失败：只记录 degraded / restart_deferred，不重启。",
            "第 3 次连续失败：先记录 restart_backoff，等待 15 秒后重新检查。",
            "退避后恢复：记录 recovered_during_backoff，清零并退出；仍失败才允许重启。",
            "刚发生过重启且仍在 600 秒窗口内：记录 restart_suppressed_cooldown，不再重复重启。",
            "服务本身不是 Running：使用阈值 1 立即恢复，避免对真实停机等待三分钟。",
        ],
    )

    add_heading(document, "3.1 Prompt 固定前缀与知识回答链", level=2)
    add_bullets(
        document,
        [
            "固定的不是模型内部 attention score，而是稳定的公共 Prompt token 前缀：身份 → 职责与使用原则 → 安全边界 → 回答逻辑与风格 → 固定工艺规则。动态炉况、MCP/数据库结果、知识证据、对话和当前问题统一后置，使运行时有机会复用相同前缀的 attention/KV 计算。",
            "8093 和 8094 每次 /api/qa/chat 都构造同一公共 Prompt。知识库不是每问必查：原理、规则、阈值、手册等知识意图进入 RAG；实时值、趋势、统计优先走 MCP/数据库；问候、身份和安全套取跳过 RAG。",
            "两个端口现均显式为 keyword。命中的 PostgreSQL bf_assistant.rag_* 词法证据继续注入 Prompt 的【知识证据】动态区；不缓存旧炉况或旧答案，也不请求 nomic-embed-text。",
        ],
    )

    add_heading(document, "3.2 2026-08-06 部署互斥与页面容错", level=2)
    add_bullets(
        document,
        [
            "统一恢复器与新页面热部署器使用全局命名互斥锁 Global\\BFV4PreviewProxy8093Deployment；同一时间只允许一个 8093 写入流程进入停服或原子替换阶段。",
            "服务恢复固定由 tools/remote_guarded_recover_8093_service.ps1 执行：识别已知哈希、暂停并 finally 恢复守卫、等待旧进程退出、15 秒退避、最多两次启动，并保护 8768/8094/8770/11434 PID。",
            "纯前端修复固定由 tools/remote_hot_deploy_8093_fetch_resilience.ps1 原子热更新，校验页面哈希并保留备份，不停止或重启 8093。",
            "页面不再直接显示英文 Failed to fetch。只读 GET 最多重试 3 次，退避 1.5 秒/3 秒；问答 POST/SSE 永不自动重发，明确提示用户等待服务稳定后手动重试，避免重复会话和模型负载。",
            "诊断器一次采集全端口监听快照，8093 不监听时立即输出不可用端点，不再逐端口或逐接口等待；SSH banner 异常使用新的独立会话按 1.5 秒/3 秒退避重试。",
        ],
    )

    add_heading(document, "4. 为什么当前方案可行可用")
    add_table(
        document,
        ("设计点", "解决的问题", "可用性证据"),
        [
            ("连续失败阈值", "过滤单次网络/HTTP 抖动", "行为测试证明前两次失败只延后，不执行 Restart-Service"),
            ("服务停止快速恢复", "避免去抖方案拖慢真实故障恢复", "Not Running 单独使用阈值 1"),
            ("退避后二次确认", "不在瞬时恢复前误重启", "实现 recovered_during_backoff 分支"),
            ("重启冷却", "避免分钟任务形成重启风暴", "实现 600 秒 cooldown 与持久化最近重启时间"),
            ("隔离部署", "避免修复 8093 时影响其它业务", "部署前后 8093/8768/8094/8770/11434 PID 与受保护页面/配置保持"),
            ("真实 SSE 验收", "状态绿灯不能代表真正能回答", "单次请求收到 start → delta → final → done，6.5837 秒完成"),
            ("keyword 词法检索", "避免 embedding 与单 27B 争用唯一模型槽", "实际进程与默认搜索均为 keyword；知识问答准备态命中 2 条证据，23.8856 秒完成"),
            ("部署全局互斥", "避免诊断、恢复和部署同时停启 8093", "新恢复器与页面热部署器共享 Global\\BFV4PreviewProxy8093Deployment"),
            ("前端失败容错", "短暂端口窗口不再只显示英文网络错误", "GET 有界退避；POST/SSE 不自动重发；页面热更新未改变服务 PID"),
        ],
        (4.2, 5.0, 7.0),
    )
    add_bullets(
        document,
        [
            "2026-08-04 22:54 通过受控部署器上线；仅暂停并恢复 8093 健康检查任务，没有停止业务服务。",
            "部署后 22:54～23:13 的每分钟健康记录连续为 ok，consecutiveFailures=0、lastRestartAt=null。",
            "2026-08-04 23:03 只发送一次新会话短问：首个 delta 6321.2 ms，final 6583.6 ms，总计 6583.7 ms；回答非空、会话 ID 存在、验收期间没有守卫重启。",
            "2026-08-05 07:25 只重启 8093，把实际进程环境显式切为 keyword；8768、8094、8770、11434 未变化。08:44 只发送一次知识问题，准备态显示知识检索启用、意图 parameter_optimization、命中 2 条词法证据，SSE 总计 23885.6 ms。",
            "验证前后 /api/ps 只驻留批准的 27.8B，未通过扩大 OLLAMA_MAX_LOADED_MODELS 规避问题。",
            "2026-08-06 10:41 页面容错以热更新方式上线，8093/8094/8768/8770/11434 PID 均未变化；备份为 logs/deploy_backups/8093_fetch_resilience_20260806_104046。",
            "2026-08-06 10:51 唯一一次知识 SSE 验收通过：preparing → prepared → delta → final → done，首 delta 12272.7 ms，总计 33993.3 ms，keyword 证据 2 条，问答前后全部受保护 PID、配置和守卫哈希不变。",
        ],
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "5. 再次出现不可用时的固定处理流程")
    document.add_paragraph("强制首选入口是 tools/assistant_8093_auto_recovery.py。它先完成连通性门禁和合并只读分类，再按已知哈希决定是否部署；以下单项检查是统一入口内部阶段或未知状态的人工审计依据。")
    add_numbered(
        document,
        [
            "保护现场：停止连续点击发送；先不要重启 BFV4PreviewProxy8093、BFOllama11434、BFV4PreviewWs8768，也不要清理 runner。",
            "确认链路边界：智能助手文字链路是 浏览器 → 8093 /api/qa/chat → 上下文/数据库/RAG/MCP → 11434 /api/chat → SSE；8768 正常只能证明实时炉况流正常。",
            "运行第一层只读健康探针，检查服务、监听、/api/ollama/status、/api/ps、关键配置和日志新鲜度。",
            "状态全绿但仍无回复时，运行日志尾部探针，把 POST /api/qa/chat 与 health.log 的 tcp_down、状态接口失败、restart_service_done，以及 runner 新启动时间对齐。",
            "运行进程环境探针，核对实际 BF_LLM_MODEL、OLLAMA_BASE_URL、BF_ALLOWED_LOADED_MODELS、BF_QA_KNOWLEDGE_SEARCH_MODE、OLLAMA_MAX_LOADED_MODELS 与当前驻留模型。",
            "若出现单次探测失败后立即重启，核对守卫脚本哈希、8093 配置 3/1/15/600 和状态文件；确认回退或漂移后使用受控部署器恢复，禁止手工覆盖或反复重启。",
            "若没有重启证据，才转查模型/RAG/MCP/数据库：重点检查 Timeout、embedding/nomic、psycopg、Traceback、/api/ps 驻留变化。实际进程或默认知识搜索若重新显示空值/hybrid，应判定为配置漂移并用 keyword 受控部署器恢复。",
            "修复后只做一次受控真实 SSE 问答，必须完整收到 start(preparing/prepared) → delta → final → done，并证明验收期间无守卫重启。",
            "把新根因、配置差异、命令、报告目录和回滚位置回写 AGENTS.md、error_traceability.md、automation_traceability.md 与本手册。",
        ],
    )

    add_heading(document, "6. 一键入口、内部工具与使用顺序")
    add_table(
        document,
        ("顺序", "工具", "用途 / 写入边界"),
        [
            ("0", "tools/assistant_8093_auto_recovery.py", "首选：prestage/diagnose/recover/docs；网络门禁、分类、已知最小修复、唯一 SSE 和报告"),
            ("0.1", "tools/remote_8093_assistant_diagnose.ps1", "包内只读合并采集；服务、端口、进程环境、守卫、知识、模型和近期异常一次输出"),
            ("1", "tools/probe_8093_assistant_health.ps1", "只读：服务、监听、状态接口、模型驻留和日志新鲜度"),
            ("2", "tools/probe_8093_assistant_log_tail.ps1", "只读：对齐问答请求、守卫失败、重启和异常关键字"),
            ("3", "tools/audit_22012_proxy_model_env.ps1", "只读：核对实际进程环境，不能只看配置文件"),
            ("4", "tools/probe_8093_health_guard_runtime.ps1", "只读：核对任务、脚本哈希、3/1/15/600、状态和日志"),
            ("5", "tools/remote_guarded_deploy_8093_health_guard.ps1", "写入：仅经基线校验后暂停健康任务、备份、原子部署、恢复任务和隔离验证"),
            ("6", "tools/remote_verify_8093_assistant_sse_once.ps1", "验收：包装单次真实 SSE 并保存 JSON 报告"),
            ("7", "tools/patch_8093_keyword_knowledge_mode.py", "写入准备：只允许把 8093 空值/hybrid 幂等改为 keyword；服务、端口或守卫漂移时拒绝"),
            ("8", "tools/remote_guarded_deploy_8093_keyword_knowledge_mode.ps1", "写入：只重启 8093，失败回滚，并验证实际进程与默认知识搜索均为 keyword"),
            ("9", "tools/remote_verify_8093_keyword_knowledge_sse_once.ps1", "验收：只发送一次知识问答，要求准备态证明 RAG 参与并保护守卫/PID/模型"),
            ("10", "tools/remote_guarded_recover_8093_service.ps1", "写入：全局互斥、已知哈希、15 秒退避、最多两次启动、finally 恢复守卫"),
            ("11", "tools/remote_hot_deploy_8093_fetch_resilience.ps1", "写入：全局互斥的原子页面热更新；备份/回滚，不重启服务"),
            ("12", "tools/patch_8093_assistant_fetch_resilience.py", "页面补丁：GET 有界重试；POST/SSE 不自动重发；中文可操作错误提示"),
        ],
        (1.5, 7.0, 8.0),
    )
    document.add_paragraph("本机统一入口（密码只从既有受控环境或交互输入取得）：")
    add_code_block(
        document,
        [
            "python tools\\assistant_8093_auto_recovery.py prestage --allow-agents-password",
            "python tools\\assistant_8093_auto_recovery.py diagnose --allow-agents-password --no-auto-prestage",
            "python tools\\assistant_8093_auto_recovery.py recover --allow-agents-password --no-auto-prestage",
            "python tools\\assistant_8093_auto_recovery.py docs --report logs\\assistant_8093_auto_recovery\\<recover.json>",
        ],
    )
    document.add_paragraph("远端不可变包预置在 C:\\ProgramData\\BFV4\\assistant-auto-recovery\\packages\\<package_id>，manifest 与每个 payload 都要远端 SHA-256 复核；PowerShell 文件使用 UTF-8 BOM。远端只执行短 powershell.exe -NoProfile -ExecutionPolicy Bypass -File 命令，禁止 -EncodedCommand、内联 ScriptBlock、长 -Command、逐次临时上传和嵌套转义。长任务的 exec_command 单次等待不超过 4000 ms，取得会话后短轮询。")

    add_heading(document, "7. 复发时的判定表")
    add_table(
        document,
        ("观察结果", "优先判定", "下一步"),
        [
            ("浏览器显示 Failed to fetch，且 8093 不监听，SCM 同期多次停启/文件哈希变化", "部署或恢复流程碰撞导致真实端口空窗", "停止并行写入，使用全局互斥恢复器；页面变更走无重启热部署器"),
            ("health 失败后立即 restart_service_done，runner 同期新启动", "守卫合同回退/配置漂移，或服务真实失效", "先核对 Running 状态、3/1/15/600、状态文件和哈希；符合误重启链时受控恢复守卫"),
            ("状态全绿；没有重启；SSE 仍无 delta", "8093 问答内部、数据库、RAG、MCP 或 11434 请求问题", "对齐 POST 日志与 Timeout / Traceback / psycopg / embedding 证据"),
            ("实际进程或默认知识搜索不是 keyword，或 /api/ps 出现 embedding", "知识模式/进程环境漂移，可能重新触发单驻留争用", "使用 keyword 受控部署器恢复；需要向量检索时迁移独立 embedding，不能把 MAX_LOADED_MODELS 改回 2"),
            ("服务 Not Running", "真实服务停机", "守卫应按阈值 1 立即恢复；随后查停机原因并做一次 SSE 验收"),
            ("浏览器失败但服务端完整输出 final/done", "客户端、代理链路或页面消费问题", "保留服务端报告，检查浏览器网络、控制台和 SSE 消费逻辑"),
        ],
        (5.5, 5.2, 5.8),
    )

    add_heading(document, "8. 验收与回滚标准")
    add_heading(document, "8.1 必须同时通过", level=2)
    add_bullets(
        document,
        [
            "BFV4PreviewProxy8093、BFOllama11434、BFV4PreviewWs8768 状态及对应监听正常。",
            "8093 /api/ollama/status HTTP 200，且 /api/ps 仍只有批准的 27.8B。",
            "守卫任务启用、最近结果为 0；配置为 3/1/15/600；成功后 consecutiveFailures=0。",
            "8093 实际监听进程为 BF_QA_KNOWLEDGE_SEARCH_MODE=keyword；默认知识搜索也返回 keyword 和非空证据。",
            "只提交一次真实智能助手短问，收到 start → delta → final → done，回答非空。",
            "知识问答验收的 prepared 状态必须显示知识检索启用且未跳过、存在知识意图；不能只用独立搜索接口代替问答全链路。",
            "验收期间 health.log 没有新的 restart_service_done；8093 页面正式名称为“智能助手”。",
            "8768、8094、8770、11434 没有因本修复产生无关 PID 或受保护文件变化。",
        ],
    )
    add_heading(document, "8.2 当前证据与回滚", level=2)
    add_table(
        document,
        ("项目", "当前值"),
        [
            ("共享守卫 SHA-256", "1A4B2C56D5E86BC6DCC7D82700142BF40B936492712A20AD34997953DCD9C2B3"),
            ("8093 现行配置 SHA-256", "8A24C83DF93008DD8E358682558E8755442D87052A09FB24F39778F2EAF51FDE"),
            ("守卫修复回滚目录", r"F:\高炉炼铁项目-real-sensor-v2_V4_8093_PREVIEW\logs\deploy_backups\8093_health_guard_20260804_225350"),
            ("keyword 配置回滚目录", r"F:\高炉炼铁项目-real-sensor-v2_V4_8093_PREVIEW\logs\deploy_backups\8093_knowledge_keyword_20260805_072445"),
            ("守卫 SSE 报告", r"F:\高炉炼铁项目-real-sensor-v2_V4_8093_PREVIEW\logs\acceptance\8093_health_guard_sse_20260804_20260804_230300\assistant_sse_once.json"),
            ("keyword 知识 SSE 报告", r"F:\高炉炼铁项目-real-sensor-v2_V4_8093_PREVIEW\logs\acceptance\8093_keyword_knowledge_20260805_20260805_084403\assistant_keyword_knowledge_sse_once.json"),
            ("当前一键工具远端包", "20260805_v2_72eff4ce55d3；manifest 6459A6DCEC292E27C2F3B85E30046251F35050A44241BB0F1F469A920155467E；13 个 payload 已验证"),
            ("一键工具真实报告", "本地 logs/assistant_8093_auto_recovery/20260805_110327_recover.json；远端 .../20260805_110253/assistant_keyword_knowledge_sse_once.json"),
            ("一键工具真实验收", "前后 healthy；无部署；服务阶段 184.618 秒；唯一 SSE 22.1859 秒；keyword 证据 2 条"),
            ("本地自动恢复组合测试", "22 passed（含真实 PowerShell 5.1 数字键 JSON 序列化）"),
            ("2026-08-06 页面 SHA-256", "4394D60B059CD65BDDA16D63A680F19B57880B6822C6ECE87A4190AFEAAFAC00；含 OPS-8093-QA-FETCH-RESILIENCE-20260806"),
            ("2026-08-06 本地恢复报告", "logs/assistant_8093_auto_recovery/20260806_105139_recover.json；前后 healthy；request_count=1"),
            ("2026-08-06 远端 SSE 报告", r"F:\高炉炼铁项目-real-sensor-v2_V4_8093_PREVIEW\logs\acceptance\8093_keyword_knowledge_20260805_20260806_104611\assistant_keyword_knowledge_sse_once.json"),
            ("2026-08-06 回归测试", "41 passed；覆盖页面容错、自动恢复、连接池复用、健康合同与 DOCX 合同"),
        ],
        (5.0, 11.5),
    )
    document.add_paragraph("若当前哈希不等于已知旧基线或上述现行哈希，受控部署器应拒绝覆盖。先对差异做只读审计并生成新的、带明确基线的部署版本；不得绕过哈希保护。")

    add_heading(document, "9. 已知边界与未改动项")
    add_bullets(
        document,
        [
            "本次没有把 OLLAMA_MAX_LOADED_MODELS 从 1 改回 2，也没有改变只允许批准 27.8B 驻留的政策。",
            "2026-08-05 已把 BF_QA_KNOWLEDGE_SEARCH_MODE 显式改为 keyword；没有迁移或加载 embedding。以后如需 vector/hybrid，必须先使用独立 embedding 运行时。",
            "本次没有把 8768 当作问答健康证明；它只负责实时炉况数据。",
            "守卫修复没有停止 8093 业务服务；keyword 是启动期环境变量，后续受控部署只停止/启动 8093，未重启 8768、8094、8770、11434 或数据库。",
            "如果将来证据指向数据库、RAG、MCP、模型或前端消费，不得机械套用守卫修复；按第 5、7 节重新分类。",
        ],
    )

    add_heading(document, "10. 项目内权威入口")
    add_table(
        document,
        ("内容", "项目相对路径"),
        [
            ("长期代理操作规程", "AGENTS.md#8093-智能助手无回复固定检查流程2026-08-04"),
            ("一键工具专项说明", "docs/8093智能助手自动诊断修复验收工具_20260805.md"),
            ("一键工具 CLI", "docs/cli_usage.md#8093-智能助手自动诊断修复验收"),
            ("自动值守索引", "docs/自动值守程序配置索引.yaml"),
            ("守卫运维记录", "docs/22012_8093_v4_guard_ops.md"),
            ("自动化追踪", "docs/automation_traceability.md#ops-8093-assistant-health-guard-fix-20260804"),
            ("keyword 变更追踪", "docs/automation_traceability.md#ops-8093-knowledge-keyword-mode-20260805"),
            ("错误追踪", "docs/error_traceability.md#err-8093-assistant-health-restart-loop-20260804"),
            ("本轮部署碰撞错误", "docs/error_traceability.md#err-8093-assistant-deployment-collision-20260806"),
            ("Prompt/RAG 与 KV 前缀", "docs/8093_8094_Prompt固定KV前缀与知识库回答链路_20260805.md"),
            ("配置参考", "docs/config_reference.md#80938094-智能助手知识检索模式"),
            ("通用排障", "docs/troubleshooting.md#8093-智能助手无回复且健康状态稍后恢复正常"),
            ("测试说明", "docs/test_reference.md#test-8093-assistant-health-guard-recovery-20260804"),
            ("共享守卫实现", "tools/check_managed_nssm_service_health.ps1"),
            ("受控部署器", "tools/remote_guarded_deploy_8093_health_guard.ps1"),
            ("keyword 受控部署器", "tools/remote_guarded_deploy_8093_keyword_knowledge_mode.ps1"),
            ("一键编排器", "tools/assistant_8093_auto_recovery.py"),
            ("远端合并诊断", "tools/remote_8093_assistant_diagnose.ps1"),
            ("服务受控恢复", "tools/remote_guarded_recover_8093_service.ps1"),
            ("页面无停机热部署", "tools/remote_hot_deploy_8093_fetch_resilience.ps1"),
            ("一键工具测试", "tests/test_8093_assistant_auto_recovery.py"),
            ("单次 SSE 验收", "tools/verify_8093_assistant_sse_once.py"),
            ("本手册生成器", "tools/generate_8093_assistant_repair_docx.py"),
        ],
        (5.5, 11.0),
    )

    document.add_paragraph()
    closing = document.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("核心原则：先用日志证明故障链，再做最小受控修复，最后只用一次真实 SSE 验收。")
    set_run_font(run, size=11, bold=True)
    run.font.color.rgb = RGBColor(24, 83, 91)
    return document


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate the 8093 assistant repair DOCX handbook.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="DOCX output path")
    args = parser.parse_args()

    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.core_properties.title = TITLE
    document.core_properties.subject = f"{REQUIREMENT_ID}; {SECONDARY_REQUIREMENT_ID}; {TERTIARY_REQUIREMENT_ID}; {INCIDENT_REQUIREMENT_ID}"
    document.core_properties.author = "冀南钢铁项目组"
    document.core_properties.keywords = "8093, 智能助手, SSE, 健康守卫, keyword, 知识库, 一键诊断, 自动恢复, Failed to fetch, 部署互斥"
    document.core_properties.comments = "由 tools/generate_8093_assistant_repair_docx.py 生成"
    document.save(output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
