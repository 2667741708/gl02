# -*- coding: utf-8 -*-
"""Directly query IMES Vastbase for 3 months of data and export to Excel."""
from __future__ import annotations

import json
import os
import sys
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

EXCEL_AVAILABLE = False
try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    pass

try:
    import psycopg
except ImportError:
    import psycopg2 as psycopg


OUTPUT_DIR = Path(os.environ.get("IMES_EXPORT_DIR", r"F:\bf_export_packages\imes_vastbase_export"))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DB_USER = os.environ.get("IMES_DB_USER")
DB_PASSWORD = os.environ.get("IMES_DB_PASSWORD")

if not DB_USER or not DB_PASSWORD:
    print("ERROR: IMES_DB_USER/IMES_DB_PASSWORD not set")
    sys.exit(1)

VASTBASE_HOST = os.environ.get("IMES_DB_HOST", "10.10.181.195")
VASTBASE_PORT = int(os.environ.get("IMES_DB_PORT", "5432"))
VASTBASE_DB = os.environ.get("IMES_DB_NAME", "vastbase")

CONN_INFO = {
    "host": VASTBASE_HOST,
    "port": VASTBASE_PORT,
    "dbname": VASTBASE_DB,
    "user": DB_USER,
    "password": DB_PASSWORD,
    "connect_timeout": 15,
    "options": "-c statement_timeout=300000",
}

TODAY = date.today().strftime("%Y%m%d")
START_DATE = os.environ.get("IMES_START_DATE", "2026-02-10")
END_DATE = os.environ.get("IMES_END_DATE", "2026-05-10")


def write_excel(rows, columns, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columns)
    for row in rows:
        ws.append([str(v) if v is not None else "" for v in row])
    for col_idx, _ in enumerate(columns, 1):
        letter = get_column_letter(col_idx)
        if len(columns) > 30:
            ws.column_dimensions[letter].width = 14
        else:
            ws.column_dimensions[letter].width = 22
    if len(rows) > 0:
        ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
    return len(rows)


def write_csv(rows, columns, filepath):
    import csv
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        for row in rows:
            writer.writerow([str(v) if v is not None else "" for v in row])
    return len(rows)


def fetch_all(cur, sql, params=None):
    cur.execute(sql, params)
    cols = [desc[0] for desc in cur.description]
    rows = cur.fetchall()
    return cols, rows


def export_result(columns, rows, filename):
    if EXCEL_AVAILABLE:
        fp = OUTPUT_DIR / f"{filename}.xlsx"
        n = write_excel(rows, columns, fp)
    else:
        fp = OUTPUT_DIR / f"{filename}.csv"
        n = write_csv(rows, columns, fp)
    print(f"  -> {fp.name}: {n} rows, {len(columns)} cols")
    return n


def main():
    print(f"Excel available: {EXCEL_AVAILABLE}")
    print(f"Vastbase: {VASTBASE_HOST}:{VASTBASE_PORT}/{VASTBASE_DB}")
    print(f"Date range: {START_DATE} ~ {END_DATE}")
    print(f"Output dir: {OUTPUT_DIR}")
    print(f"Connecting...")

    conn = psycopg.connect(**CONN_INFO)
    cur = conn.cursor()
    total = 0

    # ============================================================
    # 1. 生产实绩 - t_ipes_out_put
    # ============================================================
    print("\n=== 1. 生产实绩 (t_ipes_out_put) ===")
    cols, rows = fetch_all(cur,
        """SELECT *
           FROM public.t_ipes_out_put
           WHERE workdate >= %s AND workdate < %s
             AND prodcentercode = '2D012'
           ORDER BY workdate, meltno""",
        (START_DATE, END_DATE + " 23:59:59")
    )
    total += export_result(cols, rows, "01_生产实绩_t_ipes_out_put")

    # ============================================================
    # 2. 炉次条件 - t_ipes_cond
    # ============================================================
    print("\n=== 2. 炉次条件 (t_ipes_cond) ===")
    cols, rows = fetch_all(cur,
        """SELECT *
           FROM public.t_ipes_cond
           WHERE workdate >= %s AND workdate < %s
             AND prodcentercode = '2D012'
           ORDER BY workdate, meltno""",
        (START_DATE, END_DATE + " 23:59:59")
    )
    total += export_result(cols, rows, "02_炉次条件_t_ipes_cond")

    # ============================================================
    # 3. 铁水元素 - Step 1: 获取批次号
    # ============================================================
    print("\n=== 3. 铁水元素含量 (t_qpes_inner_batch + inner_batch_insp_bb) ===")
    cols_batch, batch_rows = fetch_all(cur,
        """SELECT batchno, heatno, judgetime, judgeclass, judgeshift, inspphyclass, inspstaff
           FROM public.t_qpes_inner_batch
           WHERE judgetime >= %s AND judgetime < %s
             AND heatno LIKE '2#%'
           ORDER BY judgetime, batchno""",
        (START_DATE, END_DATE + " 23:59:59")
    )
    print(f"  Found {len(batch_rows)} batch numbers")

    # Step 2: 按批次号读取铁水元素
    element_rows = []
    for batchno, *_ in batch_rows:
        cur.execute(
            """SELECT batchno, judgetime, publishtime, prodcentercode, inspphyclass,
                      judgeclass, inspshift, inspstaff, takesampletime,
                      value_01 AS C, value_02 AS Si, value_03 AS Mn,
                      value_04 AS P, value_05 AS S, value_06 AS Ti,
                      value_08 AS Cr, value_09 AS Cu, value_10 AS Ni, value_07 AS V
               FROM public.inner_batch_insp_bb
               WHERE batchno = %s""",
            (batchno,)
        )
        for row in cur.fetchall():
            element_rows.append(row)

    elem_cols = ["batchno", "judgetime", "publishtime", "prodcentercode", "inspphyclass",
                 "judgeclass", "inspshift", "inspstaff", "takesampletime",
                 "C", "Si", "Mn", "P", "S", "Ti", "Cr", "Cu", "Ni", "V"]
    total += export_result(elem_cols, element_rows, "03_铁水元素含量")

    # ============================================================
    # 4. 炉渣检验 - slag_inspection
    # ============================================================
    print("\n=== 4. 炉渣检验结果 (slag_inspection) ===")
    cols, rows = fetch_all(cur,
        """SELECT meltno, sampleno, prodcentercode, workdate, createtime, publishtime,
                  value_01 AS TFe, value_02 AS FeO, value_03 AS CaO,
                  value_04 AS MgO, value_05 AS SiO2, value_06 AS Al2O3,
                  value_07 AS TiO2, value_08 AS R2, value_09 AS R3,
                  value_10 AS R4, value_11 AS MgO_Al2O3, value_12 AS SiO2_Al2O3
           FROM public.slag_inspection
           WHERE workdate >= %s AND workdate < %s
             AND prodcentercode = 'JL2'
           ORDER BY workdate, meltno, sampleno""",
        (START_DATE, END_DATE + " 23:59:59")
    )
    total += export_result(cols, rows, "04_炉渣检验结果_slag_inspection")

    # ============================================================
    # Summary
    # ============================================================
    summary_path = OUTPUT_DIR / f"vastbase_export_summary_{TODAY}.json"
    summary = {
        "exported_at": datetime.now().isoformat(),
        "format": "xlsx" if EXCEL_AVAILABLE else "csv",
        "source": f"{VASTBASE_HOST}:{VASTBASE_PORT}/{VASTBASE_DB}",
        "date_range": f"{START_DATE} ~ {END_DATE}",
        "total_rows_exported": total,
        "output_dir": str(OUTPUT_DIR),
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n=== DONE ===")
    print(f"Total: {total} rows exported")
    print(f"Summary: {summary_path}")

    cur.close()
    conn.close()


if __name__ == "__main__":
    main()
