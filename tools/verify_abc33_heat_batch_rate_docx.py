"""Verify the bold 30+30 minute burden-batch speed criterion in an ABC33 DOCX."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from tools.build_abc33_handbook64_revision import (
    HEAT_BATCH_HARD_DELTA_LARGE_PER_HOUR,
    HEAT_BATCH_HARD_DELTA_SMALL_PER_HOUR,
    HEAT_BATCH_NORMAL_DELTA_LARGE_PER_HOUR,
    HEAT_BATCH_NORMAL_DELTA_SMALL_PER_HOUR,
    HEAT_BATCH_REFERENCE_LARGE_PER_HOUR,
    HEAT_BATCH_RATE_CRITERION,
    HEAT_BATCH_RATE_RULE_IDS,
    HEAT_BATCH_RATE_WEIGHT,
    HEAT_BATCH_RATE_WEIGHTED_FORMULAS,
    HEAT_BATCH_TWO_HOUR_EXCESS_LARGE_PER_HOUR,
)


def verify_document(path: Path) -> dict[str, object]:
    """Return an audit payload or raise when the criterion is absent/not bold."""
    document = Document(path)
    paragraphs = list(document.paragraphs)
    table_paragraphs = [
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    all_paragraphs = paragraphs + table_paragraphs
    matches = [
        paragraph
        for paragraph in all_paragraphs
        if HEAT_BATCH_RATE_CRITERION in paragraph.text
    ]
    expected = len(HEAT_BATCH_RATE_RULE_IDS) * 2
    if len(matches) != expected:
        raise ValueError(f"expected {expected} criterion paragraphs, got {len(matches)}")
    if any(not paragraph.runs for paragraph in matches):
        raise ValueError("criterion paragraph has no text runs")
    if any(run.bold is not True for paragraph in matches for run in paragraph.runs):
        raise ValueError("criterion paragraph is not fully bold")
    expected_headings = ("A2 炉温管理稳定分", "B4 炉热预警", "B5 炉凉/大凉预警")
    texts = [paragraph.text.strip() for paragraph in paragraphs]
    for expected_heading in expected_headings:
        if expected_heading not in texts:
            raise ValueError(f"missing rule heading: {expected_heading}")
        heading_index = texts.index(expected_heading)
        if heading_index + 1 >= len(paragraphs):
            raise ValueError(f"rule heading has no first item: {expected_heading}")
        first_item = paragraphs[heading_index + 1]
        if HEAT_BATCH_RATE_CRITERION not in first_item.text:
            raise ValueError(f"burden speed is not first item: {expected_heading}")
        expected_label = "首要维护判据" if expected_heading.startswith("A2") else "首要风险判据"
        if expected_label not in first_item.text:
            raise ValueError(f"incorrect first-item label: {expected_heading}")
    formula_first: set[str] = set()
    for rule_id, weighted_formula in HEAT_BATCH_RATE_WEIGHTED_FORMULAS.items():
        located = False
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_paragraphs = list(cell.paragraphs)
                    for index, paragraph in enumerate(cell_paragraphs):
                        if weighted_formula not in paragraph.text:
                            continue
                        if index == 0 or HEAT_BATCH_RATE_CRITERION not in cell_paragraphs[index - 1].text:
                            raise ValueError(f"burden speed is not first formula item: {rule_id}")
                        expected_label = "料速维护项" if rule_id == "A2" else "料速风险项"
                        if expected_label not in cell_paragraphs[index - 1].text:
                            raise ValueError(f"incorrect formula-item label: {rule_id}")
                        if f"权重{HEAT_BATCH_RATE_WEIGHT}" not in cell_paragraphs[index - 1].text:
                            raise ValueError(f"incorrect formula-item weight: {rule_id}")
                        located = True
                        formula_first.add(rule_id)
        if not located:
            raise ValueError(f"missing formula anchor: {rule_id}")
    return {
        "ok": True,
        "path": str(path),
        "criterion_paragraph_count": len(matches),
        "fully_bold": True,
        "sorted_first_in_rules": sorted(HEAT_BATCH_RATE_RULE_IDS),
        "sorted_first_in_formula_cells": sorted(formula_first),
        "formula_weight": HEAT_BATCH_RATE_WEIGHT,
        "rebalanced_existing_weight_total": 100 - HEAT_BATCH_RATE_WEIGHT,
        "reference_rate_large_batches_per_hour": HEAT_BATCH_REFERENCE_LARGE_PER_HOUR,
        "normal_rate_delta_per_hour": {
            "large_batches_per_hour": HEAT_BATCH_NORMAL_DELTA_LARGE_PER_HOUR,
            "small_batches_per_hour": HEAT_BATCH_NORMAL_DELTA_SMALL_PER_HOUR,
        },
        "hard_rate_delta_per_hour": {
            "large_batches_per_hour": HEAT_BATCH_HARD_DELTA_LARGE_PER_HOUR,
            "small_batches_per_hour": HEAT_BATCH_HARD_DELTA_SMALL_PER_HOUR,
        },
        "two_hour_excess_vs_mes_24h_large_batches_per_hour": HEAT_BATCH_TWO_HOUR_EXCESS_LARGE_PER_HOUR,
        "example_small_batch_range_second_30min": [6.5, 7.5],
        "rule_ids": sorted(HEAT_BATCH_RATE_RULE_IDS),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    print(json.dumps(verify_document(args.docx), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
