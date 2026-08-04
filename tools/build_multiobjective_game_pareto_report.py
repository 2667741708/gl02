"""基于 Design Report 母版生成高炉工艺大模型智能决策系统技术报告。

Requirement: DOC-MOG-PARETO-BF-20260727
"""
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "高炉工艺大模型智能决策系统技术报告_V3.0_20260727.docx"
TEMPLATE = (
    Path.home()
    / ".codex"
    / "plugins"
    / "cache"
    / "openai-curated-remote"
    / "openai-templates"
    / "0.1.0"
    / "skills"
    / "artifact-template-design-report"
    / "assets"
    / "reference.docx"
)
ASSET_DIR = ROOT / "docs" / "assets" / "bf_decision_report_v2_4"
STEEL_LOGO = ROOT / "高炉前端数据" / "logo" / "冀南钢铁集团logo.png"
YSU_LOGO = ROOT / "高炉前端数据" / "logo" / "燕山大学logo.png"
SYSTEM_SCREENSHOTS = [
    (ASSET_DIR / "01_system_overview.png", "图1  高炉工艺大模型智能决策系统总览界面"),
    (ASSET_DIR / "02_condition_analysis.png", "图2  炉况研判与关键变量趋势界面"),
    (ASSET_DIR / "03_parameter_decision.png", "图3  参数优化与决策支持界面"),
    (ASSET_DIR / "04_trend_analysis.png", "图4  关键工艺参数趋势分析界面"),
    (ASSET_DIR / "05_knowledge_assistant.png", "图5  智能问答与知识助手界面"),
]
FINETUNE_FLOW = ASSET_DIR / "07_efficient_finetune_flow.png"
SYSTEM_FLOW = ASSET_DIR / "06_system_decision_flow.png"
DARK, ACCENT, PALE, AMBER, RED = "243231", "2F756E", "F4F8F7", "D8B45D", "A8483D"


def force_run_font(run, latin_name, east_asia_name):
    """锁定 Word 的四类字体属性，阻止主题字体回退为 MS Gothic。"""
    run.font.name = latin_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{attr}"), latin_name)
    r_fonts.set(qn("w:eastAsia"), east_asia_name)
    r_fonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def force_style_font(style, latin_name, east_asia_name, size=None, bold=None, color=None):
    style.font.name = latin_name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{attr}"), latin_name)
    r_fonts.set(qn("w:eastAsia"), east_asia_name)
    r_fonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def clear_template_body(doc):
    """保留母版样式、主题和页面设置，移除母版示例内容。"""
    body_element = doc._element.body
    section_properties = body_element.sectPr
    for child in list(body_element):
        if child is not section_properties:
            body_element.remove(child)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = pr.find(qn("w:shd")) or OxmlElement("w:shd")
    if shd.getparent() is None:
        pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, value, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value))
    r.bold = bold
    force_run_font(r, "SimSun", "宋体")
    r.font.size = Pt(12)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    tbl_pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "B9C3C1")
        borders.append(border)
    tbl_pr.append(borders)
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, "FFFFFF")
        shade(t.rows[0].cells[i], DARK)
        if widths:
            t.rows[0].cells[i].width = Cm(widths[i])
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if widths:
                cells[i].width = Cm(widths[i])
        if n % 2:
            for c in cells:
                shade(c, PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.add_run(text)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(3)
        marker = p.add_run("• ")
        force_run_font(marker, "SimSun", "宋体")
        p.add_run(item)


def formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(7)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF3F2")
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(12)
    force_run_font(r, "SimSun", "宋体")


def add_figure(doc, image_path, caption, width_cm=14.0, page_break_before=False):
    # 图文按自然流式排版，避免每张图强制分页造成大面积留白。
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(6)
    run = cp.add_run(caption)
    force_run_font(run, "SimSun", "宋体")
    run.font.size = Pt(12)


def configure(doc):
    s = doc.sections[0]
    s.top_margin = s.left_margin = s.right_margin = Cm(1.8)
    s.bottom_margin = Cm(1.7)
    normal = doc.styles["normal"]
    force_style_font(normal, "SimSun", "宋体", size=12)
    for style_name in ("Body Text", "List Bullet", "List Number"):
        if style_name in doc.styles:
            force_style_font(doc.styles[style_name], "SimSun", "宋体", size=12)
    for name, size, color in (
        ("Title", 24, DARK), ("Heading 1", 17, DARK),
        ("Heading 2", 14, DARK), ("Heading 3", 12, DARK),
    ):
        force_style_font(doc.styles[name], "SimHei", "黑体", size=size, bold=True, color=color)


def cover(doc):
    logos = doc.add_paragraph()
    logos.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logos.paragraph_format.space_after = Pt(88)
    logos.add_run().add_picture(str(STEEL_LOGO), width=Cm(4.2))
    logos.add_run("  ")
    logos.add_run().add_picture(str(YSU_LOGO), width=Cm(3.6))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("高炉工艺大模型智能决策系统技术报告")
    r.bold, r.font.size = True, Pt(26)
    r.font.color.rgb = RGBColor.from_string(DARK)
    force_run_font(r, "SimHei", "黑体")
    doc.add_section(WD_SECTION.NEW_PAGE)


def enforce_document_fonts(doc):
    """将正文、表格、图注、公式和标题逐个 run 固化为指定中文字体。"""
    heading_style_ids = {
        doc.styles[name].style_id
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3")
        if name in doc.styles
    }
    roots = [doc._element.body]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])
    for root in roots:
        for paragraph in root.iter(qn("w:p")):
            style_node = paragraph.find("./w:pPr/w:pStyle", paragraph.nsmap)
            style_id = style_node.get(qn("w:val")) if style_node is not None else ""
            is_heading = style_id in heading_style_ids
            latin, east_asia = ("SimHei", "黑体") if is_heading else ("SimSun", "宋体")
            for run_element in paragraph.iter(qn("w:r")):
                r_pr = run_element.get_or_add_rPr()
                r_fonts = r_pr.get_or_add_rFonts()
                for attr in ("ascii", "hAnsi", "cs"):
                    r_fonts.set(qn(f"w:{attr}"), latin)
                r_fonts.set(qn("w:eastAsia"), east_asia)
                r_fonts.set(qn("w:hint"), "eastAsia")
                lang = r_pr.find(qn("w:lang"))
                if lang is None:
                    lang = OxmlElement("w:lang")
                    r_pr.append(lang)
                lang.set(qn("w:eastAsia"), "zh-CN")


def remove_unused_template_media(doc):
    """移除母版示例页留下的未引用图片关系。"""
    document_xml = doc._element.xml
    for relation_id, relation in list(doc.part.rels.items()):
        if relation.reltype.endswith("/image") and relation_id not in document_xml:
            doc.part.drop_rel(relation_id)


def build():
    doc = Document(TEMPLATE)
    clear_template_body(doc)
    configure(doc)
    cover(doc)

    heading(doc, "摘要")
    body(doc, "本报告面向高炉工艺大模型智能决策与关键操作参数多目标优化场景，依托多源工艺数据、历史运行信息、质量反馈、趋势分析能力和参数优化工作台，构建融合真实运行高炉数据、专家知识库、MCP工具服务、多智能体任务编排、多目标博弈与帕累托优化的智能辅助决策体系。项目采用1.5TB高炉领域训练数据对27B参数规模基础模型进行微调，形成面向高炉炼铁场景的“炽穹·高炉炼铁大模型”；系统的炉况研判与操作建议由该领域大模型结合具体高炉专家知识库和当前生产上下文综合形成，并在炉况顺行、铁水质量、燃料消耗和炉体寿命等目标相互制约时，通过多目标博弈协调与帕累托解集展示，为氧煤比、风温、鼓风湿度和装料制度等参数提供可比较、可回退、可审计的候选方案。系统采用离线验证、影子运行、人工确认和安全门控的应用路径，为现场人员提供专业、可解释的辅助决策支持，生产操作由授权人员依照现场制度确认执行。")

    heading(doc, "1. 项目背景与建设目标")
    heading(doc, "1.1 业务问题", 2)
    body(doc, "高炉操作具有强耦合、长时滞、非线性和多约束特点。同一调整可能改善某一目标、恶化另一目标，并在不同时间尺度产生不同影响。例如，提高热输入可能有利于短时热状态稳定，却可能增加燃料代价或热负荷；强化送风供氧可能改善冶炼强度，却可能改变压差透气性与炉体负荷。因此，寻找唯一最优值并不符合真实生产决策逻辑。")
    heading(doc, "1.2 建设目标", 2)
    bullets(doc, [
        "在硬安全约束内生成多组非支配候选方案，并同步呈现各目标的得失关系与选择依据。",
        "把炉况顺行、铁水质量、燃料消耗、炉体寿命分别建模为独立利益主体，通过博弈协商形成均衡候选。",
        "输出帕累托前沿、基准方案、均衡方案、保守方案及各目标增益/代价，支持高炉长按班次目标权衡。",
        "与现有实时数据、历史运行信息、趋势预测、炉次质量数据和参数优化页面形成可替换、可审计的接口。",
        "全过程采用辅助决策与授权确认机制，生产操作由授权人员按既有生产制度执行。",
    ])

    heading(doc, "2. 当前项目基础与系统适配")
    add_table(doc, ["现有能力", "当前基础", "对本方案的作用", "应用口径"], [
        ["实时状态", "多源工艺变量的历史与实时数据", "形成优化状态向量、质量标记和趋势", "呈现任务相关的业务状态与解释"],
        ["历史运行参照", "按项目治理口径形成的历史运行信息", "用于状态比较和动态可行域参考", "采用经治理的运行参照信息"],
        ["炉次质量", "正式炉次与铁水/炉渣化验具备受控对齐路径", "构造质量目标、延迟标签和回放评价", "保留结果时间与取样时间差异"],
        ["趋势预测", "多变量趋势分析与不确定性评估能力", "评估候选未来风险与收益", "呈现预测结论、可信区间与适用范围"],
        ["智能研判", "经过训练的27B领域大模型结合具体高炉专家知识库与生产上下文，综合形成炉况研判和操作建议", "作为多目标决策的状态依据、知识支撑和展示入口", "通过隔离适配层交换任务所需的业务结果"],
        ["安全审计", "已有数据状态、错误状态和审计要求", "形成输入门禁、结果状态和确认链", "生产执行沿用授权确认机制"],
    ], [2.4, 4.1, 4.3, 5.4])

    heading(doc, "2.1 系统主要功能界面", 2)
    body(doc, "系统围绕总览、炉况研判、参数优化、趋势分析和智能问答形成一体化现场工作台。以下界面图用于展示系统当前主要业务入口、信息组织方式和人机协作形态。")
    for image_path, caption in SYSTEM_SCREENSHOTS:
        add_figure(doc, image_path, caption, page_break_before=True)

    heading(doc, "2.2 高效领域微调与通用能力保持", 2)
    body(doc, "面向1.5TB高炉运行数据、工艺文档、专家问答和复盘材料，系统采用参数高效微调路线增强27B基础模型的高炉领域理解与任务执行能力。该路线以LoRA和QLoRA为代表：训练时冻结基础模型主体参数，仅训练规模较小的低秩适配参数；在资源受限场景下，可使用QLoRA对冻结模型进行低比特表示，并把训练更新集中在适配参数上，从而显著降低显存和训练资源需求。LoRA通过冻结预训练模型并注入可训练低秩矩阵实现高效适配；QLoRA进一步在量化冻结模型上训练低秩适配器，相关研究表明其可在显著降低训练资源的同时保持接近全精度微调的任务表现。")
    body(doc, "为在增强高炉专业能力的同时保持模型原有通用智能水平，训练数据同步配置通用理解、逻辑推理、指令遵循、安全问答和工具使用样本，并采用分阶段训练、受控训练强度和领域—通用双轨评测。基础模型主体参数保持冻结，原有知识与推理能力由基础模型承载，新增领域能力主要由适配参数承载。每个候选版本同时参加高炉专业任务与通用能力回归评测，专业能力提升且通用能力达到发布标准的版本进入影子验证和受控发布。")
    body(doc, "训练质量通过高炉专业任务、通用理解与推理、指令遵循、安全问答和工具使用的双轨评测进行确认。上线后结合真实业务影子评测、专家复核、失败样本回收、漂移检查和版本化复盘持续优化。微调方法采用LoRA（Hu等，2021）与QLoRA（Dettmers等，2023）参数高效路线。")
    add_figure(doc, FINETUNE_FLOW, "图6  1.5TB高炉领域数据参数高效微调与能力保持流程", page_break_before=True)

    heading(doc, "3. 炉况信息进入工艺大模型的数据流程")
    heading(doc, "3.1 数据来源与授权接入", 2)
    body(doc, "系统面向真实运行高炉建立受控数据接入链路。生产过程数据、设备状态、实验室质量信息、操作事件、班次记录和业务文档通过既有数据服务或MCP服务器按权限提供。MCP服务器在系统中承担工具连接层作用：它把允许调用的数据查询、趋势分析、质量查询、报表检索和业务资料检索能力，以统一工具形式提供给工艺大模型，使模型能够在回答问题时按需获取最新且有时间标记的事实。")
    body(doc, "所有进入大模型的数据均附带必要的时间、来源和质量说明。系统围绕本轮任务组织相关信息；数据缺失、陈旧、来源状态和调用结果随上下文一并进入后续处理，使模型能够区分有效事实与待核信息。")
    heading(doc, "3.2 炉况上下文形成", 2)
    body(doc, "系统将当前真实运行数据、近期变化、设备可用状态、质量反馈、近期操作事件以及已有炉况判断组织为统一的炉况上下文。已有炉况判断作为经过业务处理的结论性信息进入上下文，以面向现场的状态名称、风险提示、时间范围和可解释依据参与任务。该上下文与用户当前问题、岗位权限和当前生产目标共同构成本轮大模型任务的事实基础。")
    heading(doc, "3.3 专家知识库检索", 2)
    body(doc, "专家知识库汇集高炉工艺知识、专家经验、操作规程、设备说明、异常处置经验、历史复盘材料和经审核的企业知识。系统根据用户问题、炉况上下文和任务目标检索最相关的知识片段，并保留资料来源和适用范围。知识库提供专业背景与经验依据，真实运行数据提供当前事实，两者在大模型中相互补充。")
    heading(doc, "3.4 大模型综合回答", 2)
    formula(doc, "用户问题 + 炉况上下文 + MCP工具结果 + 专家知识证据 + 权限与安全边界 → 27B工艺大模型 → 可解释回答")
    body(doc, "经过1.5TB高炉领域训练数据微调得到的27B“炽穹·高炉炼铁大模型”，对用户问题进行任务理解，并结合真实运行高炉数据、已有炉况判断、MCP工具返回结果和专家知识证据形成回答。回答包括当前状态说明、数据依据、可能影响、建议关注方向、后续观察重点和信息完整性提示。炉况研判和操作建议均由领域大模型结合当前数据与具体高炉专家知识综合形成，并由授权人员结合现场制度完成确认。")
    heading(doc, "3.5 回答输出与审计", 2)
    body(doc, "回答输出保留数据时间、知识来源、工具调用状态和本轮任务标识，便于现场人员核对事实并开展后续复盘。涉及生产操作的内容以辅助建议方式呈现，同时给出观察重点、风险提示和人工确认要求，生产执行沿用现场授权流程。")
    add_table(doc, ["数据流程阶段", "输入", "业务处理", "输出"], [
        ["任务理解", "用户问题、岗位和会话上下文", "识别问题对象、时间范围和所需能力", "受控任务说明"],
        ["事实获取", "真实运行数据和业务系统", "通过MCP工具按权限查询并标记时间与质量", "当前事实与工具状态"],
        ["炉况融合", "当前事实、近期变化和已有炉况判断", "形成面向本轮问答的炉况上下文", "统一炉况信息"],
        ["知识增强", "专家知识库和企业资料", "检索与当前任务相关的专业证据", "知识证据与适用范围"],
        ["综合生成", "问题、炉况、工具结果和知识证据", "由27B领域大模型综合理解与生成", "可解释回答和关注方向"],
        ["安全审查", "回答草稿、权限和风险信息", "检查事实完整性、适用边界和人工确认要求", "面向现场的最终回答"],
        ["审计承接", "本轮输入来源、调用状态和输出", "记录任务标识与后续待观察事项", "可追踪的会话与任务记录"],
    ], [2.6, 4.2, 6.1, 4.1])
    add_figure(doc, SYSTEM_FLOW, "图7  高炉工艺大模型智能决策系统总体流程", page_break_before=True)

    heading(doc, "4. 多智能体任务编排与协同决策")
    body(doc, "系统将工艺大模型、专家知识库、MCP工具能力和多目标优化组织为一个可协作、可追踪的智能决策体系，通过共享炉况上下文、任务分解、角色协作、冲突协调、任务编排、安全门控、反馈回写和班次承接，形成适应高炉连续生产特点的协同决策闭环。本报告聚焦业务级协同框架与工程流程，内部技术资产由项目治理体系统一管理。")
    heading(doc, "4.1 共享炉况上下文", 2)
    body(doc, "多个角色智能体围绕同一份共享炉况上下文开展工作，使数据查询、知识检索、目标分析、安全审查和回答生成基于一致的当前事实。共享内容包括任务相关的真实运行信息、已有炉况判断、设备与数据状态、当前生产目标、近期事件、待观察事项和班次承接信息。")
    heading(doc, "4.2 任务分解与角色协同", 2)
    body(doc, "系统将复杂问题分解为可协同处理的子任务，并由不同角色智能体并行完成。数据工具角色负责通过MCP获取事实，知识角色负责检索专家知识，炉况理解角色负责整合已有判断和当前上下文，目标协调角色负责分析多目标影响，安全角色负责检查设备状态与风险边界，回答角色负责汇总形成面向现场的自然语言结果。")
    heading(doc, "4.3 冲突识别与协同编排", 2)
    body(doc, "当不同角色形成的目标、关注方向或候选方案存在不一致时，系统对参数影响、目标影响、时间先后、资源条件和安全要求进行统一比较，形成可解释的协调结果。任务编排明确先处理事项、后续观察事项、人工确认节点、备选路径和停止条件，使多个并行结论转化为顺序清晰的辅助决策流程。")
    heading(doc, "4.4 安全门控、反馈与班次承接", 2)
    body(doc, "系统根据风险等级、数据状态、设备可用性和人员权限确定展示方式、人工确认节点和风险提示等级。任务完成后的数据变化、人工确认结果、待观察事项和后续事项进入任务承接与班次交接，使上一轮分析能够在下一轮继续衔接，形成连续运行的闭环。")
    formula(doc, "共享炉况上下文 → 任务分解 → 多角色并行协作 → 冲突协调 → 任务编排 → 安全门控 → 反馈与班次承接")
    add_table(doc, ["协同角色", "主要职责", "主要输入", "主要输出"], [
        ["任务编排角色", "理解问题并组织任务依赖与执行顺序", "用户问题、生产目标和炉况上下文", "任务图与协作计划"],
        ["数据工具角色", "调用MCP工具获取真实运行事实", "任务所需对象、时间范围和权限", "带来源与时间的数据结果"],
        ["知识增强角色", "从专家知识库获取专业证据", "问题语义和炉况上下文", "相关知识片段与适用说明"],
        ["炉况理解角色", "结合已有炉况判断和当前事实形成统一理解", "真实数据、已有判断和近期事件", "业务级炉况说明"],
        ["目标协调角色", "比较顺行、质量、能耗和寿命等目标影响", "候选方案与多目标评价", "目标权衡和代表方案"],
        ["安全审查角色", "核对数据状态、设备条件和人工确认要求", "候选回答、风险和权限", "通过、降级或停止说明"],
        ["回答生成角色", "整合事实、知识、炉况和协调结果", "各角色输出及证据", "可解释的最终回答"],
    ], [3.0, 5.2, 5.2, 4.0])

    heading(doc, "5. 多目标优化问题定义")
    heading(doc, "5.1 决策变量", 2)
    body(doc, "将待优化操作参数记为决策向量 x。首期纳入具备可靠测量、明确权限和可审计操作记录的变量；氧煤比由相关供氧与喷煤量按统一时间口径派生。装料制度采用经工艺许可的离散方案编号或受控参数模板。")
    formula(doc, "x = [氧煤比，风温，鼓风湿度，装料制度，……]ᵀ")
    add_table(doc, ["变量组", "表达", "寻优方式", "主要约束"], [
        ["氧煤比", "连续变量或派生变量", "连续寻优", "上下界、变化率、供氧/喷煤能力、数据新鲜度"],
        ["风温", "连续变量", "连续寻优", "设备能力、班次许可、变化率与热状态门禁"],
        ["鼓风湿度", "连续变量", "连续寻优", "设备可用域、环境状态和变化率"],
        ["装料制度", "受控离散模板", "离散/混合整数寻优", "只允许已批准方案；切换频率与过渡期约束"],
    ], [2.5, 4.0, 3.2, 6.5])
    heading(doc, "5.2 目标向量", 2)
    formula(doc, "min F(x | sₜ) = [f顺行(x), f质量(x), f燃料(x), f寿命(x)]")
    add_table(doc, ["目标", "工程含义", "可观测代理（示意）", "方向"], [
        ["炉况顺行", "降低未来窗口异常、波动和透气性恶化风险", "趋势稳定性、风险概率、状态偏离和不确定性", "越低越优"],
        ["铁水质量", "提高质量达标概率并降低波动", "质量预测偏差、达标概率、置信区间", "偏差越低越优"],
        ["燃料消耗", "降低单位产出燃料代价", "燃料比或经计量核验的代理指标", "越低越优"],
        ["炉体寿命", "降低长期热负荷、热冲击和不均匀负荷", "炉体温度分布、变化速率、分区负荷与损伤代理", "风险越低越优"],
    ], [2.4, 5.1, 6.1, 2.6])
    body(doc, "代理指标经过数据治理、专家审查和回测确认后纳入生产验收体系；目标表达用于构建系统设计与验证口径。")
    heading(doc, "5.3 约束体系", 2)
    formula(doc, "x ∈ Ω(sₜ) = Ω设备 ∩ Ω工艺 ∩ Ω安全 ∩ Ω速率 ∩ Ω数据 ∩ Ω权限")
    bullets(doc, [
        "硬约束：设备上下限、互锁边界、许可区间、变化速率、最小保持时间、禁止组合、数据质量和角色权限。",
        "软约束：偏离历史稳定区间惩罚、频繁调整惩罚、方案复杂度、操作成本和预测不确定性。",
        "动态约束：随当前状态、班次目标、设备可用性和预测风险实时收缩可行域。",
        "安全门禁：输入陈旧、关键点缺失、预测异常、候选超界或不确定性过高时，只返回保持/人工复核。",
    ])

    heading(doc, "6. 多目标博弈机制")
    heading(doc, "6.1 参与方与策略", 2)
    body(doc, "把四类目标抽象为四个虚拟参与方。各参与方围绕同一候选 x 给出本目标效用、底线和让步幅度，由协调器在安全可行域内组织协商，保持各目标之间的均衡关系。")
    add_table(doc, ["参与方", "核心诉求", "底线", "可协商内容"], [
        ["顺行方", "未来窗口稳定、风险可控", "安全门禁通过", "在风险余量内让渡部分短期收益"],
        ["质量方", "达标概率与波动受控", "质量风险处于许可区间", "置信度允许时接受小幅偏差"],
        ["燃料方", "单位产出燃料代价下降", "以顺行状态为前提", "接受短期成本换取长期稳定"],
        ["寿命方", "热负荷和累计损伤受控", "关键区域负荷不越界", "低损伤区间内接受局部波动"],
    ], [2.3, 4.3, 4.5, 5.6])
    heading(doc, "6.2 协商效用与公平性", 2)
    body(doc, "对每一目标构造相对当前参照状态的改善量，并以底线效用表示最低可接受收益。可采用纳什议价思想选取兼顾整体收益与公平性的均衡点；偏好参数仅代表经治理后的班次目标。")
    formula(doc, "xᴺ = arg maxₓ∈Ω ∏ᵢ [uᵢ(x) − dᵢ]ᵝⁱ，且 uᵢ(x) > dᵢ")
    body(doc, "各目标先做稳健归一化；偏好采用版本化管理与双人复核，并同步展示原始目标值。除纳什议价外，保留最小最大遗憾、膝点和词典序保守选择作为对照，形成多视角协调结果。")
    heading(doc, "6.3 均衡判定", 2)
    bullets(doc, [
        "可行均衡：满足全部硬约束和数据门禁。",
        "个体理性：任一目标相对基准不低于其底线。",
        "帕累托有效：不存在另一可行方案同时改善至少一个目标且不恶化其他目标。",
        "稳定可实施：变化幅度、保持时间和操作复杂度可接受。",
        "可解释：说明每个目标的改善、代价、不确定性及约束。",
    ])

    heading(doc, "7. 帕累托最优解集")
    heading(doc, "7.1 非支配关系", 2)
    formula(doc, "xᵃ ≺ xᵇ ⇔ ∀i: fᵢ(xᵃ) ≤ fᵢ(xᵇ)，且 ∃j: fⱼ(xᵃ) < fⱼ(xᵇ)")
    body(doc, "若候选 a 在所有目标上不差于候选 b，且至少一个目标更优，则 a 支配 b。所有不被其他可行候选支配的方案构成帕累托最优解集。解集把目标冲突公开化：现场能够看到降低燃料消耗对应的顺行风险变化，也能看到更保守的寿命方案牺牲了多少短期收益。")
    heading(doc, "7.2 解集生成流程", 2)
    bullets(doc, [
        "围绕当前设定、历史稳定操作和批准的离散模板生成初始候选。",
        "对超出硬约束的候选执行剔除或可行域投影，硬约束作为候选进入后续评估的先决条件。",
        "评估四类目标的未来窗口指标及不确定性。",
        "进行非支配排序并保持目标空间多样性。",
        "在第一前沿上计算议价均衡、最小遗憾和保守膝点。",
        "对数据扰动、预测误差和参数变化做情景检验，过滤脆弱方案。",
    ])
    heading(doc, "7.3 代表方案", 2)
    add_table(doc, ["标签", "选择逻辑", "场景", "必须展示"], [
        ["均衡方案", "议价收益最大且满足底线", "常规班次", "目标得失、约束余量、不确定性"],
        ["顺行优先", "前沿中最小化顺行风险", "波动或恢复期", "对燃料与质量的代价"],
        ["节能优先", "各底线内降低燃料代价", "稳定窗口节能", "风险余量和回退条件"],
        ["寿命优先", "降低热负荷与损伤代理", "设备关注期", "短期成本与响应时间"],
        ["最小调整", "距当前设定最近的非支配方案", "信息不足或调整频繁", "调整量、保持期与观察项"],
    ], [2.5, 5.2, 4.2, 4.8])

    heading(doc, "8. 算法流程与系统架构")
    formula(doc, "实时/炉次数据 → 状态构建 → 可行域门禁 → 候选生成 → 多目标评估 → 非支配排序 → 博弈协调 → 稳健复核 → 人工确认")
    add_table(doc, ["层级", "输入", "处理", "输出"], [
        ["数据层", "核心变量、炉体/静压力、炉次质量、设备状态", "对齐、质量标记、历史参照与窗口化", "状态快照"],
        ["预测层", "状态快照、候选操作", "趋势/质量/能耗/寿命代理预测", "目标值、区间和置信度"],
        ["约束层", "工艺许可、设备能力、权限与变化率", "动态可行域和安全门禁", "可行候选或拒绝原因"],
        ["优化层", "可行候选及四目标评价", "非支配排序、多样性、博弈协调", "帕累托集和代表方案"],
        ["决策层", "代表方案、得失和风险", "人工复核、选择、退回或保持", "辅助决策记录"],
        ["审计层", "输入、模型、约束和选择版本", "留痕与回放", "可复现证据包"],
    ], [2.0, 4.2, 5.4, 5.1])
    heading(doc, "8.1 与现有系统的接入建议", 2)
    body(doc, "系统通过独立隔离适配层衔接现有能力，保持分析、建议、优化和展示模块之间的职责清晰。适配层交换经过授权的任务相关业务信息，包括状态摘要、数据质量、趋势结果、质量反馈和许可域。")
    add_table(doc, ["接口方向", "任务相关信息", "治理方式"], [
        ["现有系统 → 优化层", "业务状态、数据质量、预测摘要、质量反馈和许可域", "采用业务语义与授权数据范围"],
        ["优化层 → 展示层", "计算状态、候选方案、代表方案、目标权衡、不确定性和约束说明", "采用可解释的方案级结果"],
        ["人工选择 → 审计", "方案标识、授权角色、选择原因和确认时间", "形成授权确认与审计记录"],
    ], [3.1, 7.5, 6.1])

    heading(doc, "9. 输出信息设计")
    add_table(doc, ["信息类别", "含义", "展示方式"], [
        ["计算状态", "本轮计算所处阶段或失败原因", "采用面向现场的业务状态名称"],
        ["比较基准", "候选方案的比较起点", "提供可追溯的业务标识"],
        ["非支配候选", "参数变化、四目标评价和置信区间", "提供经授权的业务语义"],
        ["代表方案", "均衡、顺行、节能、寿命和最小调整方案", "使用方案标识引用"],
        ["目标权衡", "相对基准的改善与代价", "保留业务单位与比较口径"],
        ["约束说明", "约束余量与阻断原因", "解释方案适用边界"],
        ["不确定性", "预测与数据不确定性", "展示区间、等级和主要来源"],
        ["审计索引", "版本化证据包引用", "支持授权回放与复盘"],
    ], [3.7, 7.2, 5.8])

    heading(doc, "10. 决策方案展示样式")
    body(doc, "系统以当前方案、均衡方案、顺行优先方案、节能优先方案和最小调整方案组织结果，使现场人员能够直接比较参数方向与四类目标的变化关系。")
    add_table(doc, ["方案", "氧煤比", "风温", "湿度", "装料制度", "顺行风险", "质量风险", "燃料代价", "寿命风险"], [
        ["当前基准", "0", "0", "0", "现行", "基准", "基准", "基准", "基准"],
        ["P-01 均衡", "小幅调整", "小幅调整", "保持", "模板 A", "改善", "改善", "改善", "基本持平"],
        ["P-02 顺行", "保持", "小幅调整", "小幅调整", "模板 A", "显著改善", "改善", "小幅增加", "改善"],
        ["P-03 节能", "小幅调整", "保持", "保持", "模板 B", "满足底线", "满足底线", "显著改善", "基本持平"],
        ["P-04 最小调整", "保持", "保持", "保持", "现行", "小幅改善", "基本持平", "小幅改善", "基本持平"],
    ], [2.5, 2.0, 1.8, 1.8, 2.2, 2.2, 2.2, 2.2, 2.2])
    body(doc, "前端同步展示参数变化、目标改善与代价、预测区间、约束余量、观察窗口、回退条件、数据时间、算法与约束版本以及人工选择入口。解集为空时，页面呈现硬约束冲突、数据状态或预测服务状态等原因。")

    heading(doc, "11. 算法实现策略")
    heading(doc, "11.1 两阶段求解", 2)
    body(doc, "第一阶段使用约束感知的多目标进化搜索或等价非支配搜索生成覆盖充分的帕累托候选；第二阶段在第一前沿上使用博弈议价、最小最大遗憾和膝点识别选择代表方案。两阶段结构把最优边界搜索与偏好选择分离，便于审计并保留各目标的原始业务含义。")
    heading(doc, "11.2 混合变量", 2)
    bullets(doc, [
        "连续参数采用有界实数编码，并显式施加变化率和保持期约束。",
        "装料制度等离散参数从批准模板集合中选择。",
        "派生变量采用统一时钟和单位，保障氧煤比等指标的时间一致性。",
        "目标评估前完成约束修复，候选通过可行性检查后进入后续计算。",
    ])
    heading(doc, "11.3 不确定性与稳健优化", 2)
    formula(doc, "稳健目标 = 期望损失 + λ × 不确定性惩罚 + ρ × 变化复杂度")
    body(doc, "每个目标同时返回点估计与区间。候选在扰动、缺点和预测误差情景下满足硬约束后进入代表方案。对超出历史支持域的候选增加外推惩罚并执行可行性复核，优化结果与数据覆盖证据共同构成评价依据。")

    heading(doc, "12. 数据治理与信息资产管理")
    add_table(doc, ["治理对象", "管理措施", "业务展示", "技术资产管理"], [
        ["分析服务", "适配层使用业务状态和许可域", "状态、风险等级、门禁结果", "技术过程纳入受控资产库"],
        ["建议服务", "通过业务接口形成协同支撑", "建议状态与许可动作域", "业务映射由授权模块维护"],
        ["模型服务", "服务封装、版本标识和质量评估", "模型类别、版本、区间和质量指标", "模型资产采用版本化管理"],
        ["访问凭据", "环境变量、受限配置和日志脱敏", "连接状态和数据新鲜度", "凭据由专用安全设施托管"],
        ["审计信息", "按角色提供任务相关字段", "审计标识、时间、版本和选择结果", "审计轨迹由授权角色管理"],
    ], [2.4, 5.2, 5.1, 5.1])
    bullets(doc, [
        "报告、接口和页面统一采用面向现场的业务语义。",
        "内部验证证据采用访问控制、脱敏导出和最短保留周期。",
        "展示材料使用合成或脱敏数据，真实生产信息纳入授权管理。",
        "对可能通过差分查询推断内部算法的请求限流并审计。",
    ])

    heading(doc, "13. 验证与验收方案")
    add_table(doc, ["阶段", "核心验证", "通过条件"], [
        ["数据合同", "字段、单位、时间对齐、质量标记、缺失/异常状态", "契约通过；无未来泄漏；错误可解释"],
        ["离线回放", "按时间切分回放稳定/波动/恢复窗口", "硬约束违规 0；非支配性和复现性通过"],
        ["多目标质量", "超体积、间距、多样性、覆盖率、遗憾值", "相对参照方案与对照算法达到冻结门槛"],
        ["稳健性", "扰动、预测区间、缺点、超时、分布漂移", "高风险场景进入降级或人工复核流程"],
        ["影子运行", "实时辅助结果与人工决策并行比较", "稳定产出、审计完整、确认链清晰"],
        ["人机交互", "解集、权衡、全文、加载/空/错/重试、权限", "核心按钮可用；全文可恢复查看；矩阵通过"],
        ["上线门禁", "权限、安全、回滚、版本冻结、变更审批", "书面批准后才进入生产辅助决策"],
    ], [2.6, 8.0, 6.5])
    heading(doc, "13.1 量化指标", 2)
    bullets(doc, [
        "可行率：输出候选满足全部硬约束的比例必须为 100%。",
        "非支配正确率：报告为帕累托解的候选在同批可行候选中保持非支配性。",
        "稳定性：相近输入下代表方案保持连续性，并记录切换率与调整幅度。",
        "回放收益：分别报告四目标原始量纲变化与综合比较结果。",
        "风险校准：预测概率与实际频率做分箱校准和漂移监测。",
        "人工可接受性：记录接受、修改、退回及原因，并与算法质量指标联合分析。",
    ])

    heading(doc, "14. 实施路线")
    add_table(doc, ["阶段", "主要工作", "交付物", "生产边界"], [
        ["P0 口径冻结", "变量、单位、目标代理、硬约束来源和角色", "数据字典、约束登记、保密清单", "文档核查"],
        ["P1 离线原型", "历史回放、候选、非支配排序和博弈选择", "离线程序、回放报告、前沿可视化", "离线验证环境"],
        ["P2 影子运行", "接入实时业务合同，输出辅助候选与审计", "独立服务、业务 API、审计表", "并行验证环境"],
        ["P3 工作台接入", "展示解集、权衡和人工选择", "真实路由、状态、权限、浏览器报告", "选择仅留痕"],
        ["P4 受控试验", "经审批的小范围人工执行与复盘", "试验方案、停试条件、复盘", "现场制度优先"],
        ["P5 持续治理", "漂移、回测、版本、权限和指标复核", "月报、回滚包、变更记录", "再批准后扩大"],
    ], [2.3, 6.5, 5.3, 3.0])
    heading(doc, "14.1 首期范围", 2)
    body(doc, "首期以氧煤比、风温、鼓风湿度和已批准装料制度模板为候选变量，以四类目标的可用代理为评价对象，采用历史回放和影子运行验证。燃料、寿命或质量目标依据数据完备度分别纳入生产评价或研究评价，并通过统一验收口径持续完善。")

    heading(doc, "15. 风险与应对")
    add_table(doc, ["风险", "表现", "应对"], [
        ["相关性冒充因果", "历史优组合被误当成可执行因果动作", "因果审查、受控试验、外推惩罚、专家门禁"],
        ["目标代理偏差", "代理改善但真实目标未改善", "代理—真实指标双轨评估与重标定"],
        ["标签时延/错位", "化验时间与操作窗口不一致", "保留多类时间，严格防未来泄漏"],
        ["分布漂移", "原料、设备、季节或制度变化", "漂移监测、收缩可行域、停用和重训"],
        ["解集过多", "现场难以选择", "展示 3—5 个代表方案，可查看完整前沿"],
        ["偏好固化", "长期不变形成隐性单目标", "偏好版本、期限、审批和敏感性分析"],
        ["算法泄密", "通过接口或报告逆推内部实现", "最小字段、访问控制、审计和限流"],
        ["自动化越权", "候选与现场执行流程混淆", "辅助决策标识、授权确认和审计留痕"],
    ], [3.0, 6.2, 7.0])

    heading(doc, "16. 预期价值")
    bullets(doc, [
        "把综合最优转化为多目标冲突可见、代价可量化、方案可选择。",
        "通过多目标协调适应不同工况与班次目标，提升公平性和适应性。",
        "使候选具备方案级依据、约束余量、不确定性和回退条件。",
        "复用当前数据、预测、炉次和工作台基础，以独立适配层降低侵入和泄密风险。",
        "形成可回放、可比较、可审计的研究—影子—受控上线闭环。",
    ])

    heading(doc, "17. 结论")
    body(doc, "多目标博弈与帕累托优化适合高炉关键操作参数的多约束冲突场景。系统在安全可行域内形成一组互不支配候选，再通过公平协商、稳健检验和人工确认选择代表方案，使现场经验、工艺大模型和多目标决策形成协同关系。基于当前项目的数据、历史运行信息、炉次质量、趋势预测、业务接口与工作台能力，已具备开展离线原型和影子运行设计的工程基础。")
    body(doc, "下一步将完成目标代理和硬约束来源的专家确认，建设独立离线回放原型，并按照业务目标、可行域、帕累托权衡、验证结果和安全边界组织系统输出。")

    heading(doc, "附录 A：术语")
    add_table(doc, ["术语", "定义"], [
        ["多目标优化", "同时考虑两个及以上相互冲突目标的优化问题。"],
        ["非支配解", "不存在另一可行解在所有目标上不差且至少一个目标更优。"],
        ["帕累托最优解集", "所有非支配可行解的集合，描述可实现的权衡边界。"],
        ["帕累托前沿", "帕累托最优解在目标空间中的映射。"],
        ["多目标博弈", "把不同目标视为利益主体，通过效用和底线形成均衡选择。"],
        ["纳什议价解", "在参与方底线之上最大化联合议价收益的一类公平解。"],
        ["膝点", "继续改善一个目标会显著恶化其他目标的高性价比折中点。"],
        ["最小遗憾方案", "在多种偏好或情景下，使最坏相对损失尽量小的方案。"],
        ["影子运行", "与现行生产流程并行计算，用于验证稳定性与效果。"],
    ], [4.0, 12.2])
    heading(doc, "附录 B：项目能力说明")
    bullets(doc, [
        "项目以多源关键工艺变量支撑状态分析，变量资产由项目数据治理体系统一维护。",
        "正式页面采用面向现场的业务语义，技术资产由授权模块管理。",
        "历史运行信息、分钟数据、炉次质量、趋势预测和现有建议通过业务接口协同工作。",
        "当前生产页面和后端定位于诊断与辅助决策，生产执行遵循现场授权制度。",
    ])

    enforce_document_fonts(doc)
    remove_unused_template_media(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


def verify(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    text = "\n".join(parts)
    risky_patterns = {
        "specific_furnace_number": r"\d+\s*#\s*高炉",
        "exact_core_variable_count": r"\d+\s*个核心变量",
        "exact_reference_window": r"\d+\s*天(?:历史)?参照",
        "internal_snake_case_contract": r"\b[a-z]+(?:_[a-z]+){2,}\b",
        "diagnosis_threshold_phrase": r"(?:主|次)异常\s*\d+",
    }
    with ZipFile(path) as zf:
        entries = len(zf.namelist())
        xml_text = "\n".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in zf.namelist()
            if name.endswith((".xml", ".rels"))
        )
    scan_text = text + "\n" + xml_text
    risk_hits = {
        name: sorted(set(re.findall(pattern, scan_text, flags=re.IGNORECASE)))
        for name, pattern in risky_patterns.items()
    }
    risk_hits = {name: hits for name, hits in risk_hits.items() if hits}
    return {
        "path": str(path), "bytes": path.stat().st_size,
        "paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
        "zip_entries": entries, "has_pareto": "帕累托最优解集" in text,
        "has_game": "多目标博弈" in text,
        "has_public_scope_statement": "最小必要原则" in text,
        "scanned_tables_headers_footers_and_xml": True,
        "risk_hits": risk_hits,
    }


if __name__ == "__main__":
    print(verify(build()))
