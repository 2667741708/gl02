from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "docs" / "8093智能助手不可用原因与正式修复手册_20260804.docx"
BACKUP = TARGET.with_name(TARGET.stem + ".before_abc33_20260811.docx")
MARKER = "REQ-ABC33-EXPLANATION-CONTEXT-20260811"
SECURITY_MARKER = "REQ-ABC33-CONTEXTUAL-ASSISTANT-SECURITY-20260812"
PRODUCTION_MARKER = "REQ-ABC33-CONTEXTUAL-ASSISTANT-PRODUCTION-20260812"


def paragraph_text(document: Document) -> str:
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def main() -> None:
    if not TARGET.exists():
        raise FileNotFoundError(TARGET)
    document = Document(TARGET)
    existing_text = paragraph_text(document)
    if MARKER in existing_text and SECURITY_MARKER in existing_text and PRODUCTION_MARKER in existing_text:
        print(f"already_updated={TARGET}")
        return
    if not BACKUP.exists():
        shutil.copy2(TARGET, BACKUP)

    if MARKER not in existing_text:
        document.add_page_break()
        document.add_heading("ABC33 炉框上下文智能助手（2026-08-11）", level=1)
        document.add_paragraph(MARKER)
        document.add_paragraph(
            "本方案把33个A/B/C炉框接入既有智能助手会话体系。ABC33权威计算仍是唯一数值来源；"
            "模型只解释形成原因、关注点与既有处置顺序，不得改分、改权重、改安全门禁或自动下发控制。"
        )

        document.add_heading("固定链路", level=2)
        add_bullet(document, "浏览器只提交 rule_id、evaluation_id 和 conversation_id。")
        add_bullet(document, "服务端重读权威批次，生成 abc_rule_explanation_context.v1 与稳定 SHA-256。")
        add_bullet(document, "不可变上下文快照绑定会话和消息；新批次不会覆盖历史回答依据。")
        add_bullet(document, "确定性解释先展示；模型或知识库故障时公式项与五步处置仍可查看。")
        add_bullet(document, "同一上下文、Prompt版本和模型只允许一个首轮生成，其余请求等待完成缓存。")

        document.add_heading("Prompt 与知识库", level=2)
        document.add_paragraph(
            "公共Prompt前缀稳定排列为：身份→安全边界→回答规则→固定工艺规则；"
            "单炉框动态上下文和本次任务置于后部。支持前缀缓存的模型运行时可复用相同attention/KV计算，"
            "但这不是人为固定注意力分数，也不保证每次命中。8093知识检索继续固定为keyword词法检索。"
        )

        document.add_heading("固定排障顺序", level=2)
        document.add_paragraph(
            "规则批次 → 解释上下文 → 会话来源 → 上下文快照 → 分析缓存 → 模型状态 → SSE → 助手页面索引"
        )
        add_bullet(document, "解释接口403：先恢复operator/admin登录；409：核对批次是否包含该规则。")
        add_bullet(document, "弹窗有公式但无AI：检查分析缓存、单航班owner、Ollama状态和SSE。")
        add_bullet(document, "模型生成时不得长期持有PostgreSQL连接；SSE失败不得自动重发。")
        add_bullet(document, "代码回滚保留四张增量审计表，不执行破坏性DROP。")

    if SECURITY_MARKER not in existing_text:
        document.add_heading("上下文助手安全与并发加固（2026-08-12）", level=2)
        document.add_paragraph(SECURITY_MARKER)
        add_bullet(document, "QA会话按登录主体隔离；不能读取或向其他操作者的会话发送问题。")
        add_bullet(document, "浏览器QA写请求必须为application/json并精确同源，端口不同也拒绝；不得恢复通配CORS。")
        add_bullet(document, "上下文会话只接收来源、页面、规则、批次和复用策略五字段，其他字段直接拒绝。")
        add_bullet(document, "首轮分析使用PostgreSQL原子claim与短租约跨8093/8094去重，进程锁不是唯一保障。")
        add_bullet(document, "模型严格JSON必须通过公式项、数值、处置顺序和C类安全提示校验后才能写完成缓存。")
        add_bullet(document, "诊断页当前没有创建QA会话的生产入口，因此诊断来源筛选为空属于已知状态，不得伪造记录。")

    if PRODUCTION_MARKER not in existing_text:
        document.add_heading("8093 正式部署与验收（2026-08-12）", level=2)
        document.add_paragraph(PRODUCTION_MARKER)
        add_bullet(document, "Ollama format使用当前上下文派生的JSON Schema；首次校验失败只允许在同一SSE请求内做一次受控repair，repair仍必须通过原validator，失败不得降级。")
        add_bullet(document, "密封执行ID为abc33-prod-20260812-0156-866ec9d9；部署备份位于backups/abc33_contextual_assistant_8093_20260812_015656。")
        add_bullet(document, "8093 PID由9152切换为6060；8094、8768、8770、PostgreSQL和Ollama PID全部保持不变；守卫暂停后已在finally中恢复。")
        add_bullet(document, "真实验收只发送首次解释和一次人工追问，共2个SSE请求；两次均包含preparing、prepared、delta、final、done，问答期间8093 PID不变。")
        add_bullet(document, "首次解释和追问答案分别为849和501字，prepared上下文哈希与权威解释哈希一致。")
        add_bullet(document, "完整预览矩阵85/85通过；生产Chromium在1366×768和390×844的参数优化/智能问答4项轻量冒烟全部通过。")

    document.save(TARGET)
    digest = hashlib.sha256(TARGET.read_bytes()).hexdigest().upper()
    print(f"updated={TARGET}")
    print(f"backup={BACKUP}")
    print(f"sha256={digest}")


if __name__ == "__main__":
    main()
