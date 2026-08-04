from __future__ import annotations

import argparse
import collections
import re
import sys
from pathlib import Path

from docx import Document


NUMBERED_HEADING = re.compile(r"^\s*(?:第[一二三四五六七八九十百]+[章节篇部分]|\d+(?:\.\d+){0,5})(?:[、.．\s]|$)")


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect paragraph styles and numbered heading candidates in a DOCX file.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--limit", type=int, default=500)
    parser.add_argument("--major-only", action="store_true")
    parser.add_argument("--find", action="append", default=[])
    args = parser.parse_args()

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    document = Document(args.docx)
    styles: collections.Counter[str] = collections.Counter()
    nonempty = 0
    candidates: list[tuple[int, str, str]] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        nonempty += 1
        style = paragraph.style.name if paragraph.style is not None else ""
        styles[style] += 1
        if args.major_only:
            if style.lower().startswith("heading") or style == "Title":
                candidates.append((index, style, text))
        elif style.lower().startswith("heading") or NUMBERED_HEADING.match(text):
            candidates.append((index, style, text))

    if args.find:
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if text and any(term in text for term in args.find):
                style = paragraph.style.name if paragraph.style is not None else ""
                print(f"FIND\t{index:05d}\t{style}\t{text}")

    print(f"paragraphs={len(document.paragraphs)} nonempty={nonempty} tables={len(document.tables)}")
    print("styles=" + repr(styles.most_common(20)))
    print(f"heading_candidates={len(candidates)}")
    for index, style, text in candidates[: max(0, args.limit)]:
        print(f"{index:05d}\t{style}\t{text}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
