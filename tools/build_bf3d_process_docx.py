#!/usr/bin/env python3
"""Build the traceable GL02 blast-furnace 3D production record.

The report embeds every image from the scoped P35/P36/P40/P50/P60 production
folders and the three browser-QA folders. Images are downsampled only inside
the DOCX; original evidence files remain untouched and are indexed in CSV.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import tempfile
import zipfile
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageFile, ImageOps
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ImageFile.LOAD_TRUNCATED_IMAGES = True

ROOT = Path(__file__).resolve().parents[1]
WORK_ROOT = ROOT / "PT" / "高炉3D模型" / "work"
OUTPUT_ROOT = ROOT / "PT" / "高炉3D模型"
DEFAULT_OUTPUT = OUTPUT_ROOT / "高炉3D模型制作过程与内切面验收记录_20260717.docx"
DEFAULT_MANIFEST = OUTPUT_ROOT / "高炉3D模型制作过程图片清单_20260717.csv"
DEFAULT_REPORT = OUTPUT_ROOT / "高炉3D模型制作过程文档生成报告_20260717.json"
PREVIEW_URL = (
    "http://127.0.0.1:8094/"
    "frontend_dashboard_v3.server.html?ws_port=8767#overview"
)
PRODUCTION_GLB_HASH = (
    "808960f1b2703e7fb27df35f1b1b1a17063b9b10d2267acba593fc3872b62af6"
)

QA_ROOTS = (
    ROOT / "logs" / "bf3d_cutaway_runtime_20260717",
    ROOT / "logs" / "bf3d_layer_matrix_20260717",
    ROOT / "logs" / "bf3d_cutaway_cross_engine_20260717",
)
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif"}
STAGE_PATTERN = re.compile(r"^P(?:35|36|40|50|60)", re.IGNORECASE)

STATUS_GREEN = "2F6B59"
STATUS_AMBER = "A66A14"
STATUS_RED = "A13A32"
STEEL_DARK = "33413E"
STEEL_MID = "5D6B67"
STEEL_LIGHT = "E7ECEA"
ACCENT = "1D7866"


@dataclass(frozen=True)
class ImageRecord:
    index: int
    stage: str
    root_tag: str
    category: str
    absolute_path: str
    relative_path: str
    root_relative_path: str
    filename: str
    width: int
    height: int
    format: str
    frames: int
    bytes: int
    sha256: str
    duplicate_count: int
    caption: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate the complete GL02 3D production and QA DOCX."
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument(
        "--expected-image-count",
        type=int,
        default=308,
        help="Fail if the scoped evidence count changes unexpectedly.",
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def scoped_image_paths() -> list[Path]:
    roots: list[Path] = []
    if WORK_ROOT.is_dir():
        roots.extend(
            child
            for child in WORK_ROOT.iterdir()
            if child.is_dir() and STAGE_PATTERN.match(child.name)
        )
    roots.extend(root for root in QA_ROOTS if root.is_dir())
    paths = {
        path.resolve()
        for root in roots
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    }

    def sort_key(path: Path) -> tuple[int, str]:
        text = path.as_posix().lower()
        if "/p35" in text:
            order = 35
        elif "/p36" in text:
            order = 36
        elif "/p40" in text:
            order = 40
        elif "/p50" in text:
            order = 50
        elif "/p60" in text:
            order = 60
        else:
            order = 90
        return order, text

    return sorted(paths, key=sort_key)


def stage_for(path: Path) -> str:
    try:
        name = path.relative_to(WORK_ROOT).parts[0]
    except ValueError:
        return "QA"
    match = re.match(r"^(P\d+)", name, re.IGNORECASE)
    return match.group(1).upper() if match else "OTHER"


def root_tag_for(path: Path) -> str:
    try:
        return path.relative_to(WORK_ROOT).parts[0]
    except ValueError:
        for qa_root in QA_ROOTS:
            try:
                path.relative_to(qa_root)
                return qa_root.name
            except ValueError:
                continue
    return path.parent.name


def root_relative_for(path: Path, root_tag: str) -> str:
    stage_root = WORK_ROOT / root_tag
    if stage_root.exists():
        return path.relative_to(stage_root).as_posix()
    for qa_root in QA_ROOTS:
        if qa_root.name == root_tag:
            return path.relative_to(qa_root).as_posix()
    return path.name


def category_for(path: Path, stage: str) -> str:
    lowered = {part.lower() for part in path.parts}
    if stage == "QA":
        return "浏览器验收"
    if "textures" in lowered:
        return "PBR贴图"
    if "review" in lowered or "rotation_frames" in lowered:
        return "视觉复核"
    if "renders" in lowered:
        return "渲染输出"
    return "阶段输出"


def caption_for(path: Path, stage: str) -> str:
    name = path.stem
    upper = name.upper()
    if "FAIL" in upper:
        return "历史失败取证：后续已修复并通过，不代表最终效果"
    if upper == "FORMAL_CUTAWAY_L16":
        return "正式模型 L16 内切面运行时"
    if upper == "P60_4K_CUTAWAY_L16":
        return "P60 4K 预检模型 L16 内切面运行时"
    engine = ""
    if upper.startswith("CHROMIUM_"):
        engine = "Chromium"
    elif upper.startswith("FIREFOX_"):
        engine = "Firefox"
    elif upper.startswith("WEBKIT_"):
        engine = "WebKit"
    if engine:
        viewport = re.search(r"(\d{3,4}X\d{3,4})", upper)
        vp = viewport.group(1).lower() if viewport else "代表视口"
        return f"{engine} {vp}：L7～L16 分层与内切面验收"
    if "P36_L7_L16_FRONT" in upper:
        return "P36 L7～L16 十层分层正视复核"
    if "P36_L7_L16_OBLIQUE" in upper:
        return "P36 L7～L16 十层分层斜视复核"
    channel_labels = {
        "BASECOLOR": "基础色",
        "ROUGHNESS": "粗糙度",
        "METALLIC": "金属度",
        "NORMALGL": "OpenGL 法线",
        "_AO": "环境遮蔽",
        "_ORM": "ORM 合并",
    }
    if "P50_GL02_" in upper:
        label = next(
            (value for key, value in channel_labels.items() if key in upper),
            "PBR",
        )
        resolution = "4K" if "4K" in upper else "阶段"
        return f"P50 {resolution} {label}贴图"
    rotation = re.search(r"P50_ROTATION_(\d{3})$", upper)
    if rotation:
        return f"P50 4K 主材质 24 方位旋转复核，第 {rotation.group(1)} 帧"
    if "P50_ROTATION_CONTACT_SHEET" in upper:
        return "P50 4K 主材质 24 方位旋转联系表"
    if "P50_ROTATION_PREVIEW" in upper:
        return "P50 4K 主材质 24 帧旋转源动画（Word 中展示首帧）"
    p50_review = {
        "BACK_SEAM_CLOSEUP": "背部 UV 接缝近景",
        "SHELL_ZONE_BOUNDARY_CLOSEUP": "炉壳分区边界近景",
        "NEAR_MIP": "近距离 MIP 复核",
        "FAR_MIP": "远距离 MIP 复核",
        "FRONT": "成品正视复核",
        "BACK": "成品背视复核",
        "LEFT": "成品左视复核",
        "RIGHT": "成品右视复核",
    }
    if upper.startswith("P50_REVIEW_"):
        label = next(
            (value for key, value in p50_review.items() if key in upper),
            "PBR 成品复核",
        )
        return f"P50 {label}"
    if "DIAGNOS" in upper or upper.startswith(tuple(f"{n:02d}_" for n in range(1, 20))):
        return f"P50 PBR 通道诊断：{name}"
    if upper.startswith("P40_"):
        renderer = "Cycles" if "CYCLES" in upper else "EEVEE"
        lighting = "中性光" if "NEUTRAL" in upper else "展示光"
        views = {
            "GLOBAL_FRONT": "整炉正视",
            "GLOBAL_BACK": "整炉背视",
            "GLOBAL_LEFT": "整炉左视",
            "GLOBAL_RIGHT": "整炉右视",
            "DETAIL_SHELL": "炉壳细节",
            "DETAIL_TUYERE": "风口细节",
            "DETAIL_TAPHOLE": "出铁口细节",
        }
        view = next((value for key, value in views.items() if key in upper), name)
        prefix = f"{stage} " if stage in {"P35", "P40"} else ""
        return f"{prefix}{renderer}／{lighting}／{view}"
    if stage == "P50":
        return f"P50 材质烘焙与复核：{name}"
    return f"{stage} 制作过程：{name}"


def build_records(paths: Sequence[Path]) -> list[ImageRecord]:
    raw: list[dict[str, object]] = []
    hashes: list[str] = []
    for path in paths:
        with Image.open(path) as image:
            width, height = image.size
            frames = int(getattr(image, "n_frames", 1))
            image_format = str(image.format or path.suffix.lstrip(".")).upper()
        digest = sha256_file(path)
        hashes.append(digest)
        stage = stage_for(path)
        root_tag = root_tag_for(path)
        raw.append(
            {
                "stage": stage,
                "root_tag": root_tag,
                "category": category_for(path, stage),
                "absolute_path": str(path),
                "relative_path": path.relative_to(ROOT).as_posix(),
                "root_relative_path": root_relative_for(path, root_tag),
                "filename": path.name,
                "width": width,
                "height": height,
                "format": image_format,
                "frames": frames,
                "bytes": path.stat().st_size,
                "sha256": digest,
                "caption": caption_for(path, stage),
            }
        )
    duplicate_counts = Counter(hashes)
    records: list[ImageRecord] = []
    for index, item in enumerate(raw, start=1):
        digest = str(item["sha256"])
        records.append(
            ImageRecord(
                index=index,
                duplicate_count=duplicate_counts[digest],
                **item,  # type: ignore[arg-type]
            )
        )
    return records


def set_run_font(run, east_asia: str, size: float | None = None) -> None:
    run.font.name = east_asia
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_layout(table) -> None:
    table.autofit = False
    table_pr = table._tbl.tblPr
    layout = table_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def configure_styles(document: DocumentType) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.3

    heading_specs = {
        "Title": ("黑体", 26, STEEL_DARK),
        "Subtitle": ("宋体", 13, STEEL_MID),
        "Heading 1": ("黑体", 16, STEEL_DARK),
        "Heading 2": ("黑体", 13, ACCENT),
        "Heading 3": ("黑体", 11, STEEL_MID),
    }
    for style_name, (font_name, size, color) in heading_specs.items():
        style = styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    custom_specs = {
        "BF3D Caption": ("宋体", 8.5, STEEL_MID),
        "BF3D Small Path": ("Consolas", 6.5, "59635F"),
        "BF3D Code": ("Consolas", 8.5, "24312E"),
        "BF3D Note": ("宋体", 9.5, STEEL_DARK),
    }
    for name, (font_name, size, color) in custom_specs.items():
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.05


def configure_sections(document: DocumentType) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)


def add_headers_and_footers(document: DocumentType) -> None:
    # Later sections remain linked to the first section, so write the shared
    # header/footer only once and avoid duplicated runs.
    for section in (document.sections[0],):
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("GL02 高炉 3D 制作与验收记录｜隔离预览，生产模型未替换")
        set_run_font(run, "宋体", 8)
        run.font.color.rgb = RGBColor.from_string(STEEL_MID)
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("第 ")
        set_run_font(run, "宋体", 8)
        add_field(run, "PAGE")
        run = paragraph.add_run(" 页 / 共 ")
        set_run_font(run, "宋体", 8)
        add_field(run, "NUMPAGES")
        run = paragraph.add_run(" 页")
        set_run_font(run, "宋体", 8)


def set_update_fields(document: DocumentType) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_toc(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    add_field(run, 'TOC \\o "1-3" \\h \\z \\u')
    note = document.add_paragraph(
        "如目录未自动刷新：在 Word 中按 Ctrl+A，再按 F9。", style="BF3D Note"
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullets(document: DocumentType, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: DocumentType, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_code_block(document: DocumentType, text: str) -> None:
    paragraph = document.add_paragraph(style="BF3D Code")
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, "Consolas", 8.5)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF2F1")
    p_pr.append(shd)


def add_note(document: DocumentType, text: str, color: str = "E9F3F0") -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["BF3D Note"]
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_status_table(document: DocumentType) -> None:
    rows = (
        ("P35 细节几何", "已批准", "24/24", STATUS_GREEN, "E1F0EA"),
        (
            "P36 L7～L16 分层",
            "已批准",
            "19/19＋结构 8 项",
            STATUS_GREEN,
            "E1F0EA",
        ),
        (
            "P40 固定 LookDev",
            "已批准",
            "15/15＋18 张复核",
            STATUS_GREEN,
            "E1F0EA",
        ),
        (
            "P50 4K PBR 母版",
            "待批准",
            "41/41＋24 帧旋转",
            STATUS_AMBER,
            "FFF2D8",
        ),
        (
            "P60 隔离 GLB",
            "仅预检",
            "20/20；缺 Khronos Validator",
            STATUS_AMBER,
            "FFF2D8",
        ),
        ("P70／生产替换", "未批准", "未执行", STATUS_RED, "F9E2DF"),
    )
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_repeat_table_layout(table)
    headers = ("阶段", "状态", "证据")
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, STEEL_DARK)
        cell.text = label
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            set_run_font(run, "黑体", 9)
    set_table_repeat_header(table.rows[0])
    for stage, status, evidence, color, fill in rows:
        cells = table.add_row().cells
        cells[0].text = stage
        cells[1].text = status
        cells[2].text = evidence
        set_cell_shading(cells[1], fill)
        cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(color)
        cells[1].paragraphs[0].runs[0].font.bold = True
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def proxy_for(
    record: ImageRecord,
    proxy_root: Path,
    proxy_cache: dict[str, Path],
    max_pixels: tuple[int, int] = (1500, 1100),
) -> Path:
    existing = proxy_cache.get(record.sha256)
    if existing is not None:
        return existing
    destination = proxy_root / f"{record.sha256[:24]}.jpg"
    source = Path(record.absolute_path)
    with Image.open(source) as opened:
        try:
            opened.seek(0)
        except EOFError:
            pass
        image = ImageOps.exif_transpose(opened).convert("RGBA")
        background = Image.new("RGBA", image.size, (19, 27, 32, 255))
        image = Image.alpha_composite(background, image).convert("RGB")
        image.thumbnail(max_pixels, Image.Resampling.LANCZOS)
        image.save(
            destination,
            format="JPEG",
            quality=86,
            optimize=True,
            progressive=True,
            subsampling="4:2:0",
        )
    proxy_cache[record.sha256] = destination
    return destination


def fit_dimensions(
    width_px: int,
    height_px: int,
    max_width_cm: float,
    max_height_cm: float,
) -> tuple[float, float]:
    scale = min(max_width_cm / width_px, max_height_cm / height_px)
    return width_px * scale, height_px * scale


def find_record(records: Sequence[ImageRecord], *path_parts: str) -> ImageRecord | None:
    lowered_parts = tuple(part.lower().replace("\\", "/") for part in path_parts)
    for record in records:
        path = record.relative_path.lower().replace("\\", "/")
        if all(part in path for part in lowered_parts):
            return record
    return None


def add_hero_image(
    document: DocumentType,
    record: ImageRecord | None,
    proxy_root: Path,
    proxy_cache: dict[str, Path],
    figure_counter: list[int],
    width_cm: float = 16.2,
) -> None:
    if record is None:
        return
    proxy = proxy_for(record, proxy_root, proxy_cache)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(proxy), width=Cm(width_cm))
    figure_counter[0] += 1
    caption = document.add_paragraph(style="BF3D Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(f"图 {figure_counter[0]}  {record.caption}")
    path_paragraph = document.add_paragraph(style="BF3D Small Path")
    path_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    path_paragraph.add_run(record.relative_path)


def add_compact_table(
    document: DocumentType,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_cm: Sequence[float] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_layout(table)
    for index, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_shading(cell, STEEL_DARK)
        cell.text = label
        if widths_cm:
            cell.width = Cm(widths_cm[index])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            set_run_font(run, "黑体", 8.5)
    set_table_repeat_header(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, (cell, value) in enumerate(zip(row.cells, row_values)):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=55, start=65, bottom=55, end=65)
            if widths_cm:
                cell.width = Cm(widths_cm[index])
            for run in cell.paragraphs[0].runs:
                set_run_font(run, "宋体", 8)


def add_cover(
    document: DocumentType,
    records: Sequence[ImageRecord],
    proxy_root: Path,
    proxy_cache: dict[str, Path],
) -> None:
    for _ in range(2):
        document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("GL02 高炉 3D 模型制作过程、4K 材质与浏览器内切面验收记录")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("P35 细节几何｜P36 L7～L16 分层｜P40 LookDev｜P50 4K PBR｜P60 隔离 GLB")
    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run("受控工程记录 · 隔离预览 · 正式生产 GLB 未替换")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(STATUS_AMBER)
    set_run_font(run, "黑体", 11)

    hero = find_record(records, "p60_4k_cutaway_l16.png")
    if hero:
        proxy = proxy_for(hero, proxy_root, proxy_cache)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(proxy), width=Cm(16.5))
        caption = document.add_paragraph(style="BF3D Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run("封面图：P60 4K 隔离模型，L16 与内切面组合运行时")

    document.add_paragraph()
    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    values = (
        ("文档版本", "V1.0"),
        ("生成日期", str(date(2026, 7, 17))),
        ("图片范围", f"全部 {len(records)} 个过程与验收图像文件"),
        ("当前隔离入口", PREVIEW_URL),
        ("生产模型 SHA-256", PRODUCTION_GLB_HASH),
    )
    for row, (label, value) in zip(table.rows, values):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], STEEL_LIGHT)
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        for cell in row.cells:
            set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, "宋体", 8.5)
    document.add_paragraph()
    note = document.add_paragraph(style="BF3D Note")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "内部矿焦层、软熔带、气流、滴落和炉缸液相均为参数化工业示意，"
        "不是权威施工 CAD、实时 CFD 或实时断面测量。"
    )
    document.add_page_break()


def add_main_report(
    document: DocumentType,
    records: Sequence[ImageRecord],
    proxy_root: Path,
    proxy_cache: dict[str, Path],
) -> None:
    figure_counter = [0]
    document.add_heading("目录", level=1)
    add_toc(document)
    document.add_page_break()

    document.add_heading("1. 项目目标与当前结论", level=1)
    document.add_paragraph(
        "本轮目标是在不破坏正式高炉资产和生产页面的前提下，完成炉壳细节、"
        "L7～L16 十层测温结构、哑光工业材质、4K PBR 烘焙、隔离 GLB 交付"
        "以及 Three.js 内切面显示，并用多浏览器与多视口证据验证。"
    )
    add_status_table(document)
    add_note(
        document,
        "停止线：P35、P36、P40 已批准；P50 仍为 KEEP_P50_PENDING_NOT_APPROVED；"
        "P60 仅 not_granted_preflight_only。Khronos glTF Validator、P70 性能与"
        "现场生产链路完成前，不得覆盖正式 gl02_blast_furnace.glb。",
        "FFF2D8",
    )
    document.add_paragraph(
        "正式生产模型未被覆盖，当前 SHA-256 为："
    )
    add_code_block(document, PRODUCTION_GLB_HASH)

    document.add_heading("2. 原始模型来源与不可破坏合同", level=1)
    add_code_block(
        document,
        "YAML 炉型与传感器参数\n"
        "→ generate_gl02_cad.py\n"
        "→ visual_glb.py 生成 GLB\n"
        "→ build123d_shell.py 生成 STEP\n"
        "→ Three.js CadFurnaceViewer",
    )
    add_bullets(
        document,
        (
            "原始模型合同：190 节点、29 网格、23 材质。",
            "115 个真实传感器，其中炉体测温点 80 个。",
            "L7～L16 十层，每层固定 A～H 八点。",
            "3 个静压力点：P_static_20m35、P_static_23m49、P_static_28m98。",
            "26 个风口、2 个出铁口、5 个炉体工艺段。",
            "APPROX_GL02_ 与内部结构是参数化工业示意，不等同权威 CAD 或实时工艺测量。",
        ),
    )

    document.add_heading("3. P35 炉壳、焊缝与风口细节", level=1)
    document.add_paragraph("状态：已批准；机器断言 24/24。")
    add_bullets(
        document,
        (
            "新增 10 道贴合炉型的水平轮廓焊带和 80 段错缝纵焊。",
            "新增 26 组风口根部焊缝、套管、法兰和锥形喷嘴。",
            "新增 156 个带倒角六角螺栓，共增加 31,680 三角面。",
            "边界边、非流形边、零面积面和游离顶点均为 0。",
            "与最近传感器净距约 0.140929m，高于 0.02m 门槛。",
            "原始几何、UV、矩阵、材质与 115 点合同未改变。",
        ),
    )
    add_hero_image(
        document,
        find_record(records, "P35_VISUAL_PREVIEW_20260717_1405", "GLOBAL_FRONT.png"),
        proxy_root,
        proxy_cache,
        figure_counter,
    )
    add_hero_image(
        document,
        find_record(records, "P35_VISUAL_PREVIEW_20260717_1405", "DETAIL_TUYERE.png"),
        proxy_root,
        proxy_cache,
        figure_counter,
    )
    add_code_block(
        document,
        "P35_DETAIL_GEOMETRY_APPROVED.blend\n"
        "SHA-256: 1873ebbc6ae7a67f9f83bda56cba00e08eb5d8595af0ecdc155774c1af25a7c6",
    )

    document.add_heading("4. P36 L7～L16 十层拆分", level=1)
    document.add_paragraph(
        "状态：已批准；机器断言 19/19，结构 GLB 8 项检查通过。"
        "采用非破坏式炉型贴合覆盖层，不切坏正式炉壳。"
    )
    boundaries = (
        ("L7", "16.860", "16.1225", "17.5975", "8"),
        ("L8", "18.335", "17.5975", "19.2300", "8"),
        ("L9", "20.125", "19.2300", "20.9925", "8"),
        ("L10", "21.860", "20.9925", "22.7855", "8"),
        ("L11", "23.711", "22.7855", "24.5760", "8"),
        ("L12", "25.441", "24.5760", "26.3060", "8"),
        ("L13", "27.171", "26.3060", "28.0360", "8"),
        ("L14", "28.901", "28.0360", "29.7660", "8"),
        ("L15", "30.631", "29.7660", "31.4960", "8"),
        ("L16", "32.361", "31.4960", "33.2260", "8"),
    )
    add_compact_table(
        document,
        ("层", "测温标高/m", "下边界/m", "上边界/m", "点位"),
        boundaries,
        (2.0, 3.1, 3.1, 3.1, 2.0),
    )
    add_bullets(
        document,
        (
            "10 个独立网格和 10 个独立材质；九处相邻边界最大误差为 0。",
            "新增 2,944 三角面和 2,112 顶点。",
            "默认全部隐藏，运行时只显示选中层；单层只增加 1 个 draw call。",
            "彩虹十层仅供审查，正式页面使用正常、关注、严重、无数据状态色。",
        ),
    )
    add_hero_image(
        document,
        find_record(
            records,
            "P36_LAYER_SEGMENTATION_20260717_P35_INTEGRATED",
            "P36_L7_L16_front.png",
        ),
        proxy_root,
        proxy_cache,
        figure_counter,
    )
    add_hero_image(
        document,
        find_record(
            records,
            "P36_LAYER_SEGMENTATION_20260717_P35_INTEGRATED",
            "P36_L7_L16_oblique.png",
        ),
        proxy_root,
        proxy_cache,
        figure_counter,
    )

    document.add_heading("5. P40 哑光工业 LookDev", level=1)
    document.add_paragraph("状态：已批准；机器断言 15/15，视觉复核 18 张。")
    add_bullets(
        document,
        (
            "四个全景：前、后、左、右；三个细节：炉壳、风口、出铁口。",
            "EEVEE 中性光 7 张、展示光 7 张；Cycles 展示光 4 张。",
            "实际启用 RTX 5070 Laptop GPU 的 OptiX，Cycles 为 48 samples。",
            "色彩管理使用 AgX；修复正交相机宽高比，全景保留约 15% 安全边。",
            "哑光风化钢、焊缝、板缝、风口法兰和六螺栓均独立复核。",
            "限制：出铁口正面细节仍较抽象；模型纵横比高，全景横向留白属正常。",
        ),
    )
    for part in (
        "P40_CYCLES_PRESENTATION_CAM_GLOBAL_FRONT.png",
        "P40_PRESENTATION_CAM_DETAIL_SHELL.png",
        "P40_PRESENTATION_CAM_DETAIL_TUYERE.png",
        "P40_PRESENTATION_CAM_DETAIL_TAPHOLE.png",
    ):
        add_hero_image(
            document,
            find_record(records, "P40_FIXED_LOOKDEV_20260717_P36_FINAL", part),
            proxy_root,
            proxy_cache,
            figure_counter,
        )

    document.add_heading("6. P50 法线、Metallic 与局部 AO 迭代", level=1)
    add_numbered(
        document,
        (
            "旧 1K 候选 NormalGL 被量化成平坦值，结论为 ITERATE。",
            "法线修复采用 4× 编码放大和运行时 normalTexture.scale=0.25。",
            "Metallic 近二值比例从 79.87% 降到 0.043%，消除盐粒式闪点来源。",
            "炉腹—炉腰横线确认为约 7.67° 的合法炉型折角，不是贴图断裂。",
            "局部 AO 使用白名单，只允许焊缝、加强圈、风口法兰和螺栓产生接触阴影。",
        ),
    )
    add_bullets(
        document,
        (
            "P50 AO R3：41/41 断言通过。",
            "distance=0.12m，strength=0.28，floor=0.78，32 samples。",
            "AO 最小值 0.780392，均值 0.996459，标准差 0.022720。",
            "非纯白占比 5.197%，未发现平台黑带或全炉整体压暗。",
        ),
    )
    add_note(
        document,
        "P50 AO R3 的人工停止线仍是 "
        "AO_MAP_LOCALIZED_VISUAL_CONSUMPTION_PENDING_P60；approval=not_granted。",
        "FFF2D8",
    )
    ao_candidates = [
        record
        for record in records
        if "P50_LOCAL_CONTACT_AO_1K_20260717_1702_R3" in record.relative_path
        and record.category in {"视觉复核", "渲染输出"}
    ][:4]
    for record in ao_candidates:
        add_hero_image(
            document, record, proxy_root, proxy_cache, figure_counter, width_cm=14.8
        )

    document.add_heading("7. P50 4K PBR 母版", level=1)
    document.add_paragraph("状态：41/41 机器断言通过，但尚未授予 P50 批准。")
    add_bullets(
        document,
        (
            "4096×4096；16px 图集间距；3 列、10 个 UV 岛。",
            "图集有效像素占比 57.297%，有效纹素密度约 102.321 texel/m。",
            "OptiX，64 samples。",
            "贴图：BaseColor、Roughness、Metallic、AO、NormalGL（OpenGL +Y）、ORM。",
            "ORM 通道：R=AO、G=Roughness、B=Metallic。",
            "24 帧整圈审查未发现明显背部硬接缝、盐粒噪点、UV 翻转、静态摩尔纹或尺度突变。",
        ),
    )
    add_code_block(
        document,
        "transparent=false\nopacity=1\ndepthWrite=true\nroughness=1\nmetalness=1",
    )
    add_note(
        document,
        "炉壳保持不透明，空间变化由烘焙贴图表达，防止透明化抹掉粗糙质感。"
        "停止线：KEEP_P50_PENDING_NOT_APPROVED；approval=not_granted。",
        "FFF2D8",
    )
    for part in (
        "P50_REVIEW_front.png",
        "P50_REVIEW_back_seam_closeup.png",
        "P50_REVIEW_shell_zone_boundary_closeup.png",
        "P50_ROTATION_CONTACT_SHEET_24.png",
    ):
        add_hero_image(
            document,
            find_record(records, "P50_MASTER_4K_20260717_R1", part),
            proxy_root,
            proxy_cache,
            figure_counter,
        )

    texture_rows = []
    for label, part in (
        ("BaseColor", "P50_GL02_BaseColor_4K.png"),
        ("Roughness", "P50_GL02_Roughness_4K.png"),
        ("Metallic", "P50_GL02_Metallic_4K.png"),
        ("NormalGL", "P50_GL02_NormalGL_4K.png"),
        ("AO", "P50_GL02_AO_4K.png"),
        ("ORM", "P50_GL02_ORM_4K.png"),
    ):
        record = find_record(records, "P50_MASTER_4K_20260717_R1", part)
        if record:
            texture_rows.append((label, f"{record.width}×{record.height}", record.relative_path))
    add_compact_table(
        document,
        ("贴图通道", "分辨率", "源文件"),
        texture_rows,
        (3.0, 3.0, 10.2),
    )

    document.add_heading("8. P60 隔离 GLB 预检", level=1)
    add_bullets(
        document,
        (
            "隔离模型大小 32,748,944 字节，SHA-256 为 "
            "39c2ecd9455c9bf85ce13847f6d6a32ddb79bb93ccec80ada250ac186c6e56cd。",
            "20/20 内部检查和 Blender 回读通过。",
            "240 节点、42 网格、34 材质；内嵌 3 张图片和 3 个纹理。",
            "炉壳为 OPAQUE，BaseColor alpha=1；ORM 同时供 Metallic-Roughness 与 Occlusion。",
            "无外部 URI，不依赖 Draco 或 Meshopt。",
            "保留 115 个传感器、80 个测温点、L7～L16 与 45 个内部工艺示意对象。",
        ),
    )
    add_note(
        document,
        "本机没有 Khronos glTF Validator，因此 P60 只能标记 "
        "not_granted_preflight_only。P60 阶段目录本身没有离线截图，视觉取证来自"
        "浏览器运行时 p60_4k_cutaway_L16.png。",
        "FFF2D8",
    )
    add_hero_image(
        document,
        find_record(records, "p60_4k_cutaway_l16.png"),
        proxy_root,
        proxy_cache,
        figure_counter,
    )

    document.add_heading("9. Three.js 外观与内切面设计", level=1)
    document.add_paragraph(
        "默认“外观”精确隐藏 45 个内部工艺示意对象：矿层、焦层、软熔带、冷风流、"
        "逆流煤气流线、铁水池、渣层、回旋区、36 个铁水滴落对象和内部参考环带。"
    )
    add_numbered(
        document,
        (
            "裁剪平面切开炉壳。",
            "先显示料柱与矿焦层。",
            "再显示软熔带，并做克制呼吸。",
            "最后显示气流、回旋区、滴落与炉缸流体。",
            "临时隐藏塔架/桥、风口总管、煤气管、平台和支撑等 5 组遮挡附件。",
            "开启暖、冷两盏炉内补光。",
            "显示两条按真实炉型折线生成的琥珀色剖切边缘。",
            "暂停保留当前断面；复位恢复附件和外观状态。",
        ),
    )
    add_code_block(
        document,
        "setCutawayMode('exterior' | 'cutaway')\n"
        "playCutaway()\n"
        "pauseCutaway()\n"
        "resetCutaway()\n"
        "getCutawayState()",
    )
    add_note(
        document,
        "固定提示：工艺示意，非实时断面测量。内切面可与 L7～L16 任一单层"
        "A～H 八点筛选同时使用。",
    )
    add_hero_image(
        document,
        find_record(records, "formal_cutaway_l16.png"),
        proxy_root,
        proxy_cache,
        figure_counter,
    )

    document.add_heading("10. 浏览器与视口验收", level=1)
    qa_rows = (
        ("正式 GLB 运行时", "通过", "ok=true；意外页面/控制台/HTTP 错误 0"),
        ("P60 4K 隔离候选", "通过", "ok=true；45/45 内部对象、5/5 遮挡组"),
        ("L7～L16", "通过", "每层可见 A～H 八点"),
        ("Chromium 固定视口", "9/9", "桌面、平板、手机矩阵"),
        ("Firefox", "4/4", "1920×1080、1366×768、768×1024、390×844"),
        ("WebKit", "4/4", "1920×1080、1366×768、768×1024、390×844"),
        ("Edge 静态专项", "通过", "1366×768"),
    )
    add_compact_table(
        document,
        ("验收项", "结果", "说明"),
        qa_rows,
        (4.0, 2.4, 9.8),
    )
    add_note(
        document,
        "上述浏览器矩阵使用静态数据源，允许 8767 WebSocket 与 /api/* 作为"
        "预期离线；它验证模型、交互、脚本、视口与错误状态，不代表已验证生产数据库"
        "实时数值和实时状态色。",
        "E9F3F0",
    )
    for part in (
        "chromium_1920x1080_L16.png",
        "chromium_390x844_L16.png",
        "firefox_1366x768_cutaway.png",
        "webkit_390x844_cutaway.png",
    ):
        add_hero_image(
            document,
            find_record(records, part),
            proxy_root,
            proxy_cache,
            figure_counter,
            width_cm=15.0,
        )

    document.add_heading("11. 图片归档口径", level=1)
    stage_counts = Counter(record.stage for record in records)
    category_counts = Counter(record.category for record in records)
    unique_count = len({record.sha256 for record in records})
    add_bullets(
        document,
        (
            f"工作范围内全部图像文件：{len(records)}。",
            f"唯一视觉内容：{unique_count}；迭代目录重复文件：{len(records) - unique_count}。",
            "P35/P36/P40/P50/P60 阶段目录与三组浏览器 QA 日志均已递归扫描。",
            "GIF 在 Word 中以首帧显示；同目录 24 帧联系表和 24 张旋转帧完整保留。",
            "历史 webkit_390x844_FAIL.png 明确标为失败取证；最终同视口已修复并通过。",
        ),
    )
    rows = [
        (stage, str(stage_counts[stage]), "")
        for stage in ("P35", "P36", "P40", "P50", "P60", "QA")
        if stage_counts[stage]
    ]
    rows.extend(
        (category, str(count), "按文件用途")
        for category, count in sorted(category_counts.items())
    )
    add_compact_table(document, ("分类", "数量", "备注"), rows, (6.0, 3.0, 7.2))

    document.add_heading("12. 复现、核验与后续停止线", level=1)
    add_code_block(
        document,
        "python tools\\serve_bf3d_preview.py --port 8094\n"
        f"# 打开：{PREVIEW_URL}\n\n"
        "node --check tools\\verify_gl02_cutaway_runtime.cjs\n"
        "node tools\\verify_gl02_cutaway_runtime.cjs\n"
        "node tools\\verify_gl02_cutaway_runtime.cjs "
        "--model=\"PT\\高炉3D模型\\work\\P60_PREFLIGHT_4K_20260717_R1\\"
        "P60_PREFLIGHT_4K_UNCOMPRESSED.glb\"\n\n"
        "python tools\\verify_gl02_layered_model_ui.py "
        "--base-url http://127.0.0.1:18092 "
        "--page-path /frontend_dashboard_v3.server.html "
        "--ws-port 8767 --allow-offline-data-source",
    )
    add_bullets(
        document,
        (
            "尚未完成：P50 最终独立批准。",
            "尚未完成：Khronos glTF Validator 与 P60 正式批准。",
            "尚未完成：P70 性能、现场 Edge、8767 与数据库实时链路。",
            "尚未完成：权威 CAD 尺寸和工艺结构复核。",
            "尚未完成：受控替换生产 GLB 的审批、备份与回滚。",
        ),
    )


def add_image_catalogue(
    document: DocumentType,
    records: Sequence[ImageRecord],
    proxy_root: Path,
    proxy_cache: dict[str, Path],
) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    document.add_heading("附录 A：全部 308 个过程与验收图像", level=1)
    document.add_paragraph(
        "本附录按实际目录逐图嵌入。为控制 Word 体积，文档内使用压缩代理图，"
        "原始图像不做任何修改；每项保留原尺寸、格式、相对路径和重复次数。"
    )

    grouped: dict[str, list[ImageRecord]] = defaultdict(list)
    for record in records:
        grouped[record.root_tag].append(record)

    for root_tag, group in grouped.items():
        document.add_heading(
            f"{root_tag}（{len(group)} 个文件）",
            level=2,
        )
        table = document.add_table(rows=0, cols=4)
        table.style = "Table Grid"
        set_repeat_table_layout(table)
        for start in range(0, len(group), 4):
            row = table.add_row()
            prevent_row_split(row)
            for column, cell in enumerate(row.cells):
                set_cell_margins(cell, top=45, start=45, bottom=45, end=45)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                record_index = start + column
                if record_index >= len(group):
                    set_cell_shading(cell, "F7F8F7")
                    continue
                record = group[record_index]
                proxy = proxy_for(record, proxy_root, proxy_cache)
                image_width, image_height = fit_dimensions(
                    record.width,
                    record.height,
                    max_width_cm=6.05,
                    max_height_cm=4.9,
                )
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(
                    str(proxy),
                    width=Cm(image_width),
                    height=Cm(image_height),
                )
                caption = cell.add_paragraph(style="BF3D Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = caption.add_run(
                    f"图 A-{record.index:03d}  {record.caption}\n"
                    f"{record.width}×{record.height}｜{record.format}"
                    + (f"｜{record.frames} 帧" if record.frames > 1 else "")
                    + (
                        f"｜重复 {record.duplicate_count} 份"
                        if record.duplicate_count > 1
                        else ""
                    )
                )
                set_run_font(run, "宋体", 7.1)
                path_paragraph = cell.add_paragraph(style="BF3D Small Path")
                path_run = path_paragraph.add_run(record.root_relative_path)
                set_run_font(path_run, "Consolas", 5.7)

    document.add_page_break()
    document.add_heading("附录 B：证据文件与追溯入口", level=1)
    evidence_rows = (
        ("总体制作说明", "PT/3D模型构建显示.md"),
        ("高炉 3D 总设计", "PT/高炉3D模型/总设计详细规划.md"),
        ("自动化追溯", "docs/automation_traceability.md"),
        ("测试索引", "docs/test_reference.md"),
        (
            "P60 预检报告",
            "PT/高炉3D模型/work/P60_PREFLIGHT_4K_20260717_R1/"
            "p60_preflight_report.json",
        ),
        (
            "正式模型内切面报告",
            "logs/bf3d_cutaway_runtime_20260717/formal_report.json",
        ),
        (
            "P60 4K 内切面报告",
            "logs/bf3d_cutaway_runtime_20260717/p60_4k_report.json",
        ),
        (
            "Chromium 视口矩阵",
            "logs/bf3d_layer_matrix_20260717/manifest.json",
        ),
        (
            "Firefox/WebKit 矩阵",
            "logs/bf3d_cutaway_cross_engine_20260717/manifest.json",
        ),
        ("完整图片 CSV", DEFAULT_MANIFEST.relative_to(ROOT).as_posix()),
        ("隔离预览启动器", "tools/serve_bf3d_preview.py"),
    )
    add_compact_table(
        document,
        ("证据类型", "仓库相对路径"),
        evidence_rows,
        (6.0, 20.5),
    )
    add_note(
        document,
        "历史失败截图保留用于说明问题如何被发现和修复；最终验收结论只以最终"
        "manifest、最终截图和 ok=true 报告为准。",
        "FFF2D8",
    )


def write_manifest(path: Path, records: Sequence[ImageRecord]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(asdict(records[0]).keys()) if records else []
    with path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=fields)
        writer.writeheader()
        for record in records:
            writer.writerow(asdict(record))


def build_document(
    output: Path,
    records: Sequence[ImageRecord],
    proxy_root: Path,
    proxy_cache: dict[str, Path],
) -> DocumentType:
    document = Document()
    document.core_properties.title = (
        "GL02 高炉 3D 模型制作过程、4K 材质与浏览器内切面验收记录"
    )
    document.core_properties.subject = "P35/P36/P40/P50/P60 过程与浏览器验收"
    document.core_properties.author = "冀南钢铁高炉项目 / Codex"
    document.core_properties.keywords = (
        "GL02, 高炉, Three.js, Blender, PBR, L7-L16, 内切面, DOCX"
    )
    document.core_properties.comments = (
        "隔离预览工程记录；正式生产 GLB 未替换。"
    )
    configure_styles(document)
    configure_sections(document)
    set_update_fields(document)
    add_cover(document, records, proxy_root, proxy_cache)
    add_main_report(document, records, proxy_root, proxy_cache)
    add_image_catalogue(document, records, proxy_root, proxy_cache)
    add_headers_and_footers(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return document


def verify_docx(
    output: Path,
    manifest: Path,
    records: Sequence[ImageRecord],
    document: DocumentType,
) -> dict[str, object]:
    with zipfile.ZipFile(output) as archive:
        bad_member = archive.testzip()
        media_members = [
            name for name in archive.namelist() if name.startswith("word/media/")
        ]
        document_xml = archive.read("word/document.xml").decode("utf-8")
        field_xml = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name == "word/document.xml"
            or name.startswith("word/header")
            or name.startswith("word/footer")
        )
    reopened = Document(output)
    return {
        "ok": bad_member is None and len(records) > 0,
        "generated_at": "2026-07-17",
        "docx": str(output),
        "docx_bytes": output.stat().st_size,
        "manifest": str(manifest),
        "source_image_count": len(records),
        "unique_source_sha256_count": len({record.sha256 for record in records}),
        "duplicate_source_file_count": len(records)
        - len({record.sha256 for record in records}),
        "source_stage_counts": dict(Counter(record.stage for record in records)),
        "source_category_counts": dict(
            Counter(record.category for record in records)
        ),
        "docx_media_part_count": len(media_members),
        "inline_shape_count": len(reopened.inline_shapes),
        "paragraph_count": len(reopened.paragraphs),
        "table_count": len(reopened.tables),
        "section_count": len(reopened.sections),
        "zip_bad_member": bad_member,
        "has_toc_field": 'TOC \\o "1-3"' in document_xml,
        "has_page_fields": "NUMPAGES" in field_xml and "PAGE" in field_xml,
        "has_production_hash": PRODUCTION_GLB_HASH in document_xml,
        "all_catalogue_indices_present": all(
            f"图 A-{record.index:03d}" in document_xml for record in records
        ),
    }


def main() -> int:
    args = parse_args()
    paths = scoped_image_paths()
    if len(paths) != args.expected_image_count:
        raise RuntimeError(
            f"Expected {args.expected_image_count} scoped images, found {len(paths)}"
        )
    records = build_records(paths)
    write_manifest(args.manifest, records)
    with tempfile.TemporaryDirectory(prefix="bf3d_docx_") as temp_dir:
        proxy_root = Path(temp_dir)
        proxy_cache: dict[str, Path] = {}
        document = build_document(
            args.output,
            records,
            proxy_root,
            proxy_cache,
        )
        report = verify_docx(
            args.output,
            args.manifest,
            records,
            document,
        )
        report["proxy_image_count"] = len(proxy_cache)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
