# -*- coding: utf-8 -*-
"""Export batch_input view from Vastbase in weekly batches to avoid timeout."""
from __future__ import annotations

import json
import os
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
import psycopg

OUTPUT_DIR = Path(r"D:\文件\服务器实际运行版\导出数据包\IMES_3months")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

CONN_INFO = {
    "host": "10.10.181.195",
    "port": 5432,
    "dbname": "vastbase",
    "user": "gl2#dmx",
    "password": "gl2#dmx!",
    "connect_timeout": 15,
    "options": "-c statement_timeout=120000",
}

TODAY = date.today().strftime("%Y%m%d")
START_DATE = date(2026, 2, 10)
END_DATE = date(2026, 5, 10)


def write_excel(rows, columns, filepath):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(columns)
    for row in rows:
        ws.append([str(v) if v is not None else "" for v in row])
    for col_idx, _ in enumerate(columns, 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 18
    if len(rows) > 0:
        ws.auto_filter.ref = ws.dimensions
    wb.save(filepath)
    return len(rows)


def batch_export(cur, sql, params, columns, filename):
    cur.execute(sql, params)
    rows = cur.fetchall()
    fp = OUTPUT_DIR / f"{filename}.xlsx"
    n = write_excel(rows, columns, fp)
    print(f"  -> {fp.name}: {n} rows")
    return n


def main():
    print(f"Connecting to Vastbase...")
    conn = psycopg.connect(**CONN_INFO)
    cur = conn.cursor()

    # Get columns
    cur.execute(
        "SELECT column_name FROM information_schema.columns "
        "WHERE table_schema='public' AND table_name='batch_input' ORDER BY ordinal_position"
    )
    all_columns = [c[0] for c in cur.fetchall()]
    print(f"Columns: {len(all_columns)}")

    # Collect all rows by week
    all_rows = []
    cursor_date = START_DATE
    week_num = 0

    while cursor_date <= END_DATE:
        week_end = min(cursor_date + timedelta(days=6), END_DATE)
        week_num += 1
        d1 = cursor_date.strftime("%Y-%m-%d")
        d2 = (week_end + timedelta(days=1)).strftime("%Y-%m-%d")

        # Use a fresh connection per week to avoid transaction issues
        cur.close()
        conn.close()
        conn = psycopg.connect(**CONN_INFO)
        cur = conn.cursor()

        print(f"  Week {week_num}: {d1} ~ {week_end.strftime('%Y-%m-%d')} ...", end=" ", flush=True)
        cur.execute(
            """SELECT * FROM public.batch_input
               WHERE workdate2 >= %s AND workdate2 < %s
               ORDER BY workdate2, lot, charge""",
            (d1, d2)
        )
        batch = cur.fetchall()
        all_rows.extend(batch)
        print(f"{len(batch)} rows (total: {len(all_rows)})")

        cursor_date = week_end + timedelta(days=1)

    # Write one Excel file
    print(f"\nWriting {len(all_rows)} rows to Excel...")
    fp = OUTPUT_DIR / f"05_批次投料详情_batch_input_{TODAY}.xlsx"
    n = write_excel(all_rows, all_columns, fp)
    print(f"  -> {fp.name}: {n} rows, {len(all_columns)} cols")

    # Summary
    summary = {
        "exported_at": datetime.now().isoformat(),
        "source": "10.10.181.195:5432/vastbase public.batch_input",
        "date_range": f"{START_DATE} ~ {END_DATE}",
        "total_rows": n,
        "columns": len(all_columns),
    }
    sp = OUTPUT_DIR / f"batch_input_summary_{TODAY}.json"
    sp.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nDone. Summary: {sp}")

    cur.close()
    conn.close()


if __name__ == "__main__":
    main()
