"""Build the heat-centric data model and dashboard operating guide.

Requirement: REQ-HEAT-CENTRIC-DASHBOARD-20260725
Output: docs/以炉次为中心的五张核心数据表与133点时间窗口说明_20260725.docx
"""

from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
from urllib.parse import quote
from urllib.request import urlopen
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "以炉次为中心的五张核心数据表与133点时间窗口说明_20260725.docx"
SCREENSHOT = (
    ROOT
    / "logs"
    / "heat_dashboard_20260725"
    / "heat_si_distribution_1440x900_20260726.png"
)
DASHBOARD_URL = "http://127.0.0.1:8890/heat"
ACCENT = "2F756E"
DARK = "243231"
LIGHT = "E7F0EE"
AMBER = "D8B45D"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: object, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[idx], DARK)
        if widths:
            table.rows[0].cells[idx].width = Cm(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            if widths:
                cells[idx].width = Cm(widths[idx])
        if len(table.rows) % 2 == 1:
            for cell in cells:
                shade(cell, "F4F8F7")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        lead, rest = text[: len(bold_prefix)], text[len(bold_prefix) :]
        paragraph.add_run(lead).bold = True
        paragraph.add_run(rest)
    else:
        paragraph.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_code(doc: Document, code: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF3F2")
    p_pr.append(shd)
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(8.5)


def fetch_json(path: str, timeout: int = 180) -> dict | None:
    try:
        with urlopen(f"http://127.0.0.1:8890{path}", timeout=timeout) as response:
            payload = json.load(response)
        return payload if payload.get("ok") else None
    except Exception:
        return None


def current_evidence() -> tuple[dict | None, dict | None]:
    heats = fetch_json("/api/heats?limit=24")
    if not heats or not heats.get("heats"):
        return heats, None
    preferred = next(
        (
            item
            for item in heats["heats"]
            if item.get("hot_metal_sample_count") and item.get("slag_sample_count")
        ),
        heats["heats"][0],
    )
    meltno = quote(str(preferred["meltno"]), safe="")
    detail = fetch_json(
        f"/api/heat-detail?meltno={meltno}&window=pre_tap&group=core",
        timeout=240,
    )
    return heats, detail


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for name, size, color in (
        ("Title", 24, DARK),
        ("Heading 1", 17, DARK),
        ("Heading 2", 14, ACCENT),
        ("Heading 3", 12, ACCENT),
    ):
        style = styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("以炉次为中心的五张核心数据表\n与133点传感器时间窗口说明")
    run.bold = True
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor.from_string(DARK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("包含 Vastbase 铁水、炉渣、烧结矿化验及本地炉次分析仪表盘").font.size = Pt(13)
    doc.add_paragraph()
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.add_run("GL02 2号高炉 · 只读数据分析设计").bold = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"生成时间：{datetime.now():%Y-%m-%d %H:%M}\n文档状态：本机可运行版本")
    doc.add_page_break()


def build_document() -> Path:
    heats, detail = current_evidence()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "1. 结论与使用边界", 1)
    add_body(
        doc,
        "本文把口语中的“高炉次数”统一规范为“炉次”。炉次是一次正式出铁作业的业务主键，"
        "当前以 IMES t_ipes_cond.meltno 为权威标识，以 opentime/closetime 为时间锚点。",
    )
    add_body(
        doc,
        "结论：你关于烧结矿化验的判断是正确的。Vastbase 中存在 "
        "public.v_qpes_sinter_machine_sample_insp_final，包含1号、2号烧结机的烧结矿试样及 "
        "TFe、CaO、MgO、SiO₂、Al₂O₃、P、TiO₂、MnO、Zn、Cr、R2、FeO、S、镁铝比、铝硅比、QD 等结果。",
        bold_prefix="结论：",
    )
    add_body(
        doc,
        "重要边界：烧结矿视图没有高炉 meltno，也没有“该批烧结矿实际进入哪一炉”的完整谱系。"
        "因此当前仪表盘只把开铁口前12小时样品作为上游原料背景；只有补齐烧结试样→料仓→矿批→"
        "batch_input→t_ipes_cond.sumbatchStart/sumbatchEnd 后，才能提升为炉次精确归属。",
        bold_prefix="重要边界：",
    )

    add_heading(doc, "2. 数据流与五张核心逻辑表", 1)
    add_code(
        doc,
        "IMES炉次/生产实绩 ─┐\n"
        "铁水与炉渣化验 ────┼─> heat_master ─> 质量表/窗口特征/对齐审计 ─> 炉次API ─> 8890炉次仪表盘\n"
        "烧结矿化验背景 ─────┤\n"
        "PostgreSQL 133点 ───┘",
    )
    add_body(
        doc,
        "这五张表是建议的数据产品合同。当前页面采用实时只读聚合，不在生产 Vastbase 或 PostgreSQL "
        "中创建实体表；后续若要落库，应创建独立分析 schema，并沿用同样的字段和审计规则。",
    )

    add_heading(doc, "2.1 heat_master：炉次主表", 2)
    add_table(
        doc,
        ["字段", "来源", "含义/规则"],
        [
            ["meltno", "t_ipes_cond.meltno", "主键，例如 2#20260725-327"],
            ["open_ts", "t_ipes_cond.opentime", "开铁口时间；预测窗口的终点"],
            ["close_ts", "t_ipes_cond.closetime", "堵铁口时间；出铁过程窗口终点"],
            ["duration_minutes", "close-open 或 tappingtime", "仅接受0—600分钟；异常负值不参与统计"],
            ["batch_start/end/count", "sumbatchStart/End/sumbatch", "炉次关联矿批范围的关键线索"],
            ["actual_iron_qty", "t_ipes_out_put.ironquan", "同炉次多条实绩求和，并保留明细条数"],
            ["slag_rate/shift", "t_ipes_cond", "渣比、班次；空值保持为空，不填0"],
        ],
        [3.3, 4.5, 8.6],
    )

    add_heading(doc, "2.2 heat_hot_metal_chemistry：铁水质量表", 2)
    add_table(
        doc,
        ["字段组", "来源", "对齐规则"],
        [
            ["sample_no/result_ts/sample_seq", "铁水命名视图", "保留每个样品，不覆盖复验样"],
            ["C/Si/Mn/P/S/Ti/V/Cr/Ni/Cu/As", "化验结果文本列", "安全转换为数值；空值不是0"],
            ["hot_metal_si_summary", "全部有效Si样本派生", "输出valid_count/min/median/max/spread；原始样本仍逐条保留"],
            ["meltno", "派生", "试样号 FYYMM-NNN-SSS 与炉号、年月、NNN匹配，再加24小时保护"],
            ["alignment_method", "审计字段", "sample_no_furnace_yymm_heat_tail_with_24h_guard"],
        ],
        [4.6, 4.8, 7.0],
    )
    add_body(
        doc,
        "“发布时间”是结果判定/审核时间，不是取样时间。它只能用于结果版本和周转时长，不能作为"
        "传感器特征窗口终点。多试样建模时必须显式选择最终放行样、指定序号或统计聚合。",
    )

    add_heading(doc, "2.3 heat_slag_chemistry：炉渣质量表", 2)
    add_table(
        doc,
        ["字段", "业务含义", "使用提示"],
        [
            ["sampleno/meltno/publish_ts", "炉渣样品和正式炉次", "meltno 精确关联；同炉多样全部保留"],
            ["TFe/FeO", "含铁与氧化性", "TFe 历史覆盖较低；不能设为每炉必有"],
            ["CaO/MgO/SiO2/Al2O3/TiO2", "主要氧化物", "用于炉渣制度、黏度和含钛背景分析"],
            ["R2/R3/R4", "二/三/四元碱度", "源系统结果优先；公式仍以化验室标准为准"],
            ["MgO_Al2O3/SiO2_Al2O3", "比值指标", "SiO2_Al2O3 当前历史可能全空"],
        ],
        [4.5, 4.7, 7.2],
    )

    add_heading(doc, "2.4 heat_sensor_window_features：炉次传感器窗口特征表", 2)
    add_table(
        doc,
        ["字段", "含义"],
        [
            ["meltno/window_kind/window_start/window_end", "炉次、窗口类型及左闭右开时间边界"],
            ["variable_name/chinese_name/tag_long_name", "稳定变量ID、中文语义与pSpace物理点"],
            ["mean/min/max/std/latest", "窗口均值、极值、标准差和末值"],
            ["slope_per_minute", "以时间为自变量的每分钟线性斜率"],
            ["sample_count/expected_minutes/coverage_ratio", "实有样本、应有分钟数和覆盖率"],
            ["status", "good≥95%；partial为部分数据；missing为无数据"],
        ],
        [6.5, 9.9],
    )

    add_heading(doc, "2.5 heat_alignment_audit：关联与质量审计表", 2)
    add_table(
        doc,
        ["审计对象", "状态/方法", "必须记录的异常"],
        [
            ["炉次", "exact_meltno", "缺失或不符合正式炉次格式"],
            ["铁水", "试样号语义映射＋24h保护", "无匹配、多个候选、结果时间过远"],
            ["炉渣", "exact_meltno", "meltno为空、同炉多样、关键成分缺失"],
            ["烧结矿", "time_context_only，低置信度", "不得升级成该炉实际入炉结论"],
            ["传感器", "窗口边界＋覆盖率", "缺点、低覆盖、末值陈旧、异常时长"],
        ],
        [4.1, 5.8, 6.5],
    )

    add_heading(doc, "3. 烧结矿及其他Vastbase化验对象", 1)
    add_table(
        doc,
        ["对象", "数据性质", "与炉次关系"],
        [
            ["v_qpes_inner_batch_insp_final_sample", "高炉铁水C/Si/Mn/P/S等", "试样号映射，非直接meltno"],
            ["v_qpes_slag_insoection_final", "高炉炉渣氧化物和碱度", "有meltno，优先精确关联"],
            ["v_qpes_sinter_machine_sample_insp_final", "1/2号烧结机烧结矿化验", "无meltno；当前仅时间背景"],
            ["v_qpes_mat_final", "铁水视图精简投影", "与铁水完整视图同源，禁止重复累计"],
            ["v_qpes_steel_final", "炼钢样品", "不是高炉铁水，不进入本炉次模型"],
        ],
        [6.6, 5.1, 4.7],
    )
    add_heading(doc, "3.1 烧结矿关键指标", 2)
    add_table(
        doc,
        ["指标", "解释", "当前注意事项"],
        [
            ["TFe", "烧结矿总铁含量", "核心品位指标"],
            ["FeO", "氧化亚铁", "反映氧化程度、还原性背景"],
            ["CaO/MgO/SiO2/Al2O3", "主要碱/酸性氧化物", "决定入炉脉石和炉渣制度背景"],
            ["R2", "CaO/SiO2二元碱度", "源值与计算关系一致，仍以化验标准为准"],
            ["P/S/Zn/Cr", "有害或残余元素", "关注入炉负荷与循环富集"],
            ["TiO2/MnO", "含钛/含锰背景", "用于护炉和元素平衡分析"],
            ["QD", "强度类指标", "确切试验名称、算法和单位待正式字典确认"],
        ],
        [3.2, 6.0, 7.2],
    )

    add_heading(doc, "4. 如何映射不同传感器的时间窗口", 1)
    add_table(
        doc,
        ["窗口", "公式", "适用分析"],
        [
            ["pre_tap", "[opentime−120min, opentime)", "质量预测、因果分析；默认窗口"],
            ["tapping", "[opentime, closetime]", "出铁过程、铁口温度、操作复盘"],
            ["inter_heat", "[上一炉closetime, 本炉opentime)", "两炉之间的恢复和过渡"],
        ],
        [3.5, 6.2, 6.7],
    )
    add_heading(doc, "4.1 传感器分组", 2)
    add_table(
        doc,
        ["API group", "点数", "内容", "建议窗口"],
        [
            ["core", "17", "顶压、压差、透气性、风量/风压/风温、富氧、喷煤、煤气利用率、料线、两铁口温度", "三种窗口均可"],
            ["static", "18", "20.35m/23.49m/28.98m三层×A—F静压力", "pre_tap、inter_heat"],
            ["body", "80", "L7—L16十层×A—H炉体温度", "pre_tap为主；可扩展4—8小时"],
            ["all", "133", "全部物理点", "建模/离线分析；先检查覆盖率"],
        ],
        [2.7, 1.6, 8.0, 4.1],
    )
    add_bullets(
        doc,
        [
            "温度点使用℃，压力与压差使用kPa，料线使用m；不同单位不得直接混画成同一数值轴。",
            "覆盖率低于95%的点标记partial；缺点不能填0。炉体80点可按层、方位分别聚合，不能只看总平均。",
            "铁水/炉渣化验发布时间不得进入pre_tap特征；它们是标签版本时间，不是因果输入。",
            "烧结矿若只有时间背景，建议形成过去6h/12h/24h滚动均值与波动，并保留low confidence。",
        ],
    )

    add_heading(doc, "5. API与访问方式", 1)
    add_body(doc, f"炉次仪表盘地址：{DASHBOARD_URL}")
    add_code(
        doc,
        "GET /api/heats?limit=24\n"
        "GET /api/heats?limit=24&format=xlsx\n"
        "GET /api/heat-detail?meltno=2%2320260725-327&window=pre_tap&group=core\n"
        "GET /api/heat-detail?meltno=2%2320260725-327&window=pre_tap&group=static\n"
        "GET /api/heat-detail?meltno=2%2320260725-327&window=pre_tap&group=body\n"
        "GET /api/heat-detail?meltno=2%2320260725-327&window=tapping&group=all",
    )
    add_table(
        doc,
        ["参数", "可选值", "说明"],
        [
            ["limit", "1—200", "最近炉次数"],
            ["meltno", "URL编码正式炉次号", "# 必须编码为 %23"],
            ["window", "pre_tap/tapping/inter_heat", "传感器时间窗口"],
            ["group", "core/static/body/all/custom", "物理点分组"],
            ["variables", "逗号分隔变量ID", "group=custom时使用"],
            ["pre_tap_minutes", "15—720", "默认120分钟"],
            ["format", "json/csv/xlsx", "当前炉次列表支持导出"],
        ],
        [3.5, 5.3, 7.6],
    )
    add_body(
        doc,
        "实现采用三类只读来源：Vastbase炉次/实绩账号、Vastbase化验视图账号和GL02 PostgreSQL只读账号。"
        "凭据只存在于本机环境或被.gitignore排除的配置文件，不返回浏览器。",
    )

    add_heading(doc, "6. 本地炉次仪表盘", 1)
    add_bullets(
        doc,
        [
            "炉次列表：开堵口、时长、实绩铁量、铁水Si中位数、最小—最大范围、样本数、炉渣R2/FeO和关联完整性。",
            "Si趋势：逐点显示全部铁水Si样本，叠加每炉中位数曲线和最小—最大范围；不再只用最后一个样本代表整炉。",
            "质量详情：保留同炉次的全部铁水、炉渣样品，并显示Si有效样本数、中位数、范围和极差，不静默覆盖或平均复验样。",
            "原料背景：显示开铁口前12小时烧结矿样品，并始终标注“仅时间背景”。",
            "传感器分析：支持核心17点、静压力18点、炉体温度80点和全部133点，输出均值、极值、标准差、斜率、末值和覆盖率。",
            "交互闭环：加载、失败、重试、空状态、刷新和XLSX导出均有真实动作。",
        ],
    )
    if SCREENSHOT.exists():
        doc.add_picture(str(SCREENSHOT), width=Cm(16.5))
        caption = doc.add_paragraph("图1  炉次质量与传感器分析页面（Chromium 1440×900）")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "7. 当前真实数据样例", 1)
    if heats:
        summary = heats.get("summary", {})
        add_table(
            doc,
            ["指标", "当前实读值"],
            [
                ["查询炉次", summary.get("heat_count")],
                ["实绩铁量合计", summary.get("actual_iron_qty_total")],
                ["有铁水化验炉次", summary.get("hot_metal_heat_count")],
                ["有炉渣化验炉次", summary.get("slag_heat_count")],
                ["铁水+炉渣完整炉次", summary.get("complete_heat_count")],
                ["平均有效出铁时长/分钟", summary.get("average_duration_minutes")],
                ["最新炉次", summary.get("latest_meltno")],
            ],
            [7.4, 9.0],
        )
    else:
        add_body(doc, "生成文档时本地API不可用，因此未嵌入当前实读快照；表结构和调用方式不受影响。")
    if detail:
        heat = detail["heat"]
        window = detail["sensor_window"]
        si = heat.get("hot_metal_si_summary") or {}
        add_body(
            doc,
            f"示例炉次 {heat['meltno']}：铁水样 {heat['hot_metal_sample_count']} 个，炉渣样 "
            f"{heat['slag_sample_count']} 个，开口前12小时烧结矿背景样 {len(heat.get('sinter_context') or [])} 个；"
            f"铁水Si中位数 {si.get('median')}%，范围 {si.get('min')}%—{si.get('max')}%；"
            f"{window['group']} 组 {window['point_count']} 点，平均覆盖率 "
            f"{window['average_coverage'] * 100:.2f}%，窗口炉况记录 {detail['diagnosis']['count']} 条。",
        )

    add_heading(doc, "8. 生产化前的下一步", 1)
    add_bullets(
        doc,
        [
            "与MES/化验室确认铁水试样序号、最终放行样规则及烧结矿QD正式定义。",
            "补齐烧结试样到料仓、物料、矿批和炉次批次范围的谱系，形成可审计heat_material_context。",
            "若落库，使用独立分析schema和幂等批次，不对Vastbase源表或bf_sensor生产表写入。",
            "对每炉特征固定数据版本、窗口边界、点位清单和覆盖率，模型训练集必须能够重放。",
        ],
    )
    add_body(
        doc,
        "相关程序：db_dashboard/heat_service.py、db_dashboard/heat.html、db_dashboard/server.py、"
        "tools/build_heat_centric_docx.py；测试：tests/test_heat_service.py。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"DOCX ZIP damaged at {bad}")
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
    print(f"bytes={path.stat().st_size}")
