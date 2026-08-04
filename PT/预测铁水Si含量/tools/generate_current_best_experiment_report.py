"""Generate the formal Chinese V13 best-method experiment report."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from generate_algorithm_flowchart import build_flowchart


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "当前最优铁水Si预测实验报告_20260727.docx"

EXPERIMENTS = {
    "V9": ROOT / "reports/experiments/EXP-SI-V9-MULTIMODEL-PARETO-8H-HISTORYFIXR12-004_20260727/metrics.json",
    "V13": ROOT / "reports/experiments/EXP-SI-V13-LAG-MOMENT-HISTORYFIXR12-001_20260727/metrics.json",
    "V14": ROOT / "reports/experiments/EXP-SI-V14-SEMANTIC-LATENT-HISTORYFIXR12-001_20260727/metrics.json",
    "V15": ROOT / "reports/experiments/EXP-SI-V15-BEST-RESPONSE-LAG-HISTORYFIXR12-001_20260727/metrics.json",
    "V16": ROOT / "reports/experiments/EXP-SI-V16-RESIDUAL-ALIGNED-HISTORYFIXR12-001_20260727/metrics.json",
    "V18": ROOT / "reports/experiments/EXP-SI-V18-CATBOOST-PLAIN-HISTORYFIXR12-001_20260727/metrics.json",
}


def load_metrics() -> dict[str, dict]:
    return {
        name: json.loads(path.read_text(encoding="utf-8"))
        for name, path in EXPERIMENTS.items()
    }


def set_run_font(run, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 12, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_body(document: Document, text: str, *, indent: bool = True):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 12)
    return paragraph


def add_bullet(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 12)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, "黑体", 16 if level == 1 else 14, bold=True)
    return paragraph


def add_code_block(document: Document, lines: list[str]):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    cell.text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        set_run_font(run, "宋体", 10.5)


def add_results_table(document: Document, metrics: dict[str, dict]):
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["实验", "核心方法", "预测试稳健分数", "历史MAE", "±0.02命中率", "结论"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True)
        set_cell_shading(cell, "D9EAD3")

    methods = {
        "V9": ("多专家Pareto", "历史点MAE最低"),
        "V13": ("时延矩＋Huber残差", "当前推荐主方法"),
        "V14": ("七组PLS潜因子", "负结果"),
        "V15": ("单一最佳时延", "退回V13"),
        "V16": ("残差对齐排序", "退回V13"),
        "V18": ("Plain CatBoost", "负结果"),
    }
    for name in ("V9", "V13", "V14", "V15", "V16", "V18"):
        selected = metrics[name]["selected"]
        row = table.add_row().cells
        method, conclusion = methods[name]
        values = [
            name,
            method,
            f"{selected['robust_mae_score']:.6f}",
            f"{selected['historical_test']['mae']:.6f}",
            f"{selected['historical_test']['hit_rate_abs_le_002']:.2%}",
            conclusion,
        ]
        for cell, text in zip(row, values):
            set_cell_text(cell, text)
        if name == "V13":
            for cell in row:
                set_cell_shading(cell, "FFF2CC")
    return table


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.bold = True


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("冀南钢铁2号高炉　铁水Si含量软测量实验")
    set_run_font(run, "宋体", 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("内部离线实验资料　未经未来盲测不得用于生产闭环")
    set_run_font(run, "宋体", 9)


def build_report() -> Path:
    metrics = load_metrics()
    flowchart = build_flowchart()
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    add_header_footer(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("当前最优铁水Si含量预测实验报告")
    set_run_font(run, "黑体", 22, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("基于133点趋势数据的可解释热状态软测量（V13）")
    set_run_font(run, "黑体", 16, bold=True)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("报告日期：2026年7月27日　　状态：离线实验候选")
    set_run_font(run, "宋体", 12)

    add_heading(document, "一、摘要")
    add_body(
        document,
        "在当前严格历史合同、固定时间折和既有数据条件下，V13是推荐的主实验方法。"
        "其核心是：从133点传感器趋势中恢复30～480分钟非重叠时延带，构造均值、"
        "波动、均方根、变异系数、标准化冲击和空间场模态；以预测时刻已经发布的"
        "前炉Si作为基线，使用LightGBM Huber残差模型预测本炉代表Si，并同步输出"
        "P10/P50/P90分布。",
    )
    add_body(
        document,
        "V13在4月、5月和固定6月验证折上的稳健分数为0.057979，历史7月344炉"
        "MAE为0.047108% Si，±0.02%命中率为29.36%。V9的单一历史段MAE"
        "0.047081略低，但其稳健分数0.058470、±0.02%命中率26.45%均不如V13。"
        "因此，按预先固定的跨月稳健准则，V13优于仅按历史点MAE选择V9。",
    )
    add_body(
        document,
        "当前结果尚未达到0.02% Si目标；真实、带炉次和取样时刻的铁水温度标签仍"
        "为0行，所以尚不能训练或宣称已获得铁水温度/炉内热状态联合预测模型。",
    )

    add_heading(document, "二、数据与防泄漏合同")
    for item in (
        "正式样本共2,291炉；训练、固定验证、历史测试分别为1,603、344、344炉。",
        "4月、5月采用扩展训练折，6月使用固定验证折；7月仅作历史比较，不参与特征、模型或阈值选择。",
        "使用MES正式meltno；同一prediction_cutoff_ts的炉次不能互相进入历史。",
        "所有传感器统计严格早于预测截止时刻；历史Si只允许使用截止时刻前已发布的结果。",
        "出铁口温度T_taphole_1/2仅为辅助代理，不得冒充真实铁水温度。",
        "当前完整多窗口统计覆盖115个物理变量；其余点位按数据可用性保持缺失，不填造伪值。",
    ):
        add_bullet(document, item)

    add_heading(document, "三、当前推荐的V13实验方法")
    picture = document.add_picture(str(flowchart), width=Cm(15.6))
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run(
        "图1　可解释热状态神经元驱动的联合预测算法流程"
    )
    set_run_font(run, "宋体", 10.5)
    add_heading(document, "3.1 输入分支", level=2)
    for item in (
        "热量输入：热风温度、理论燃烧温度、风量、风压、富氧和喷煤。",
        "煤气利用与炉顶状态：煤气利用率、炉顶温度、炉顶压力及A～D上升管压力。",
        "透气性与煤气流：PI、上/下/总压差及三高度A～F的18点静压力场。",
        "炉体热场：7～16层各方位炉身温度及圆周、纵向场模态。",
        "料柱运动：总料线、南北探尺及其多窗口变化。",
        "历史状态与化学成分：预测时刻可见的前炉Si、变化速度和已发布化验背景。",
    ):
        add_bullet(document, item)

    add_heading(document, "3.2 可解释语义神经元", level=2)
    add_body(
        document,
        "V10阶段新增2,660个非重叠时延带神经元和1,390个空间场模态；V13再从"
        "累计计数、均值和样本标准差恢复4,123个非重叠二阶矩神经元。特征覆盖"
        "0～30、30～60、60～120、120～240、240～480分钟，并表达短期冲击、"
        "长时蓄热、波动增强、圆周偏流和上下部不均匀性。",
    )

    add_heading(document, "3.3 训练与输出", level=2)
    for item in (
        "每个时间折内独立重建特征重要性，最终V13配方使用1,316项特征。",
        "先由已发布历史Si形成预测基线，再用Huber LightGBM学习残差。",
        "候选按三折平均MAE＋0.5×折间标准差选择，并以±0.02%命中率作为Pareto次目标。",
        "输出代表Si点预测及P10/P50/P90分布；历史P10～P90覆盖率为75.29%，平均宽度0.131554% Si。",
    ):
        add_bullet(document, item)

    add_heading(document, "四、V9～V18实验结果对比")
    add_results_table(document, metrics)
    add_body(
        document,
        "V14把七类业务组压缩为1～3个PLS潜因子，丢失了局部空间和多时标信息；"
        "V15每个变量只保留一个最佳滞后，但跨折完全稳定率仅41.77%；V16虽然让"
        "重要性目标与残差训练目标一致，仍未超过V13。V17 Ordered CatBoost两次"
        "在1,004秒时限内未完成，并留下最高约3.31 GB的残留进程；V18 Plain "
        "CatBoost完成后稳健分数0.060185、历史MAE 0.049539，也未超过V13。",
    )

    add_heading(document, "五、结论与适用边界")
    for item in (
        "当前推荐主方法：V13多时标时延矩＋空间场＋历史残差Huber LightGBM。",
        "当前单一历史段点MAE最低：V9，MAE为0.047081% Si；但不作为跨月总体最优。",
        "目前没有任何方法达到0.02% Si，不能把0.047～0.050%改写成已达到目标。",
        "V13仍是离线候选，不写生产设定值、不控制高炉、不自动发布操作建议。",
        "只有新未来数据累计不少于300炉且覆盖30天，并完成盲测后，才允许正式比较冻结候选。",
    ):
        add_bullet(document, item)

    add_heading(document, "六、220.12服务器复现实验计划")
    add_body(
        document,
        "2026年7月27日本机对10.30.220.12:22的TCP检查失败，固定SSH身份探测"
        "超时，因此本报告不包含伪造的远端运行结果。恢复SSLVPN后，应先只读核查"
        "远端Python、磁盘、代码与数据哈希，再在独立实验目录复现V13，不能重启"
        "8093、8768或覆盖生产模型。",
    )
    add_code_block(
        document,
        [
            "python -m si_semantic_engine.train_v13 `",
            "  --dataset data/processed/formal_v3_historyfix_r12_20260727/formal_heat_training_dataset.csv `",
            "  --heat-targets data/processed/formal_v3_historyfix_r12_20260727/formal_heat_targets.csv `",
            "  --samples data/processed/formal_v3_historyfix_r12_20260727/formal_samples.csv `",
            "  --sensor-catalog ../高炉3D模型/docs/GL02传感器点位清单.v1.json `",
            "  --temporal-dir data/processed/formal_v6_temporal_stats_8h_20260727 `",
            "  --output-dir reports/experiments/EXP-SI-V13-22012-REPRO-001",
        ],
    )

    add_heading(document, "七、下一阶段建议")
    for item in (
        "优先取得真实铁水温度、取样时刻、铁口、铁罐号和出铁阶段，建立温度与Si联合标签。",
        "继续运行V4前瞻盲测账本，停止使用当前7月历史段选择新模型。",
        "基于未来炉次分析V13在低Si、高Si、月份漂移和不同热状态代理区间的误差。",
        "多任务模型采用共享热状态概念层，分别输出铁水温度分布、Si分布和热状态概率。",
        "热损失分支需补充冷却壁水温差、流量和热负荷；当前数据不能用缺失信号训练伪神经元。",
    ):
        add_bullet(document, item)

    add_heading(document, "八、复现证据")
    evidence = (
        "V13指标：reports/experiments/EXP-SI-V13-LAG-MOMENT-HISTORYFIXR12-001_20260727/metrics.json",
        "V13选择审计：reports/experiments/EXP-SI-V13-LAG-MOMENT-HISTORYFIXR12-001_20260727/selection_audit.json",
        "V13无目标回放：reports/prospective_replay_v4_lagmoment_20260727.json",
        "V14～V15报告：reports/2026-07-27_V14_V15语义压缩与响应时延实验.md",
        "V16指标：reports/experiments/EXP-SI-V16-RESIDUAL-ALIGNED-HISTORYFIXR12-001_20260727/metrics.json",
        "V18指标：reports/experiments/EXP-SI-V18-CATBOOST-PLAIN-HISTORYFIXR12-001_20260727/metrics.json",
    )
    for item in evidence:
        add_bullet(document, item)

    document.core_properties.title = "当前最优铁水Si含量预测实验报告"
    document.core_properties.subject = "V13可解释热状态软测量"
    document.core_properties.author = "冀南钢铁高炉智能助手项目"
    document.core_properties.comments = "离线实验报告"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
