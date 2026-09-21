"""Private source freeze fails closed before writing or publishing source text."""
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace
import sys

from docx import Document
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import freeze_qa_source_scope_candidate as freeze


def fixture_document():
    source = Document()
    def role(code):
        return '喷吹工' if code == 19 else f'岗位甲{code}' if code < 27 else '岗位交接班制度' if code == 27 else '生产联系确认制'
    for code in range(1, 29):
        title = role(code)
        source.add_paragraph(f'{code}. {title}……{code}')
    for code in range(1, 29):
        title = role(code)
        source.add_paragraph(f'{code}. {title}')
        if code < 27:
            source.add_paragraph('安全操作规程')
        source.add_paragraph(f'1. 原书测试条款{code}。')
    return source


def test_changed_original_identity_blocks_before_extraction(monkeypatch):
    monkeypatch.setattr(freeze, 'extract_scope', lambda *a, **kw: pytest.fail('Must reject identity first'))
    with pytest.raises(ValueError, match='identity changed'):
        freeze.prepare(None, 'different-source')


@pytest.mark.parametrize('gate', ['items', 'chunks'])
def test_independent_gate_failure_prevents_candidate(gate, monkeypatch):
    monkeypatch.setattr(freeze, 'validate_' + gate, lambda *a: {'verified': False})
    with pytest.raises(ValueError, match='Independent source ' + gate):
        freeze.prepare(fixture_document(), freeze.SOURCE_SHA)


def test_report_and_manifest_do_not_contain_original_body(monkeypatch):
    source = fixture_document()
    contents = []
    freeze.extract_scope(source, freeze.builder.REGULATION_HEADING_ALIASES,
                         freeze.builder.CHAPTER_HEADING_ALIASES, content_sink=contents.append)
    monkeypatch.setattr(freeze, 'NEW_AUTHORITY_SHA', freeze.digest('\n'.join(contents)))
    manifest, candidate, report = freeze.prepare(source, freeze.SOURCE_SHA)
    assert candidate['full_text'] == '\n'.join(contents)
    assert '原书测试条款' not in json.dumps(manifest, ensure_ascii=False)
    assert '原书测试条款' not in json.dumps(report, ensure_ascii=False)
    assert report['database_writes'] == report['model_calls'] == report['question_posts'] == 0
    assert candidate['knowledge_search_mode'] == 'keyword'
    assert candidate['embedding_generation'] is report['production_applied'] is False
    assert report['model_digest'] == 'e4ad74c41d68de1c8004419d8141a2b2df2275fa08f0dcf326ca0e63fb6d8124'
    assert report['semantic_passed'] == 0


@pytest.mark.parametrize('location', ['outside', 'root', 'existing'])
def test_public_or_reused_output_is_rejected(location, tmp_path, monkeypatch):
    monkeypatch.setattr(freeze, 'ROOT', tmp_path)
    private_root = tmp_path / '.codex_runtime/qa-source-scope-20260917'
    target = tmp_path / 'public' if location == 'outside' else private_root if location == 'root' else private_root / 'existing'
    if location == 'existing':
        target.mkdir(parents=True)
    monkeypatch.setattr(freeze.subprocess, 'run', lambda *a, **kw: pytest.fail('Path must reject before Git'))
    with pytest.raises(ValueError, match='unused private'):
        freeze.private_target(target)


def test_unignored_candidate_is_rejected_without_creation(tmp_path, monkeypatch):
    monkeypatch.setattr(freeze, 'ROOT', tmp_path)
    target = tmp_path / '.codex_runtime/qa-source-scope-20260917/new'
    monkeypatch.setattr(freeze.subprocess, 'run', lambda *a, **kw: SimpleNamespace(returncode=1))
    with pytest.raises(ValueError, match='Git ignored'):
        freeze.private_target(target)
    assert not target.exists()


def test_main_parses_the_exact_bytes_used_for_identity(tmp_path, monkeypatch, capsys):
    source = tmp_path / 'source.docx'
    raw = b'original snapshot'
    source.write_bytes(raw)
    target = tmp_path / 'candidate'
    monkeypatch.setattr(sys, 'argv', ['freeze', '--docx', str(source), '--output', str(target)])
    monkeypatch.setattr(freeze, 'private_target', lambda value: target)
    def parse(value):
        assert isinstance(value, BytesIO) and value.getvalue() == raw
        source.write_bytes(b'changed after read')
        return 'parsed-original'
    def prepare(document, sha):
        assert document == 'parsed-original' and sha == freeze.hashlib.sha256(raw).hexdigest()
        raise ValueError('Stop before freeze')
    monkeypatch.setattr(freeze, 'Document', parse)
    monkeypatch.setattr(freeze, 'prepare', prepare)
    with pytest.raises(ValueError, match='Stop before freeze'):
        freeze.main()
    assert not target.exists() and capsys.readouterr().out == ''
