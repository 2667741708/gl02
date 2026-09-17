"""BUG-QA-REGULATION-CLAUSE-DROP-20260917: headings cannot consume clauses/tables."""
import importlib.util
from pathlib import Path
import sys

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('source_heading_builder', ROOT / 'tools/build_three_rules_hierarchical_kb.py')
builder = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = builder
spec.loader.exec_module(builder)


@pytest.mark.parametrize('role,label', [('岗位甲', '安全操作规程'), ('岗位甲', '技术操作规程'),
                                      ('岗位甲', '设备使用维护规程'), ('岗位交接班制度', '岗位交接班制度'),
                                      ('生产联系确认制', '生产联系确认制')])
@pytest.mark.parametrize('prefix', ['', '1. ', '一、', '（一）'])
def test_explicit_heading_is_recognized(role, label, prefix):
    assert builder.match_regulation(prefix + label + '：', role) == label


def test_prefixed_role_heading_and_process_alias():
    assert builder.match_regulation('岗位甲安全操作规程', '岗位甲') == '安全操作规程'
    assert builder.match_regulation('工艺操作规程', '岗位甲') == '技术操作规程'


@pytest.mark.parametrize('heading,expected', [('工艺技术规程', '技术操作规程'), ('工艺技术操作规程', '技术操作规程'),
                                            ('工艺技术技作规程', '技术操作规程'), ('设备维护规程', '设备使用维护规程'),
                                            ('设备维护操作规程', '设备使用维护规程'), ('设备操维护规程', '设备使用维护规程'),
                                            ('三.、设备维护规程', '设备使用维护规程'), ('三、.设备操维护规程', '设备使用维护规程')])
def test_exact_heading_variants_from_original_source(heading, expected):
    assert builder.match_regulation(heading, '岗位甲') == expected
    assert builder.match_regulation('必须遵守' + heading + '执行。', '岗位甲') is None


@pytest.mark.parametrize('text', ['1. 双方确认交接记录。', '根据岗位交接班制度做好交接记录。',
                                '必须遵守生产联系确认制。', '按照安全操作规程执行，先确认后操作。',
                                '检查设备使用维护规程规定的项目。', '事项 | 要求\n交接 | 双方确认',
                                '安全操作规程 | 已确认', '安全操作规程\n1. 确认后操作。'])
@pytest.mark.parametrize('role', ['岗位甲', '岗位交接班制度', '生产联系确认制'])
def test_clause_or_table_is_not_a_heading(text, role):
    assert builder.match_regulation(text, role) is None


def source_document():
    document = Document()
    names = {i: f'岗位{i}' for i in range(1, 29)}
    names[27], names[28] = '岗位交接班制度', '生产联系确认制'
    for number, name in names.items():
        document.add_paragraph(f'{number}. {name}……{number}')
    for number, name in names.items():
        document.add_paragraph(f'{number}. {name}' + ('三规二制' if number < 27 else ''))
        if number < 27:
            document.add_paragraph('安全操作规程')
        document.add_paragraph('1. 双方确认交接记录。')
        if number == 27:
            document.add_paragraph('2. 根据岗位交接班制度做好交接记录。')
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = '事项', '要求'
            table.cell(1, 0).text, table.cell(1, 1).text = '交接', '双方确认'
        if number == 28:
            document.add_paragraph('2. 必须遵守生产联系确认制。')
    return document


def test_actual_source_parser_preserves_short_system_clauses_and_table():
    chapters, items = builder.parse_source_items(source_document())
    assert len(chapters) == 28
    handover = [item for item in items if item.chapter_code == '27']
    contact = [item for item in items if item.chapter_code == '28']
    assert [item.text for item in handover] == ['1. 双方确认交接记录。', '2. 根据岗位交接班制度做好交接记录。', '事项 | 要求\n交接 | 双方确认']
    assert [item.text for item in contact] == ['1. 双方确认交接记录。', '2. 必须遵守生产联系确认制。']
    assert all(item.regulation_type == '岗位交接班制度' for item in handover)
    assert all(item.regulation_type == '生产联系确认制' for item in contact)
    chunks = builder.build_chunks(items)
    for original in handover + contact:
        atomic = [chunk for chunk in chunks if chunk.granularity == 'atomic' and chunk.chapter_code == original.chapter_code]
        assert any(chunk.content == original.text for chunk in atomic)


def test_mentions_do_not_reassign_regulation_or_discard_original_clause():
    document = source_document()
    # Append a new body for role 1 after all preceding roles to isolate the boundary.
    document.add_paragraph('1. 岗位1三规二制')
    document.add_paragraph('技术操作规程')
    clause = '1. 按照安全操作规程执行，先确认后操作。'
    document.add_paragraph(clause)
    _, items = builder.parse_source_items(document)
    assert items[-1].text == clause and items[-1].regulation_type == '技术操作规程'


@pytest.mark.parametrize('text', ['1. 岗位甲负责交接检查。', '1. 根据岗位甲要求检查。',
                                '1. 岗位', '1. 岗位甲 | 操作要求', '1. 岗位甲……1'])
def test_role_mentions_prefixes_tables_and_toc_are_not_body_chapters(text):
    assert builder.match_chapter(text, {1: '岗位甲'}) is None


@pytest.mark.parametrize('text,chapters,expected', [('1. 工长', {1: '高炉工长'}, ('1', '高炉工长')),
                                                 ('18.喷煤制粉工“三规一制”', {18: '喷煤制粉工'}, ('18', '喷煤制粉工')),
                                                 ('19.喷煤喷吹工“三规一制”', {19: '喷吹工'}, ('19', '喷吹工'))])
def test_exact_role_aliases_and_quoted_source_marker(text, chapters, expected):
    assert builder.match_chapter(text, chapters) == expected


def test_source_role_alias_is_bound_to_its_chapter_only():
    assert builder.match_chapter('18.喷煤喷吹工“三规一制”', {18: '喷吹工'}) is None
    assert builder.match_chapter('19.喷煤喷吹工负责确认。', {19: '喷吹工'}) is None


def test_conflicting_source_directory_fails_closed():
    document = source_document()
    document.paragraphs[1].insert_paragraph_before('1. 另一个岗位……1')
    with pytest.raises(ValueError, match='目录编号'):
        builder.toc_chapters(document)


def test_toc_does_not_assign_prebody_text_to_last_system_chapter():
    document = source_document()
    document.paragraphs[28].insert_paragraph_before('目录后的说明。')
    _, items = builder.parse_source_items(document)
    assert all(item.text != '目录后的说明。' for item in items)
