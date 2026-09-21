"""Independent DOCX position oracle catches source omissions and wrong scope."""
import copy
import importlib.util
from pathlib import Path
import sys

from docx import Document
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools'))
import qa_source_scope_contract as contract
import build_three_rules_hierarchical_kb as builder


def document():
    source = Document()
    for i in range(1, 29):
        name = f'岗位甲{i}' if i < 27 else '岗位交接班制度' if i == 27 else '生产联系确认制'
        source.add_paragraph(f'{i}. {name}……{i}')
    source.add_paragraph('前言说明。')
    for i in range(1, 29):
        name = f'岗位甲{i}' if i < 27 else '岗位交接班制度' if i == 27 else '生产联系确认制'
        source.add_paragraph(f'{i}. {name}')
        if i < 27:
            source.add_paragraph('安全操作规程')
        source.add_paragraph('1. 确认交接事项。')
        if i == 27:
            table = source.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = '事项', '要求'
            table.cell(1, 0).text, table.cell(1, 1).text = '交接', '双方确认'
    return source


def scope(source):
    return contract.extract_scope(source, builder.REGULATION_HEADING_ALIASES)


def test_all_original_body_positions_and_table_match_actual_source_items():
    source = document()
    oracle = scope(source)
    _, items = builder.parse_source_items(source)
    result = contract.validate_items(oracle, items)
    assert result['verified'] and result['order_matches']
    assert result['expected_items'] == result['actual_items'] == 29
    assert not result['semantic_verified']
    assert len(oracle['roles']) == 28 and len(oracle['regulation_headings']) == 26
    assert len([row for row in oracle['items'] if row['kind'] == 'table']) == 1
    assert '确认交接事项' not in str(oracle) and '前言说明' not in str(oracle)


@pytest.mark.parametrize('mutation', ['omission', 'duplicate', 'order', 'role', 'role_title', 'regulation', 'block', 'content'])
def test_independent_oracle_rejects_item_mutation(mutation):
    source = document()
    oracle = scope(source)
    _, original = builder.parse_source_items(source)
    items = copy.deepcopy(original)
    if mutation == 'omission':
        items.pop(0)
    elif mutation == 'duplicate':
        items.append(copy.deepcopy(items[0]))
    elif mutation == 'order':
        items[0], items[1] = items[1], items[0]
    else:
        field = {'role': 'chapter_code', 'role_title': 'chapter_title', 'regulation': 'regulation_type', 'block': 'block_index', 'content': 'text'}[mutation]
        setattr(items[0], field, 999 if mutation == 'block' else 'invalid')
    assert not contract.validate_items(oracle, items)['verified']


def test_oracle_does_not_call_index_extractors_or_metadata_functions(monkeypatch):
    source = document()
    for name in ('parse_source_items', 'match_chapter', 'match_regulation', 'iter_blocks', 'toc_chapters'):
        monkeypatch.setattr(builder, name, lambda *args: pytest.fail('Independent oracle called index builder'))
    assert len(scope(source)['items']) == 29


@pytest.mark.parametrize('change', ['duplicate', 'out_of_order', 'missing', 'directory_inside_body', 'untyped_content'])
def test_ambiguous_source_layout_blocks_candidate(change):
    source = document()
    if change == 'duplicate':
        source.add_paragraph('28. 生产联系确认制')
    elif change == 'out_of_order':
        source.paragraphs[29].text = '2. 岗位甲2'
    elif change == 'missing':
        source.paragraphs[29].text = '非岗位标题'
    elif change == 'directory_inside_body':
        source.add_paragraph('1. 岗位甲1……1')
    else:
        source.paragraphs[30].text = '未分类正文。'
    with pytest.raises(contract.SourceScopeUnverified):
        scope(source)


def test_source_heading_inside_multiline_paragraph_has_independent_position():
    source = document()
    source.paragraphs[29].text += '\n安全操作规程\n1. 额外确认。'
    oracle = scope(source)
    _, items = builder.parse_source_items(source)
    assert contract.validate_items(oracle, items)['verified']
    assert any(row['block_index'] == 29 and row['line_index'] == 2 for row in oracle['items'])
    contents = []
    oracle = contract.extract_scope(source, builder.REGULATION_HEADING_ALIASES, content_sink=contents.append)
    assert contract.validate_chunks(oracle, builder.build_chunks(items), contents)['verified']


def test_role_alias_must_be_explicitly_bound_to_source_directory():
    with pytest.raises(contract.SourceScopeUnverified, match='alias'):
        contract.extract_scope(document(), builder.REGULATION_HEADING_ALIASES, {(19, '别名'): '不是源目录岗位'})


def test_invalid_source_heading_policy_fails_closed():
    with pytest.raises(contract.SourceScopeUnverified, match='heading policy'):
        contract.extract_scope(document(), {'未知标题': '不是正式类别'})


def test_single_cell_table_is_content_even_when_it_looks_like_a_heading():
    source = document()
    source.add_table(rows=1, cols=1).cell(0, 0).text = '安全操作规程'
    oracle = scope(source)
    _, items = builder.parse_source_items(source)
    assert items[-1].kind == 'table' and items[-1].text == '安全操作规程'
    assert contract.validate_items(oracle, items)['verified']


@pytest.mark.parametrize('mutation', ['none', 'missing_atomic', 'missing_section', 'section_tail', 'wrong_header', 'role_title',
                                    'duplicate_id', 'range', 'kind', 'topic_reorder'])
def test_independent_chunk_contract_rejects_omissions_and_wrong_bindings(mutation):
    source = document()
    contents = []
    oracle = contract.extract_scope(source, builder.REGULATION_HEADING_ALIASES, content_sink=contents.append)
    _, items = builder.parse_source_items(source)
    chunks = builder.build_chunks(items)
    if mutation == 'missing_atomic':
        chunks.pop(0)
    elif mutation == 'missing_section':
        chunks.pop(next(i for i, c in enumerate(chunks) if c.granularity == 'section'))
    elif mutation == 'section_tail':
        item = next(c for c in chunks if c.granularity == 'section')
        item.content = item.content[:-1]
        item.content_hash = contract.digest(item.content)
    elif mutation == 'wrong_header':
        chunks[0].enriched_content = chunks[0].enriched_content.replace('【岗位/制度】1.', '【岗位/制度】2.')
    elif mutation == 'role_title':
        chunks[0].chapter_title = '错误岗位'
    elif mutation == 'duplicate_id':
        chunks[1].chunk_id = chunks[0].chunk_id
    elif mutation == 'range':
        chunks[0].source_block_end += 200
    elif mutation == 'kind':
        chunks[0].granularity = 'unknown'
    elif mutation == 'topic_reorder':
        # Add a forged aggregate under the two-row system scope; the lines are
        # genuine source lines but their order is wrong.
        base = copy.deepcopy(next(c for c in chunks if c.chapter_code == '27' and c.granularity == 'section'))
        base.granularity, base.chunk_id = 'topic', 'synthetic-reversed-topic'
        base.content = '\n'.join(reversed(base.content.splitlines()))
        base.content_hash = contract.digest(base.content)
        chunks.append(base)
    result = contract.validate_chunks(oracle, chunks, contents)
    assert result['verified'] is (mutation == 'none')
    assert not result['topic_completeness_verified'] and not result['semantic_verified']


def test_unbound_private_source_rows_cannot_validate_chunks():
    source = document()
    contents = []
    oracle = contract.extract_scope(source, builder.REGULATION_HEADING_ALIASES, content_sink=contents.append)
    _, items = builder.parse_source_items(source)
    contents[0] = '改写原文'
    with pytest.raises(contract.SourceScopeUnverified, match='content binding'):
        contract.validate_chunks(oracle, builder.build_chunks(items), contents)


def test_empty_table_cells_keep_delimiters_under_declared_whitespace_normalization():
    source = document()
    table = source.add_table(rows=2, cols=3)
    table.cell(0, 1).text = '确认要求'
    table.cell(1, 0).text = '事项'
    oracle = scope(source)
    _, items = builder.parse_source_items(source)
    assert items[-1].text == '| 确认要求 |\n事项 | |'
    assert contract.validate_items(oracle, items)['verified']
