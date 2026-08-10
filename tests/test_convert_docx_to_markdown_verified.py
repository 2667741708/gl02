from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "tools" / "convert_docx_to_markdown_verified.py"
SPEC = importlib.util.spec_from_file_location("convert_docx_verified", MODULE_PATH)
assert SPEC is not None and SPEC.loader is not None
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class VerifiedDocxMarkdownConversionTests(unittest.TestCase):
    def test_preserves_ordered_paragraph_and_real_table_cell_text(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / "source.docx"
            output = temp / "output.md"
            report_path = temp / "report.json"

            document = Document()
            document.add_heading("三规二制示例", level=1)
            document.add_paragraph("5.3 炉况调剂方法")
            document.add_paragraph("5.3.1 上部调剂")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "条件|一"
            table.cell(0, 1).text = "动作\n复核"
            document.save(source)

            report = MODULE.convert_document(source, output, report_path)
            markdown = output.read_text(encoding="utf-8")
            saved_report = json.loads(report_path.read_text(encoding="utf-8"))

            self.assertTrue(report["exact_ordered_text_payload_match"])
            self.assertEqual(
                report["source_text_fragment_sha256"],
                report["emitted_text_fragment_sha256"],
            )
            self.assertEqual(saved_report, report)
            self.assertEqual(
                report["high_furnace_foreman_section_5_3_codes"],
                ["5.3", "5.3.1"],
            )
            self.assertIn("### 5.3 炉况调剂方法", markdown)
            self.assertIn("#### 5.3.1 上部调剂", markdown)
            self.assertIn("<td>条件|一</td>", markdown)
            self.assertIn("<td>动作<br>复核</td>", markdown)


if __name__ == "__main__":
    unittest.main(verbosity=2)
