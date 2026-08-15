from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from tools.build_abc33_handbook64_revision import (
    CHAPTERS,
    HEAT_BATCH_RATE_CRITERION,
    HEAT_BATCH_RATE_FORMULA_ANCHORS,
    HEAT_BATCH_RATE_RULE_IDS,
    HEAT_BATCH_RATE_WEIGHTED_FORMULAS,
    RULES,
    build_document,
    validate_catalog,
    write_mapping_json,
)
from tools.verify_abc33_heat_batch_rate_docx import verify_document


def test_catalog_covers_33_rules_and_64_handbook_chapters() -> None:
    audit = validate_catalog()
    assert audit["rule_count"] == 33
    assert audit["chapter_count"] == 64
    assert audit["mapped_chapter_count"] == 64
    assert audit["unmapped_chapters"] == []
    assert [rule["id"] for rule in RULES[:9]] == [f"A{i}" for i in range(1, 10)]
    assert set(CHAPTERS) == set(range(1, 65))


def test_each_rule_has_mechanism_and_ordered_intervention_steps() -> None:
    for rule in RULES:
        assert str(rule["mechanism"]).strip()
        steps = list(rule["steps"])
        assert len(steps) >= 3
        assert all(str(step).strip() for step in steps)


def test_build_revision_docx_and_mapping_json(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "revision.docx"
    mapping = tmp_path / "mapping.json"
    source_document = Document()
    source_document.add_heading("原规则文档", level=1)
    source_document.add_paragraph("原公式保持不变。")
    source_document.add_paragraph("偏热或偏凉时分别进入B5/B6。")
    formula_table = source_document.add_table(rows=3, cols=1)
    for row, rule_id in zip(formula_table.rows, ("A2", "B4", "B5")):
        row.cells[0].text = HEAT_BATCH_RATE_FORMULA_ANCHORS[rule_id] + ")"
    source_document.save(source)

    audit = build_document(source, output)
    write_mapping_json(mapping)

    assert audit["rule_count"] == 33
    assert output.exists() and output.stat().st_size > 0
    revised = Document(output)
    texts = [paragraph.text for paragraph in revised.paragraphs]
    all_paragraphs = list(revised.paragraphs) + [
        paragraph
        for table in revised.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    assert sum("【新增】形成原理" in text for text in texts) == 33
    assert sum("【新增】干预处置流程" in text for text in texts) == 33
    criterion_paragraphs = [
        paragraph
        for paragraph in all_paragraphs
        if HEAT_BATCH_RATE_CRITERION in paragraph.text
    ]
    assert len(criterion_paragraphs) == len(HEAT_BATCH_RATE_RULE_IDS) * 2 == 6
    assert all(paragraph.runs for paragraph in criterion_paragraphs)
    assert all(run.bold is True for paragraph in criterion_paragraphs for run in paragraph.runs)
    assert "昨日平均料速作为正常料速基线" in HEAT_BATCH_RATE_CRITERION
    assert "MES滚动24 h平均料速" in HEAT_BATCH_RATE_CRITERION
    assert "≤0.25大批/h或≤0.5小批/h" in HEAT_BATCH_RATE_CRITERION
    assert "≤0.5大批/h或≤1小批/h" in HEAT_BATCH_RATE_CRITERION
    assert "后30 min应处于6.5—7.5小批" in HEAT_BATCH_RATE_CRITERION
    assert "连续2 h大批料速高于MES 24 h平均值1大批/h及以上" in HEAT_BATCH_RATE_CRITERION
    verification = verify_document(output)
    assert verification["criterion_paragraph_count"] == 6
    assert verification["fully_bold"] is True
    assert verification["sorted_first_in_rules"] == ["A2", "B4", "B5"]
    assert verification["sorted_first_in_formula_cells"] == ["A2", "B4", "B5"]
    assert verification["formula_weight"] == 20
    assert verification["rebalanced_existing_weight_total"] == 80
    formula_text = "\n".join(cell.text for table in revised.tables for row in table.rows for cell in row.cells)
    assert all(formula in formula_text for formula in HEAT_BATCH_RATE_WEIGHTED_FORMULAS.values())
    assert "原公式保持不变。" in texts
    assert any("偏热或偏凉时分别进入B4/B5" in text for text in texts)
    assert all("偏热或偏凉时分别进入B5/B6" not in text for text in texts)

    payload = json.loads(mapping.read_text(encoding="utf-8"))
    assert len(payload["chapters"]) == 64
    assert len(payload["rules"]) == 33
    assert all(item["primary_rule"] for item in payload["chapters"])
